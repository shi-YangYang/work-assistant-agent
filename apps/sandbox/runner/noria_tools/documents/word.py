from pathlib import Path
import subprocess
from tempfile import TemporaryDirectory
from .blocks import validate_layout, cell_text, check_page_limit
from ..common.paths import output_file, image_path
from ..common.fonts import FONT_NAME, require_font
from ..common.results import result


def rendered_page_count(path):
    from pypdf import PdfReader
    # Only toolkit-generated documents enter this path, inside the isolated runner.
    with TemporaryDirectory(prefix='word-pages-', dir=path.parent) as directory:
        root = Path(directory).resolve()
        try:
            completed = subprocess.run([
                'libreoffice', '-env:UserInstallation=' + (root / 'profile').as_uri(),
                '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(root), str(path.resolve()),
            ], capture_output=True, timeout=25, check=False)
        except (OSError, subprocess.TimeoutExpired) as error:
            raise ValueError('无法渲染DOCX核验页数，未发布文件；请检查执行镜像的LibreOffice') from error
        preview = root / (path.stem + '.pdf')
        if completed.returncode != 0 or not preview.is_file():
            raise ValueError('DOCX渲染失败，无法核验页数，未发布文件')
        return len(PdfReader(preview).pages)


def create_word(*, filename, title, blocks, max_pages=None):
    from docx import Document
    from docx.shared import Inches, Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    require_font(); validate_layout(title, blocks)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    compact = max_pages is not None
    if compact:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Inches(.6)
        section.left_margin = section.right_margin = Inches(.65)
    for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet'):
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style.font.color.rgb = RGBColor.from_string('333333')
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('asciiTheme', 'eastAsiaTheme', 'hAnsiTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + key), None)
        fonts.set(qn('w:eastAsia'), FONT_NAME)
        paragraph = style.element.get_or_add_pPr()
        border = paragraph.find(qn('w:pBdr'))
        if border is not None:
            paragraph.remove(border)
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.space_after = Pt(8)
    if compact:
        for name, size in [('Normal', 11), ('Title', 18), ('Heading 1', 13), ('Heading 2', 12),
                           ('Heading 3', 11), ('List Bullet', 11)]:
            style = document.styles[name]
            style.font.size = Pt(size)
            style.paragraph_format.line_spacing = 1.1
            style.paragraph_format.space_before = Pt(6 if name.startswith('Heading') else 0)
            style.paragraph_format.space_after = Pt(3)
    document.core_properties.title = title
    document.add_heading(title, 0)
    for block in blocks:
        kind = block['type']
        if kind == 'heading':
            document.add_heading(block['text'], block.get('level', 1))
        elif kind == 'paragraph':
            document.add_paragraph(block['text'])
        elif kind == 'list':
            for item in block['items']:
                document.add_paragraph(item, style='List Bullet')
        elif kind == 'table':
            data = block['table']
            table = document.add_table(rows=1, cols=len(data['columns']))
            table.style = 'Table Grid'
            for cell, column in zip(table.rows[0].cells, data['columns'], strict=True):
                cell.text = column
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in data.get('rows', []):
                for cell, value in zip(table.add_row().cells, row, strict=True):
                    cell.text = cell_text(value)
            if compact:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(0)
                            paragraph.paragraph_format.line_spacing = 1.15
                            for run in paragraph.runs:
                                run.font.size = Pt(10.5)
        elif kind == 'image':
            path, width, height = image_path(block['input'])
            scale = min(6.7 / width, 8.3 / height)
            document.add_picture(str(path), width=Inches(width * scale), height=Inches(height * scale))
            if block.get('caption'):
                document.add_paragraph(block['caption'])
    data = {'filename': filename, 'blockCount': len(blocks), 'editable': True}
    with output_file(filename, '.docx') as path:
        document.save(path)
        Document(path)
        if compact:
            pages = rendered_page_count(path)
            check_page_limit(pages, max_pages)
            data.update(pageCount=pages, pageCountRenderer='LibreOffice', maxPages=max_pages, layout='compact')
    warnings = ['DOCX使用镜像中文字体；其他设备的分页可能随字体变化，未声明字体嵌入。']
    if not compact:
        warnings.append('未渲染核验DOCX页数，不能声称满足页数限制。')
    return result(data, warnings)
