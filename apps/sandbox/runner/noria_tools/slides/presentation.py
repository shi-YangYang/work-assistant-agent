"""Conservative, editable layouts with explicit pagination rather than clipping."""
from ..common.paths import output_file, image_path
from ..common.fonts import FONT_NAME, require_font
from ..common.results import result
from ..documents.blocks import cell_text


def lines(text, width_points, size):
    from PIL import ImageFont
    font = ImageFont.truetype(require_font(), size * 4)
    result, current = [], ''
    for character in text.expandtabs(4):
        if character == '\n' or font.getlength(current + character) > width_points * 4:
            result.append(current)
            current = ''
        if character != '\n':
            current += character
    result.append(current)
    return result


def create_slides(*, filename, title, pages):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    require_font()
    deck = Presentation(); deck.slide_width = Inches(13.333); deck.slide_height = Inches(7.5)
    deck.core_properties.title = title
    warnings = []
    def text_box(slide, text, x, y, width, height, size, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame; frame.word_wrap = False
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, line in enumerate(text.split('\n')):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line; paragraph.font.name = FONT_NAME; paragraph.font.size = Pt(size)
            paragraph.font.bold = bold; paragraph.font.color.rgb = RGBColor.from_string('333333')
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = Pt(size * 1.45)
            for tag in ('a:ea', 'a:cs'):
                face = OxmlElement(tag); face.set('typeface', FONT_NAME)
                paragraph.font._rPr.append(face)
        return box
    def new_slide(heading):
        wrapped = lines(heading, 11.9 * 72 - 8, 28)
        if len(wrapped) > 2:
            raise ValueError('幻灯片标题过长，请精简至两行以内')
        if len(deck.slides) >= 120:
            raise ValueError('内容分页超过120页，请拆分演示文稿')
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        text_box(slide, '\n'.join(wrapped), .7, .45, 11.9, 1.2, 28, True)
        text_box(slide, str(len(deck.slides)), 12, 7.05, .5, .25, 10)
        return slide
    top, bottom, gap = 1.9, 6.65, .18
    for page in pages:
        before = len(deck.slides)
        slide, cursor = new_slide(page['title']), top
        for block in page.get('blocks', []):
            kind = block['type']
            if kind in ('paragraph', 'heading', 'list'):
                parts = block['items'] if kind == 'list' else [block['text']]
                content = []
                size = 26 if kind == 'heading' else 24
                for part in parts:
                    content.extend(lines(('• ' if kind == 'list' else '') + part, 11.7 * 72 - 8, size))
                size = 26 if kind == 'heading' else 24
                line_height = size * 1.45 / 72
                cursor += .04
                while content:
                    capacity = int((bottom - cursor) / line_height)
                    if capacity < 1:
                        slide, cursor = new_slide(page['title']), top
                        capacity = int((bottom - cursor) / line_height)
                    chunk, content = content[:capacity], content[capacity:]
                    height = len(chunk) * line_height
                    text_box(slide, '\n'.join(chunk), .8, cursor, 11.7, height, size, kind == 'heading')
                    cursor += height + gap
            elif kind == 'table':
                data = block['table']; columns = data['columns']; width = len(columns)
                if width > 6:
                    raise ValueError('幻灯片表格最多6列；请拆分宽表')
                available = 11.7 * 72 / width - 16
                header = ['\n'.join(lines(value, available, 18)) for value in columns]
                header_height = max(value.count('\n') + 1 for value in header) * .37 + .14
                if header_height > 1.6:
                    raise ValueError('表头过长，请缩短列名')
                rows = []
                for row in data.get('rows', []):
                    texts = ['\n'.join(lines(cell_text(value), available, 18)) for value in row]
                    height = max(value.count('\n') + 1 for value in texts) * .37 + .14
                    if height + header_height > bottom - top:
                        raise ValueError('单行表格内容过长，请拆分内容或使用文档')
                    rows.append((texts, height))
                pending = True
                while pending:
                    required = header_height + (rows[0][1] if rows else 0)
                    if cursor + required > bottom:
                        slide, cursor = new_slide(page['title']), top
                    group, used = [], header_height
                    while rows and cursor + used + rows[0][1] <= bottom:
                        item = rows.pop(0); group.append(item); used += item[1]
                    table = slide.shapes.add_table(len(group) + 1, width, Inches(.8), Inches(cursor), Inches(11.7), Inches(used)).table
                    for index, (values, height) in enumerate([(header, header_height), *group]):
                        table.rows[index].height = Inches(height)
                        for cell, value in zip(table.rows[index].cells, values, strict=True):
                            cell.text = value; cell.margin_top = cell.margin_bottom = Inches(.04)
                            cell.margin_left = cell.margin_right = Inches(.05)
                            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string('EEEEEE' if index == 0 else 'FFFFFF')
                            for paragraph in cell.text_frame.paragraphs:
                                paragraph.font.name = FONT_NAME; paragraph.font.size = Pt(18)
                                paragraph.font.color.rgb = RGBColor.from_string('333333')
                                paragraph.alignment = PP_ALIGN.LEFT
                                paragraph.line_spacing = Pt(18 * 1.45)
                                paragraph.space_before = paragraph.space_after = Pt(0)
                                for tag in ('a:ea', 'a:cs'):
                                    face = OxmlElement(tag); face.set('typeface', FONT_NAME)
                                    paragraph.font._rPr.append(face)
                    cursor += used + gap
                    pending = bool(rows)
            elif kind == 'image':
                path, width, height = image_path(block['input'])
                caption = lines(block.get('caption', ''), 11.7 * 72 - 8, 16)
                if len(caption) > 2:
                    raise ValueError('图片说明过长，请单独放入文字块')
                caption_height = len(caption) * .34 if block.get('caption') else 0
                scale = min(11.7 / width, (bottom - top - caption_height) / height)
                rendered_width, rendered_height = width * scale, height * scale
                if cursor + rendered_height + caption_height > bottom:
                    slide, cursor = new_slide(page['title']), top
                slide.shapes.add_picture(str(path), Inches((13.333 - rendered_width) / 2), Inches(cursor), width=Inches(rendered_width), height=Inches(rendered_height))
                cursor += rendered_height
                if caption_height:
                    text_box(slide, '\n'.join(caption), .8, cursor, 11.7, caption_height, 16)
                    cursor += caption_height
                cursor += gap
        added = len(deck.slides) - before
        if added > 1:
            warnings.append(f'“{page["title"][:80]}”按内容拆成{added}页，未删除内容。')
    with output_file(filename, '.pptx') as path:
        deck.save(path); Presentation(path)
    return result({'filename': filename, 'pageCount': len(deck.slides), 'editable': True}, warnings)
