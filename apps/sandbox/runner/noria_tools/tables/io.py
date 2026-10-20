"""Bounded, loss-aware table I/O. No implicit grouping or row removal."""
import csv
import math
import re
from decimal import Decimal, InvalidOperation
from ..common.paths import input_path, output_file
from ..common.results import result

MAX_ROWS = 20000
MAX_CELLS = 400000


def _number(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        raise ValueError('布尔值不能隐式转换为数值')
    if type(value) is int:
        return value
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError('数值列包含无法转换的内容') from None
    if not number.is_finite():
        raise ValueError('数值必须有限')
    if number == number.to_integral_value():
        return int(number)
    converted = float(number)
    if not math.isfinite(converted) or (converted == 0 and number != 0):
        raise ValueError('数值超出支持的精度范围')
    return converted



def _load_table(data=None, source=None):
    warnings = []
    if data is not None:
        columns, rows = list(data['columns']), [list(row) for row in data.get('rows', [])]
    else:
        path = input_path(source['input'], {'.csv', '.xlsx'})
        if path.suffix.lower() == '.csv':
            if source.get('sheet'):
                raise ValueError('CSV没有工作表，sheet请留空')
            with path.open(encoding=source.get('encoding', 'utf-8-sig'), newline='') as stream:
                reader = csv.reader(stream, strict=True)
                columns = next(reader, None)
                rows = []
                for row in reader:
                    if len(rows) >= MAX_ROWS:
                        raise ValueError('表格超过20000行，请使用run_python按需求分批处理')
                    rows.append(row)
            warnings.append('CSV未指定类型的列保留原文本，不自动推断日期或删除前导零。')
        else:
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                sheet = source.get('sheet') or book.sheetnames[0]
                if sheet not in book.sheetnames:
                    raise ValueError('指定工作表不存在')
                page = book[sheet]
                if page.max_column > 40 or page.max_row > MAX_ROWS + 1:
                    raise ValueError('工作表超过20000行或40列，请使用run_python分批处理')
                values = page.iter_rows(values_only=True)
                columns = list(next(values, []))
                rows = [list(row) for row in values]
                if any(isinstance(cell, str) and cell.startswith('=') for row in [columns, *rows] for cell in row):
                    warnings.append('源文件公式保留为文本，未执行或重新计算。')
            finally:
                book.close()
    return columns, rows, warnings


def _valid_header(columns):
    return (bool(columns) and len(columns) <= 40 and
            all(isinstance(c, str) and c and len(c) <= 120 for c in columns) and
            len(set(columns)) == len(columns))


def _normalize_rows(rows, width, warnings):
    if len(rows) * width > MAX_CELLS:
        raise ValueError('表格单元格超过400000个')
    dates = 0
    for row in rows:
        for index, cell in enumerate(row):
            if hasattr(cell, 'isoformat'):
                row[index] = cell.isoformat()
                dates += 1
            elif isinstance(cell, float) and not math.isfinite(cell):
                raise ValueError('表格包含非有限数字')
            elif cell is not None and type(cell) not in (str, int, float, bool):
                raise ValueError('表格包含不支持的数据类型')
            if isinstance(row[index], str) and len(row[index]) > 2000:
                raise ValueError('单元格超过2000字，请通过Python处理长文本')
    if dates:
        warnings.append(f'{dates}个日期/时间单元格以ISO文本表示。')


def _prepare_table(columns, rows, warnings, source):
    if not _valid_header(columns):
        raise ValueError('表头必须有1至40个不重复的非空文本列名')
    if any(len(row) != len(columns) for row in rows):
        raise ValueError('表格存在列数不一致的行，未跳过任何数据')
    _normalize_rows(rows, len(columns), warnings)
    for column, kind in (source or {}).get('column_types', {}).items():
        if column not in columns:
            raise ValueError('指定类型的列不存在：' + column)
        index = columns.index(column)
        for row in rows:
            value = row[index]
            if value is None or value == '':
                continue
            if kind == 'number':
                row[index] = _number(value)
            elif kind == 'string':
                row[index] = str(value)
            elif kind == 'boolean':
                if type(value) is bool:
                    continue
                if str(value).lower() not in ('true', 'false'):
                    raise ValueError('布尔列只接受true/false，不推断其他内容')
                row[index] = str(value).lower() == 'true'
    return columns, rows, warnings


def read_table(data=None, source=None):
    return _prepare_table(*_load_table(data, source), source)


def inspection_rows(source):
    """One file read; invalid headers are data, never inferred or repaired."""
    columns, rows, warnings = _load_table(source=source)
    if _valid_header(columns):
        columns, rows, warnings = _prepare_table(columns, rows, warnings, source)
        return columns, rows, warnings, True
    if source.get('column_types'):
        raise ValueError('没有有效表头，不能按列名转换类型；请清空column_types查看原始位置')
    # Even a blank first CSV record is retained. Empty files have no record.
    if columns is not None:
        rows.insert(0, columns)
    width = max((len(row) for row in rows), default=0)
    if width > 40 or len(rows) > MAX_ROWS:
        raise ValueError('工作表超过20000行或40列，请使用run_python分批处理')
    _normalize_rows(rows, width, warnings)
    warnings.append('首行不是有效表头，已按原始行列位置检查；首行和空行均保留，未推断列名。')
    return list(range(1, width + 1)), rows, warnings, False


def _csv_text(value):
    if not isinstance(value, str):
        return value, False
    stripped = value.lstrip()
    danger = stripped.startswith(('=', '+', '@', '-')) and not re.fullmatch(r'-?\d+(\.\d+)?', stripped)
    return ("'" + value, True) if danger else (value, False)


def export_table(*, filename, format='xlsx', data=None, source=None):
    columns, rows, warnings = read_table(data, source)
    with output_file(filename, '.' + format) as path:
        if format == 'csv':
            escaped = 0
            with path.open('w', encoding='utf-8-sig', newline='') as stream:
                writer = csv.writer(stream)
                for row in [columns, *rows]:
                    converted = [_csv_text(value) for value in row]
                    escaped += sum(flag for _, flag in converted)
                    writer.writerow([value for value, _ in converted])
            if escaped:
                warnings.append(f'{escaped}个公式样式文本已添加单引号，避免电子表格执行公式。')
        else:
            from openpyxl import Workbook, load_workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            book = Workbook(); sheet = book.active; sheet.title = '数据'
            precise = 0
            for row in [columns, *rows]:
                converted = []
                for value in row:
                    if type(value) is int and len(str(abs(value))) > 15:
                        converted.append(str(value)); precise += 1
                    else:
                        converted.append(value)
                sheet.append(converted)
            if precise:
                warnings.append(f'{precise}个超过Excel精确整数范围的值按文本保存，未舍入。')
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='333333')
            sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
            for index, column in enumerate(columns, 1):
                width = max([len(column), *(len(str(row[index - 1] or '')) for row in rows[:100])])
                sheet.column_dimensions[get_column_letter(index)].width = min(42, max(12, width + 2))
            book.save(path)
            check = load_workbook(path, read_only=True); check.close()
    return result({'filename': filename, 'rowCount': len(rows), 'columns': columns}, warnings)
