"""Treat all container output, including archive metadata, as untrusted."""
import csv
import re
import hashlib
import io
import json
import tarfile
import zipfile
from pathlib import PurePosixPath
from .schemas import MAX_OUTPUT_BYTES

MIMES = {'.txt': 'text/plain', '.md': 'text/markdown', '.json': 'application/json', '.csv': 'text/csv',
         '.png': 'image/png', '.pdf': 'application/pdf',
         '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
         '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation'}


def _validate_office_instruction(instruction, *, word_field=False):
    if re.search(r'(DDE|WEBSERVICE|RTD\s*\(|CALL\s*\(|EXEC\s*\(|HYPERLINK\s*\(|\[[^]]+\])', instruction, re.I):
        raise ValueError('Unsafe formula or field')
    if word_field and re.match(r'\s*(INCLUDETEXT|INCLUDEPICTURE|LINK|DATABASE)\b', instruction, re.I):
        raise ValueError('Unsafe formula or field')


def validate(data, suffix):
    if suffix == '.png':
        from PIL import Image
        with Image.open(io.BytesIO(data)) as image:
            if image.format != 'PNG' or image.width * image.height > 16_000_000:
                raise ValueError('Invalid PNG dimensions')
            image.verify()
        with Image.open(io.BytesIO(data)) as image:
            image.load()
    if suffix == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(io.BytesIO(data), strict=True)
        if pdf.is_encrypted or not 1 <= len(pdf.pages) <= 200:
            raise ValueError('Invalid PDF pages')
        from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject, NameObject
        blocked = {'/JavaScript', '/JS', '/Launch', '/EmbeddedFile', '/EmbeddedFiles', '/OpenAction', '/AA',
                   '/RichMedia', '/RichMediaContent', '/XFA', '/SubmitForm', '/ImportData', '/GoToR', '/GoToE',
                   '/Rendition', '/Sound', '/Movie', '/3D', '/Filespec', '/EF'}
        pending, seen = [pdf.trailer], set()
        while pending:
            node = pending.pop()
            if isinstance(node, IndirectObject):
                token = (node.idnum, node.generation)
                if token in seen:
                    continue
                seen.add(token)
                if len(seen) > 50000:
                    raise ValueError('PDF object graph exceeds limits')
                pending.append(node.get_object())
            elif isinstance(node, DictionaryObject):
                standard_fonts = {'Helvetica', 'Helvetica-Bold', 'Helvetica-Oblique', 'Helvetica-BoldOblique',
                    'Times-Roman', 'Times-Bold', 'Times-Italic', 'Times-BoldItalic', 'Courier', 'Courier-Bold',
                    'Courier-Oblique', 'Courier-BoldOblique', 'Symbol', 'ZapfDingbats'}
                subtype = node.get('/Subtype')
                needs_embedding = subtype in ('/CIDFontType0', '/CIDFontType2', '/TrueType') or (subtype in ('/Type1', '/MMType1') and str(node.get('/BaseFont', '')).lstrip('/') not in standard_fonts)
                if node.get('/Type') == '/Font' and needs_embedding:
                    descriptor = node.get('/FontDescriptor')
                    descriptor = descriptor.get_object() if descriptor else {}
                    if not any(descriptor.get(key) for key in ('/FontFile', '/FontFile2', '/FontFile3')):
                        raise ValueError('PDF fonts must be embedded for portable display')
                if blocked.intersection(node.keys()):
                    raise ValueError('Active PDF content is forbidden')
                if node.get('/URI') and not str(node['/URI']).startswith(('http://', 'https://')):
                    raise ValueError('Unsafe PDF link')
                pending.extend(node.values())
            elif isinstance(node, ArrayObject):
                pending.extend(node)
            elif isinstance(node, NameObject) and str(node) in blocked:
                raise ValueError('Active PDF content is forbidden')
        for page in pdf.pages:
            if float(page.mediabox.width) <= 0 or float(page.mediabox.height) <= 0:
                raise ValueError('Invalid PDF page bounds')
    if suffix in ('.txt', '.md', '.csv', '.json'):
        text = data.decode('utf-8-sig')
        if suffix == '.json':
            json.loads(text)
        if suffix == '.csv':
            for row in csv.reader(io.StringIO(text)):
                for cell in row:
                    value = cell.lstrip()
                    if value.startswith(('=', '+', '-', '@')) and not re.fullmatch(r'[+-]?[0-9]+(?:\.[0-9]+)?', value):
                        raise ValueError('Quote formula-like CSV text with an apostrophe')
    if suffix in ('.xlsx', '.docx', '.pptx'):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 4096 or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
                raise ValueError('Office archive exceeds limits')
            names = {item.filename for item in entries}
            required = {'.xlsx': 'xl/workbook.xml', '.docx': 'word/document.xml', '.pptx': 'ppt/presentation.xml'}[suffix]
            if '[Content_Types].xml' not in names or required not in names:
                raise ValueError('Invalid Office document')
            if len(names) != len(entries):
                raise ValueError('Duplicate Office entries')
            from xml.etree import ElementTree
            for item in entries:
                path = PurePosixPath(item.filename)
                if path.is_absolute() or '..' in path.parts or 'vbaproject' in item.filename.lower() or item.flag_bits & 1:
                    raise ValueError('Unsafe Office document')
                if item.filename.endswith(('.xml', '.rels')):
                    if item.file_size > 16 * 1024 * 1024:
                        raise ValueError('Oversized Office XML')
                    xml = archive.read(item).decode('utf-8-sig')
                    if '<!DOCTYPE' in xml.upper() or '<!ENTITY' in xml.upper():
                        raise ValueError('XML entities are forbidden')
                    root = ElementTree.fromstring(xml)
                    # Complex Word fields can span runs and contain nested fields.
                    field_instructions = []
                    word_namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
                    for node in root.iter():
                        tag = node.tag.rsplit('}', 1)[-1]
                        if tag in ('oleObject', 'altChunk', 'externalLink'):
                            raise ValueError('Embedded active content is forbidden')
                        if tag == 'Relationship' and node.attrib.get('TargetMode') == 'External':
                            if not node.attrib.get('Type', '').endswith('/hyperlink') or not node.attrib.get('Target', '').startswith(('https://', 'http://')):
                                raise ValueError('External document resources are forbidden')
                        if tag == 'fldSimple':
                            _validate_office_instruction(node.attrib.get(word_namespace + 'instr', ''), word_field=True)
                        elif tag == 'fldChar':
                            field_type = node.attrib.get(word_namespace + 'fldCharType')
                            if field_type == 'begin':
                                field_instructions.append([])
                            elif field_type in ('separate', 'end') and field_instructions:
                                instruction = field_instructions[-1]
                                if instruction is not None:
                                    _validate_office_instruction(''.join(instruction), word_field=True)
                                if field_type == 'end':
                                    field_instructions.pop()
                                else:
                                    field_instructions[-1] = None
                        elif tag in ('f', 'instrText'):
                            if tag == 'instrText' and field_instructions and field_instructions[-1] is not None:
                                field_instructions[-1].append(node.text or '')
                            else:
                                _validate_office_instruction(node.text or '', word_field=tag == 'instrText')
                    for instruction in field_instructions:
                        if instruction is not None:
                            _validate_office_instruction(''.join(instruction), word_field=True)

        if suffix == '.xlsx':
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if not 1 <= len(workbook.worksheets) <= 100:
                    raise ValueError('Invalid workbook sheets')
                for sheet in workbook.worksheets:
                    if (sheet.max_row or 0) > 200000 or (sheet.max_column or 0) > 1000:
                        raise ValueError('Workbook dimensions exceed limits')
                    for _ in sheet.iter_rows():
                        pass
            finally:
                workbook.close()
        elif suffix == '.docx':
            from docx import Document
            document = Document(io.BytesIO(data))
            if document.element.tag != '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}document':
                raise ValueError('Invalid Word document')
            list(document.paragraphs)
        elif suffix == '.pptx':
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(data))
            if not 1 <= len(presentation.slides) <= 100 or presentation.slide_width <= 0 or presentation.slide_height <= 0:
                raise ValueError('Invalid presentation')
            for slide in presentation.slides:
                list(slide.shapes)


def export_archive(data, directory, *, validate_files=True):
    total = 0
    files = []
    with tarfile.open(fileobj=io.BytesIO(data)) as archive:
        for entry in archive:
            path = PurePosixPath(entry.name)
            if path.is_absolute() or '..' in path.parts or len(path.parts) > 2:
                raise ValueError('Unsafe output path')
            if entry.isdir() and path.name == 'output':
                continue
            if not entry.isfile() or path.parts[0] != 'output' or len(path.parts) != 2:
                raise ValueError('Only regular output files are allowed')
            if any(ord(c) < 32 for c in path.name) or '\\' in path.name or len(path.name) > 180:
                raise ValueError('Unsafe output filename')
            suffix = path.suffix.lower()
            if suffix not in MIMES:
                raise ValueError('Unsupported output format')
            total += entry.size
            if len(files) >= 12 or entry.size <= 0 or total > MAX_OUTPUT_BYTES:
                raise ValueError('Output exceeds file limits')
            content = archive.extractfile(entry).read(entry.size + 1)
            if len(content) != entry.size:
                raise ValueError('Incomplete output')
            if validate_files:
                validate(content, suffix)
            identifier = hashlib.sha256(content).hexdigest()
            if any(item['name'] == path.name for item in files):
                raise ValueError('Duplicate output filename')
            (directory / identifier).write_bytes(content)
            files.append({'id': identifier, 'name': path.name, 'size': len(content), 'mimeType': MIMES[suffix], 'sha256': identifier})
    return files
