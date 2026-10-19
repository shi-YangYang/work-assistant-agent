"""Deterministic source traversal. Output limits bound a batch, not the source.

The cursor is a stable 3000-character source fragment ordinal. Reopening a
source may scan its prefix; parsing remains inside the disposable worker.
"""
import csv
import io
import json
import zipfile
from pathlib import Path
from lxml import etree
try:
    from .document_parser import ParseFailure, MAX_EXPANDED, office_check, text_decode, MAX_TEXT, MAX_CHUNKS, MAX_PAGES
except ImportError:
    from document_parser import ParseFailure, MAX_EXPANDED, office_check, text_decode, MAX_TEXT, MAX_CHUNKS, MAX_PAGES

XML = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def office_visuals(path):
    with zipfile.ZipFile(path) as archive:
        if any('/media/' in name or '/charts/' in name or '/drawings/' in name for name in archive.namelist()):
            return True
        # Themes and styles contain geometry defaults (e.g. orthographicFront),
        # not visible document objects. Inspect only actual content parts.
        parts = (name for name in archive.namelist() if name.endswith('.xml') and
                 (name == 'word/document.xml' or name.startswith(('word/header', 'word/footer', 'word/footnotes', 'word/endnotes', 'ppt/slides/'))))
        for part in parts:
            root = etree.fromstring(archive.read(part), XML)
            for node in root.iter():
                if not isinstance(node.tag, str):
                    continue
                kind = etree.QName(node).localname
                if kind in ('drawing', 'pict', 'txbxContent', 'cxnSp', 'graphicFrame', 'grpSp'):
                    return True
                if kind == 'prstGeom' and node.get('prst') not in ('rect', 'line', None):
                    return True
        return False


def word_units(path):
    with zipfile.ZipFile(path) as archive:
        parts = ['word/document.xml', *sorted(name for name in archive.namelist()
                 if name.startswith(('word/header', 'word/footer', 'word/footnotes', 'word/endnotes')) and name.endswith('.xml'))]
        for part in parts:
            root = etree.fromstring(archive.read(part), XML)
            label = '正文' if part == 'word/document.xml' else Path(part).stem
            tables = {table: number for number, table in enumerate(root.iter(W + 'tbl'), 1)}
            for index, paragraph in enumerate(root.iter(W + 'p'), 1):
                # Text boxes may contain nested paragraphs. Never duplicate them.
                separators = {W + 'tab': '\t', W + 'br': '\n', W + 'cr': '\n'}
                texts = [node.text or '' if node.tag == W + 't' else separators[node.tag]
                         for node in paragraph.iter()
                         if node.tag in (W + 't', *separators) and
                         (node.tag != W + 'br' or node.get(W + 'type', 'textWrapping') == 'textWrapping') and
                         next(node.iterancestors(W + 'p'), None) is paragraph]
                if texts:
                    location = label
                    cell = next(paragraph.iterancestors(W + 'tc'), None)
                    if cell is not None:
                        row = next(cell.iterancestors(W + 'tr'), None)
                        table = next(cell.iterancestors(W + 'tbl'), None)
                        if row is not None and table is not None:
                            location += f' · 表格 {tables[table]} · 行 {list(table.iterchildren(W + "tr")).index(row) + 1} 列 {list(row.iterchildren(W + "tc")).index(cell) + 1}'
                    yield f'{location} · 段落 {index}', ''.join(texts), {'unit': 'paragraph', 'part': part, 'index': index}


def sheet_units(path, info):
    warnings = info['warnings']
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple
    book = load_workbook(io.BytesIO(Path(path).read_bytes()), read_only=True, data_only=False, keep_links=False)
    try:
        with zipfile.ZipFile(path) as archive:
            strings = []
            if 'xl/sharedStrings.xml' in archive.namelist():
                root = etree.fromstring(archive.read('xl/sharedStrings.xml'), XML)
                strings = [''.join(si.itertext()) for si in root]
            for sheet in book.worksheets:
                if sheet.sheet_state != 'visible':
                    warnings.append(f'工作表「{sheet.title}」已隐藏，未读取')
                    continue
                root = etree.fromstring(archive.read(sheet._worksheet_path), XML)
                hidden_rows = {int(row.get('r')) for row in root.iter(S + 'row') if row.get('hidden') in ('1', 'true')}
                hidden_cols = [(int(col.get('min')), int(col.get('max'))) for col in root.iter(S + 'col') if col.get('hidden') in ('1', 'true')]
                if hidden_rows or hidden_cols:
                    warnings.append(f'工作表「{sheet.title}」隐藏的行／列未读取')
                merges = list(root.iter(S + 'mergeCell'))
                for merge in merges:
                    yield f'工作表「{sheet.title}」· 合并 {merge.get("ref")}', '合并范围仅以左上角单元格表示：' + merge.get('ref'), {'unit': 'sheet', 'sheet': sheet.title}
                content_rows, cells = set(), 0
                # Traverse physical cells, not the untrusted dimension rectangle.
                for cell in root.iter(S + 'c'):
                    address = cell.get('r')
                    row, col = coordinate_to_tuple(address)
                    if row in hidden_rows or any(lo <= col <= hi for lo, hi in hidden_cols):
                        continue
                    value, formula = cell.find(S + 'v'), cell.find(S + 'f')
                    text = value.text or '' if value is not None else ''
                    kind = cell.get('t')
                    if kind == 's' and text:
                        text = strings[int(text)]
                    elif kind == 'inlineStr':
                        text = ''.join(node.text or '' for node in cell.iter(S + 't'))
                    elif kind == 'b':
                        text = 'True' if text == '1' else 'False'
                    elif kind == 'e':
                        text = '错误值：' + text
                    elif kind in (None, 'n') and text :
                        from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
                        from openpyxl.utils.datetime import from_excel
                        style = book._cell_styles[int(cell.get('s', '0'))]
                        number_format = BUILTIN_FORMATS.get(style.numFmtId, '') if style.numFmtId < 164 else book._number_formats[style.numFmtId - 164]
                        if is_timedelta_format(number_format):
                            text = str(from_excel(float(text), book.epoch, timedelta=True))
                        elif is_date_format(number_format):
                            text = from_excel(float(text), book.epoch).isoformat()
                    if formula is not None:
                        if text:
                            text = '=' + (formula.text or '') + '；文件保存值（非实时计算）：' + text
                        else:
                            text = '=' + (formula.text or '') + '；文件未保存计算结果，未计算'
                            warning = '部分公式缺少保存值；未执行公式或重算'
                            if warning not in warnings:
                                warnings.append(warning)
                    if text:
                        content_rows.add(row)
                        cells += 1
                        yield f'工作表「{sheet.title}」!{address}', f'{address}：{text}', {'unit': 'cell', 'sheet': sheet.title, 'cell': address}
                # Reached only after the entire visible sheet was traversed.
                # Do not turn an unfinished batch or a styled dimension into
                # a full-sheet statistic; no second source scan is required.
                structure = info.setdefault('tableStructure', {
                    'basis': '完整遍历可见单元格值或公式；行间数量仅计首末内容行之间无值的可见行。隐藏内容未读；合并表不判空行，格式不算内容。',
                    'sheets': [], 'omittedSheets': 0,
                })
                if len(structure['sheets']) < 20:
                    first, last = min(content_rows, default=None), max(content_rows, default=None)
                    hidden_between = sum(first < row < last for row in hidden_rows) if content_rows else 0
                    structure['sheets'].append({
                        'sheet': sheet.title, 'firstContentRow': first, 'lastContentRow': last,
                        'visibleContentRows': len(content_rows), 'visibleContentCells': cells,
                        'visibleRowsWithoutValuesBetween': (last - first + 1 - len(content_rows) - hidden_between if content_rows else 0) if not merges else None,
                        'mergedCellsPresent': bool(merges),
                        'hiddenRowsExcluded': bool(hidden_rows), 'hiddenColumnsExcluded': bool(hidden_cols),
                    })
                else:
                    structure['omittedSheets'] += 1
    finally:
        book.close()


def source_units(path, suffix, info):
    warnings = info['warnings']
    if suffix in ('.docx', '.xlsx', '.pptx'):
        office_check(path, suffix)
        info['hasVisuals'] = office_visuals(path)
    if suffix == '.docx':
        with zipfile.ZipFile(path) as archive:
            if any(name.startswith('word/comments') for name in archive.namelist()):
                warnings.append('批注不计入正文范围；需在原文件中查看')
            root = etree.fromstring(archive.read('word/document.xml'), XML)
            if any(etree.QName(node).localname in ('altChunk', 'delText', 'sym') for node in root.iter() if isinstance(node.tag, str)):
                warnings.append('部分嵌入内容、修订删除内容或特殊符号未作为文字提取；可查看页面或原文件')
        yield from word_units(path)
    elif suffix == '.xlsx':
        yield from sheet_units(path, info)
    elif suffix == '.pptx':
        from pptx import Presentation
        slides = Presentation(path).slides
        info['totalPages'] = len(slides)
        def texts(shapes):
            for shape in shapes:
                if hasattr(shape, 'shapes'):
                    yield from texts(shape.shapes)
                if shape.has_text_frame:
                    yield shape.text
                if shape.has_table:
                    for n, row in enumerate(shape.table.rows, 1):
                        yield f'表格行 {n}：' + ' | '.join(f'列 {i}：{c.text}' for i, c in enumerate(row.cells, 1))
        for number, slide in enumerate(slides, 1):
            value = '\n'.join(texts(slide.shapes))
            yield f'幻灯片 {number}', value, {'unit': 'slide', 'page': number}
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                yield f'幻灯片 {number} · 演讲备注', slide.notes_slide.notes_text_frame.text, {'unit': 'slide', 'page': number}
    elif suffix == '.pdf':
        from pypdf import PdfReader, overwrite_configuration
        from pypdf.errors import LimitReachedError
        overwrite_configuration(zlib_maximum_output_length=MAX_EXPANDED, lzw_maximum_output_length=MAX_EXPANDED, run_length_maximum_output_length=MAX_EXPANDED, array_based_stream_maximum_output_length=MAX_EXPANDED, disable_legacy_handling=True)
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ParseFailure('不支持加密 PDF，请先解密')
        info['totalPages'] = len(reader.pages)
        for number, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text() or ''
            except LimitReachedError:
                raise ParseFailure('PDF 内容解压后超过限制，请拆分文件') from None
            if not text.strip():
                warning = '部分页面无可提取文字，可通过页面图片读取'
                if warning not in warnings: warnings.append(warning)
            resources = page.get('/Resources', {})
            if hasattr(resources, 'get_object'):
                resources = resources.get_object()
            visual = not text.strip() or bool(resources.get('/XObject'))
            if not visual:
                from pypdf.generic import ContentStream
                contents = page.get_contents()
                visual = contents is not None and any(operator in (b'm', b'l', b'c', b're') for _, operator in ContentStream(contents, reader).operations)
            if visual:
                info['hasVisuals'] = True
                if len(info.setdefault('visualPages', [])) < 100:
                    info['visualPages'].append(number)
            yield f'第 {number} 页', text, {'unit': 'page', 'page': number, 'visual': visual}
    else:
        text = text_decode(Path(path).read_bytes())
        if suffix == '.json':
            try:
                value = json.loads(text, parse_constant=lambda _: (_ for _ in ()).throw(ValueError()))
            except (ValueError, RecursionError):
                raise ParseFailure('JSON 语法无效，请修正后重新上传') from None
            def visit(value, pointer='', depth=0):
                if depth > 100:
                    raise ParseFailure('JSON 嵌套超过 100 层，请简化结构')
                if isinstance(value, (dict, list)) and value:
                    yield 'JSON ' + (pointer or '/'), '{对象}' if isinstance(value, dict) else '[数组]', {'unit': 'json', 'pointer': pointer}
                    for key, child in (value.items() if isinstance(value, dict) else enumerate(value)):
                        yield from visit(child, pointer + '/' + str(key).replace('~', '~0').replace('/', '~1'), depth + 1)
                else:
                    yield 'JSON ' + (pointer or '/'), json.dumps(value, ensure_ascii=False), {'unit': 'json', 'pointer': pointer}
            yield from visit(value)
        elif suffix == '.csv':
            # Source upload and parser memory remain bounded; fields are split below.
            csv.field_size_limit(20 * 1024 * 1024)
            reader = csv.reader(io.StringIO(text, newline=''), strict=True)
            for number, row in enumerate(reader, 1):
                yield f'记录 {number} · 至行 {reader.line_num}', '\n'.join(f'列 {i}：{v}' for i, v in enumerate(row, 1)), {'unit': 'record', 'index': number}
        else:
            lines = text.splitlines()
            for start in range(0, len(lines), 30):
                yield f'行 {start + 1}–{min(start + 30, len(lines))}', '\n'.join(lines[start:start + 30]), {'unit': 'line', 'start': start + 1, 'end': min(start + 30, len(lines))}


def extract_batch(path, suffix, start=0):
    if not isinstance(start, int) or not 0 <= start <= 10_000_000:
        raise ValueError('Invalid extraction cursor')
    info = {'warnings': [], 'schema': 1, 'hasVisuals': False, 'sourceComplete': False, 'nextExtractCursor': None}
    chunks, count, ordinal, pages = [], 0, 0, set()
    finished = True
    for location, text, source in source_units(path, suffix, info):
        # Preserve an empty page's position and make scanned pages discoverable.
        if not text.strip() and source['unit'] not in ('page', 'slide'):
            continue
        for offset in range(0, max(1, len(text)), 3000):
            part = text[offset:offset + 3000]
            current = ordinal
            ordinal += 1
            if current < start:
                continue
            page = source.get('page')
            if count + len(part) > MAX_TEXT or len(chunks) >= MAX_CHUNKS or page and page not in pages and len(pages) >= MAX_PAGES:
                info['nextExtractCursor'] = current
                finished = False
                break
            if page:
                pages.add(page)
            label = location + (f' · 字符 {offset + 1}–{offset + len(part)}' if len(text) > 3000 else '')
            chunks.append({'ordinal': current, 'location': label[:300], 'text': part, 'source': source})
            count += len(part)
        if not finished:
            break
    if finished:
        info['sourceEnd'] = ordinal
    info.update(sourceComplete=finished, extractedRange=[start, chunks[-1]['ordinal'] + 1 if chunks else start], characters=count, chunks=len(chunks), scope='已提取文字；页面图片通过视觉读取补充，隐藏区域不纳入范围')
    if not finished:
        info['warnings'].append('部分内容尚未读取，可继续读取后续范围')
    visual_only = chunks and not any(chunk['text'].strip() for chunk in chunks)
    status = 'partial' if info['warnings'] or visual_only else 'ready' if chunks else 'empty'
    return {'status': status, 'chunks': chunks, 'info': info}
