"""Release-only real Writer/Calc/Impress checks, before switching live services."""
import asyncio
from pathlib import Path
import subprocess
import tempfile
from app.integrations.parsing.office_conversion import converter, environment, parse_office


async def check():
    executable, version = converter()
    if not executable:
        raise RuntimeError('LibreOffice is required in the worker image')
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    with tempfile.TemporaryDirectory(prefix='noria-office-preflight-') as directory:
        root = Path(directory)
        doc = Document(); doc.add_paragraph('NORIA-OFFICE-READY'); doc.save(root / 'check.docx')
        book = Workbook(); book.active['A1'] = 'NORIA-OFFICE-READY'; book.save(root / 'check.xlsx')
        slides = Presentation(); slides.slides.add_slide(slides.slide_layouts[0]).shapes.title.text = 'NORIA-OFFICE-READY'; slides.save(root / 'check.pptx')
        for modern, old in (('docx', 'doc'), ('xlsx', 'xls'), ('pptx', 'ppt')):
            profile = root / ('profile-' + old)
            exported = await asyncio.to_thread(subprocess.run, [executable, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', old, '--outdir', str(root), str(root / ('check.' + modern))], capture_output=True, timeout=30, env=environment())
            path = root / ('check.' + old)
            if exported.returncode or not path.is_file():
                raise RuntimeError('LibreOffice legacy export check failed: ' + old)
            result = await parse_office(path, '.' + old)
            if result['status'] not in ('ready', 'partial') or 'NORIA-OFFICE-READY' not in str(result['chunks']):
                raise RuntimeError('LibreOffice conversion/read check failed: ' + old)
    print('Office conversion ready: ' + version)


if __name__ == '__main__':
    asyncio.run(check())
