"""Small synthetic media only: no user files or network are used."""
from datetime import date
import io
from pathlib import Path
import subprocess
import zipfile
from PIL import Image, ImageDraw


def image_samples():
    from pillow_heif import register_heif_opener
    register_heif_opener(thumbnails=False)
    result = {}
    transparent = Image.new('RGBA', (420, 180), (0, 0, 0, 0))
    ImageDraw.Draw(transparent).text((20, 20), 'TRANSPARENT TEXT 12345', fill='black')
    rotated = Image.new('RGB', (160, 80), 'white')
    ImageDraw.Draw(rotated).rectangle((0, 0, 79, 79), fill='red')
    exif = Image.Exif(); exif[274] = 6; exif[270] = 'private metadata'
    output = io.BytesIO(); rotated.save(output, 'JPEG', exif=exif); result['rotated.jpg'] = output.getvalue()
    for name, image, kind in [('transparent.png', transparent, 'PNG'), ('phone.heic', rotated, 'HEIF')]:
        output = io.BytesIO(); image.save(output, kind); result[name] = output.getvalue()
    for name, height in [('screenshot.png', 1800), ('long.png', 9000)]:
        image = Image.new('RGB', (500, height), 'white'); drawing = ImageDraw.Draw(image)
        for y in range(0, height, 25): drawing.text((12, y), f'Row {y // 25 + 1:04d}: Small text 123456789', fill='black')
        output = io.BytesIO(); image.save(output, 'PNG'); result[name] = output.getvalue()
    return result


def spreadsheet_bytes():
    from openpyxl import Workbook
    workbook = Workbook(); sheet = workbook.active; sheet.title = '项目进度'
    sheet.append(['任务', '完成量', '日期', '缓存公式', '无缓存公式', '错误'])
    sheet.append(['界面优化', 3, date(2026, 9, 16), '=B2+2', '=B2+7', '#DIV/0!'])
    sheet['A4'] = '隐藏行秘密'; sheet.row_dimensions[4].hidden = True
    sheet['G2'] = '隐藏列秘密'; sheet.column_dimensions['G'].hidden = True
    sheet.merge_cells('A6:C6'); sheet['A6'] = '合并说明'
    hidden = workbook.create_sheet('隐藏资料'); hidden['A1'] = '隐藏表秘密'; hidden.sheet_state = 'hidden'
    output = io.BytesIO(); workbook.save(output)
    # Saved cache5 for D2; E2 intentionally missing. Bad dimensions must not hide data.
    final = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as source, zipfile.ZipFile(final, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<dimension ref="A1:G6"/>', b'<dimension ref="A1:A1"/>').replace(b'<f>B2+2</f><v></v>', b'<f>B2+2</f><v>5</v>')
            target.writestr(item.filename, data)
    return final.getvalue()


def mp3_bytes(directory: Path, ffmpeg='ffmpeg'):
    path = directory / 'voice.mp3'
    subprocess.run([ffmpeg, '-nostdin', '-hide_banner', '-loglevel', 'error', '-f', 'lavfi', '-i', 'sine=frequency=440:duration=0.25', '-c:a', 'libmp3lame', '-y', str(path)], check=True, timeout=15)
    return path.read_bytes()
