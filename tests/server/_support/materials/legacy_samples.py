"""Self-authored fixtures, exported through real LibreOffice import/export filters.

No staff documents or renamed ZIPs. XLS/PPT/DOC are actual Office 97 containers.
Generate once outside measured runs; all answers are fixed and independently known.
"""
from functools import lru_cache
import io
from pathlib import Path
import subprocess
import tempfile

TEXT = 'SPEC071 金额 1234.50 到期 2026-10-16'


@lru_cache(maxsize=1)
def legacy_samples():
    from app.integrations.parsing.office_conversion import converter, environment
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    executable, _ = converter()
    if not executable:
        raise RuntimeError('Real LibreOffice required for legacy fixtures')
    with tempfile.TemporaryDirectory(prefix='spec071-fixtures-') as directory:
        root = Path(directory)
        doc = Document(); doc.add_paragraph(TEXT); table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '事项'; table.cell(0, 1).text = '负责人'
        table.cell(1, 0).text = '材料核对'; table.cell(1, 1).text = '小林'
        doc.save(root / 'sample.docx')
        workbook = Workbook(); sheet = workbook.active; sheet.title = '金额'
        sheet.append([TEXT, '金额', '日期', '公式'])
        from datetime import datetime
        sheet.append(['支出', 1234.50, datetime(2026, 10, 16), '=B2*2'])
        sheet['B2'].number_format = '0.00'; sheet['C2'].number_format = 'yyyy-mm-dd'
        sheet['E2'] = '=1=1'; sheet['F2'] = '=1/0'; sheet['G2'] = '="缓存中文"'
        sheet.merge_cells('A4:B4'); sheet['A4'] = '合并说明'
        other = workbook.create_sheet('后续'); other['A1'] = '第二工作表标记'
        hidden = workbook.create_sheet('隐藏'); hidden.sheet_state = 'hidden'; hidden['A1'] = '不可见标记'
        workbook.save(root / 'sample.xlsx')
        slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1]); slide.shapes.title.text = TEXT
        slide.placeholders[1].text = '演示正文'; slide.notes_slide.notes_text_frame.text = '备注标记'
        slides.save(root / 'sample.pptx')
        result = {}
        for source, targets in [('docx', ('doc', 'rtf', 'odt')), ('xlsx', ('xls', 'ods')), ('pptx', ('ppt', 'odp'))]:
            for target in targets:
                profile = root / ('profile-' + target)
                output = subprocess.run([executable, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', target, '--outdir', str(root), str(root / ('sample.' + source))], capture_output=True, env=environment(), timeout=30)
                path = root / ('sample.' + target)
                if output.returncode or not path.is_file():
                    raise RuntimeError('Could not generate genuine ' + target + ' fixture')
                result['sample.' + target] = path.read_bytes()
        return result


def stale_xls(data, value=42.125):
    import olefile
    import struct
    raw = io.BytesIO(data)
    with olefile.OleFileIO(raw, write_mode=True) as ole:
        stream = bytearray(ole.openstream('Workbook').read()); offset = 0
        while offset + 4 <= len(stream):
            kind, size = struct.unpack_from('<HH', stream, offset)
            if kind == 6 and stream[offset + 16:offset + 18] != b'\xff\xff':
                struct.pack_into('<d', stream, offset + 10, value)
            offset += size + 4
        ole.write_stream('Workbook', bytes(stream))
    return raw.getvalue()


def xls_formula_cache(data, *, value=None):
    """Change only the first numeric FORMULA's saved result in genuine BIFF8."""
    import olefile
    import struct
    raw = io.BytesIO(data)
    with olefile.OleFileIO(raw, write_mode=True) as ole:
        stream = bytearray(ole.openstream('Workbook').read()); offset = 0
        while offset + 4 <= len(stream):
            kind, size = struct.unpack_from('<HH', stream, offset)
            if kind == 6 and stream[offset + 16:offset + 18] != b'\xff\xff':
                if value is None:
                    # BIFF8's explicit empty cached result, not a numeric zero.
                    stream[offset + 10:offset + 18] = b'\x03\x00\x00\x00\x00\x00\xff\xff'
                else:
                    struct.pack_into('<d', stream, offset + 10, value)
                ole.write_stream('Workbook', bytes(stream))
                return raw.getvalue()
            offset += size + 4
    raise AssertionError('Fixture must contain a numeric formula cache')


def ods_formula_cache(data, *, value=None):
    """Change only one formula cache, preserving its formula and ODF container."""
    import zipfile
    from lxml import etree
    office = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
    table = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
    text = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, 'w') as target:
        for entry in source.infolist():
            raw = source.read(entry)
            if entry.filename == 'content.xml':
                root = etree.fromstring(raw)
                cell = root.find('.//*[@{' + table + '}formula]')
                assert cell is not None, 'Fixture must contain a formula'
                for key in list(cell.attrib):
                    if key.startswith('{' + office + '}') and 'value' in key:
                        del cell.attrib[key]
                for child in list(cell):
                    cell.remove(child)
                if value is not None:
                    cell.set('{' + office + '}value-type', 'float')
                    cell.set('{' + office + '}value', str(value))
                    etree.SubElement(cell, '{' + text + '}p').text = str(value)
                raw = etree.tostring(root, encoding='UTF-8', xml_declaration=True)
            target.writestr(entry, raw)
    return output.getvalue()


def ods_repeated_tail(data, *, formula=False):
    """Real ODF package with either harmless blank padding or formula expansion."""
    import zipfile
    from lxml import etree
    office = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
    table = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, 'w') as target:
        for entry in source.infolist():
            raw = source.read(entry)
            if entry.filename == 'content.xml':
                root = etree.fromstring(raw)
                sheet = root.find('.//{' + table + '}table')
                assert sheet is not None
                row = etree.SubElement(sheet, '{' + table + '}table-row')
                row.set('{' + table + '}number-rows-repeated', '500001' if formula else '1000000')
                cell = etree.SubElement(row, '{' + table + '}table-cell')
                if formula:
                    cell.set('{' + table + '}formula', 'of:=1')
                    cell.set('{' + office + '}value-type', 'float')
                    cell.set('{' + office + '}value', '1')
                else:
                    cell.set('{' + table + '}number-columns-repeated', '1024')
                raw = etree.tostring(root, encoding='UTF-8', xml_declaration=True)
            target.writestr(entry, raw)
    return output.getvalue()


def ods_boolean_cache(data, *, numeric=False, boolean_value='true'):
    """Explicitly select ODF's two valid encodings of a saved TRUE result."""
    import zipfile
    from lxml import etree
    office = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
    table = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
    calc = 'urn:org:documentfoundation:names:experimental:calc:xmlns:calcext:1.0'
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, 'w') as target:
        for entry in source.infolist():
            raw = source.read(entry)
            if entry.filename == 'content.xml':
                root = etree.fromstring(raw)
                cells = root.xpath('//*[@table:formula="of:=1=1"]', namespaces={'table': table})
                assert len(cells) == 1
                cell = cells[0]
                for key in list(cell.attrib):
                    if key.startswith('{' + office + '}') and 'value' in key:
                        del cell.attrib[key]
                kind = 'float' if numeric else 'boolean'
                cell.set('{' + office + '}value-type', kind)
                cell.set('{' + calc + '}value-type', kind)
                if numeric:
                    cell.set('{' + office + '}value', '1')
                elif boolean_value is not None:
                    cell.set('{' + office + '}boolean-value', boolean_value)
                if boolean_value in ('false', '0'):
                    cell.set('{' + table + '}formula', 'of:=1=0')
                    for paragraph in cell:
                        paragraph.text = 'FALSE'
                raw = etree.tostring(root, encoding='UTF-8', xml_declaration=True)
            target.writestr(entry, raw)
    return output.getvalue()


def ods_boolean_saved_text(data):
    """Oracle for the genuine fixture: 7.4 stores 1, newer Calc stores true."""
    import zipfile
    from lxml import etree
    office = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
    table = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = etree.fromstring(archive.read('content.xml'))
    cells = root.xpath('//*[@table:formula="of:=1=1"]', namespaces={'table': table})
    assert len(cells) == 1
    cell = cells[0]
    if cell.get('{' + office + '}value-type') == 'boolean':
        assert cell.get('{' + office + '}boolean-value') == 'true'
        return 'True'
    assert cell.get('{' + office + '}value-type') == 'float'
    assert cell.get('{' + office + '}value') == '1'
    return '1'


@lru_cache(maxsize=1)
def formula_boundary_samples():
    """Real Calc exports, with deterministic original-cache oracles for tests.

    The numeric snapshots of NOW/RAND deliberately differ from a fresh result.
    Conversion may preserve them or fail; it must never silently publish a new
    value. Unknown functions carry Calc's real saved error, not mocked parsing.
    """
    from app.integrations.parsing.office_conversion import converter, environment
    from openpyxl import Workbook
    executable, _ = converter()
    if not executable:
        raise RuntimeError('Real LibreOffice required for formula fixtures')
    result = {}
    with tempfile.TemporaryDirectory(prefix='spec071-formulas-') as directory:
        root = Path(directory)
        for name, formula, expected in (
            ('now', '=NOW()', 40123.5),
            ('rand', '=RAND()', 0.125),
            ('unknown', '=NORIA_UNKNOWN_FUNCTION(1)', '#NAME?'),
        ):
            book = Workbook(); sheet = book.active; sheet.title = '公式边界'
            sheet['A1'] = 'SPEC071 原始保存值'; sheet['A2'] = formula
            source = root / (name + '.xlsx'); book.save(source)
            for suffix in ('xls', 'ods'):
                profile = root / ('profile-' + name + '-' + suffix)
                completed = subprocess.run([executable, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', suffix, '--outdir', str(root), str(source)], capture_output=True, env=environment(), timeout=30)
                path = root / (name + '.' + suffix)
                if completed.returncode or not path.is_file():
                    raise RuntimeError('Could not generate genuine formula fixture ' + path.name)
                data = path.read_bytes()
                if name != 'unknown':
                    update = xls_formula_cache if suffix == 'xls' else ods_formula_cache
                    data = update(data, value=expected)
                result[name + '.' + suffix] = (data, expected)
        normal = legacy_samples()
        result['missing.xls'] = (xls_formula_cache(normal['sample.xls']), None)
        result['missing.ods'] = (ods_formula_cache(normal['sample.ods']), None)
        result['stale.ods'] = (ods_formula_cache(normal['sample.ods'], value=42.125), 42.125)
    return result
