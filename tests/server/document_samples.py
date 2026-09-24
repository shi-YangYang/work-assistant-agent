"""Small reproducible non-sensitive samples; never depend on artifacts or Office."""
import io
import json
from pathlib import Path
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject



def pdf_bytes(texts=('Demo progress', 'Later page: agreed delivery Friday')):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 14 Tf 40 740 Td (' + text.encode('ascii').replace(b'(', b'\\(').replace(b')', b'\\)') + b') Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO(); writer.write(output)
    return output.getvalue()


def samples():
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    document = Document()
    document.add_paragraph('演示项目：资料整理。Demo project review.')
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip([c for row in table.rows for c in row.cells], ['事项', '进展', '样本核对', '等待确认']):
        cell.text = value
    document.add_paragraph('后续安排：周五核对交付范围。')
    word = io.BytesIO(); document.save(word)
    deck = Presentation()
    for title in ('演示项目', '后半部分：下周安排'):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = title
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(7), Inches(1)).table
        table.cell(0, 0).text, table.cell(1, 1).text = '负责人', '待确认'
        slide.notes_slide.notes_text_frame.text = '演讲备注：这是非敏感演示材料，尚未完成工作。'
    ppt = io.BytesIO(); deck.save(ppt)
    return {
        '演示项目.pdf': pdf_bytes(),
        '演示项目.docx': word.getvalue(),
        '演示项目.pptx': ppt.getvalue(),
        '演示项目.txt': ('项目资料\n' * 95 + '后半部分标记：周五讨论交付范围，不代表已完成。\n').encode(),
        '演示项目.json': json.dumps({'项目': {'名称': '演示计划', '安排': ['核对材料', '周五讨论']}, '已完成': False}, ensure_ascii=False).encode(),
        '演示项目.md': '# 演示项目\n\n材料已经整理，等待确认。\n<script>不会执行</script>\n'.encode(),
        '演示项目.csv': '事项,状态,下一步\n"资料,样本",等待确认,周五讨论\n安全示例,未执行,=1+1\n'.encode(),
    }


def export_samples(directory):
    directory = Path(directory); directory.mkdir(parents=True, exist_ok=True)
    for name, data in samples().items():
        (directory / name).write_bytes(data)
