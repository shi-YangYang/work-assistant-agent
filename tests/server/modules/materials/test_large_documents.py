"""Large image-heavy Office input through the real upload/worker/page pipeline."""
import hashlib
import io
import json
import random
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from app.agent.context.document_input import request_pages, assemble_pages
from app.integrations.parsing.office_conversion import converter
from app.integrations.parsing.process import parse_process
from app.tasks.processing.documents import prepare_document
from _support.http.management import conversation, message
from _support.materials.documents import upload, running_context

pytestmark = pytest.mark.asyncio


def large_deck(path, count=6):
    deck = Presentation()
    random_source = random.Random(740)
    for index in range(count):
        picture = io.BytesIO()
        with Image.frombytes('RGB', (2400, 2400), random_source.randbytes(2400 * 2400 * 3)) as image:
            image.save(picture, 'PNG')
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(picture, Inches(1), Inches(1), height=Inches(5))
        slide.shapes.add_textbox(Inches(1), Inches(0), Inches(7), Inches(1)).text = f'SLIDE-{index + 1}: amount {100 + index}'
    deck.save(path)
    if count == 6:
        assert 97 * 1024 * 1024 < path.stat().st_size < 100 * 1024 * 1024


async def test_image_heavy_99mb_pptx_text_pages_and_reuse(setup, tmp_path, monkeypatch):
    if not converter()[0]:
        pytest.skip('Requires real LibreOffice for the large document pipeline')
    settings, sessions, users, clients = setup
    source = tmp_path / 'large.pptx'
    large_deck(source)
    with source.open('rb') as stream:
        original = hashlib.file_digest(stream, 'sha256').hexdigest()
    client = clients['employee']
    with source.open('rb') as stream:
        item = await upload(client, 'large.pptx', stream)
    assert item['extraction']['status'] == 'unsent'
    sent = await message(client, await conversation(client), '核对最后一页金额与图片', [item['id']])
    context = await running_context(settings, sessions, users['employee'], sent)
    prepared = await prepare_document(context, item['id'])
    assert prepared['extraction']['status'] in ('ready', 'partial'), prepared
    text = (await client.get('/api/v1/uploads/' + item['id'] + '/extraction')).json()
    # Inspect all small page chunks, not just the first API preview page.
    from app.agent.tools.resources.documents import read_document_result
    received = json.loads(await read_document_result(context, item['id'], start=0))
    assert 'SLIDE-6: amount 105' in str(received) and text['items']
    pages = await request_pages(context, item['id'], [1, 6])
    assert pages['status'] == 'ready' and pages['totalPages'] == 6, pages
    blocks, supplied = await assemble_pages(context, [{'role': 'tool', 'content': json.dumps(pages)}])
    images = [part for block in blocks if isinstance(block.get('content'), list) for part in block['content'] if part.get('type') == 'image_url']
    assert len(images) == 2 and len(supplied) == 2
    assert all(part['image_url']['url'].startswith('data:image/') for part in images)
    monkeypatch.setattr('app.modules.attachments.document_assets.parse_office', AsyncMock(side_effect=AssertionError('Do not rerender saved pages')))
    assert (await request_pages(context, item['id'], [1, 6]))['pageRefs'] == pages['pageRefs']
    with (settings.media_dir / item['id']).open('rb') as stream:
        assert hashlib.file_digest(stream, 'sha256').hexdigest() == original


async def test_expansion_bomb_and_large_xml_fail_without_harming_next_parse(tmp_path):
    from _support.materials.document_samples import samples
    original = samples()['演示项目.pptx']
    for name, compression in [('ppt/media/bomb.png', zipfile.ZIP_DEFLATED), ('ppt/large.xml', zipfile.ZIP_STORED)]:
        path = tmp_path / 'bad.pptx'
        path.write_bytes(original)
        with zipfile.ZipFile(path, 'a', compression=compression) as archive:
            archive.writestr(name, b'x' * (17 * 1024 * 1024))
        result = await parse_process(path, '.pptx')
        assert result['status'] == 'failed' and result['info']['error'], result
    path.write_bytes(original)
    assert (await parse_process(path, '.pptx'))['status'] in ('ready', 'partial')


async def test_legacy_ppt_above_old_limit_still_converts_and_reads(tmp_path):
    import asyncio
    from app.integrations.parsing.office_conversion import parse_office, environment
    executable, _ = converter()
    if not executable:
        pytest.skip('Requires real LibreOffice')
    modern = tmp_path / 'large.pptx'
    large_deck(modern, count=2)
    profile = (tmp_path / 'fixture-profile').as_uri()
    process = await asyncio.create_subprocess_exec(executable, '-env:UserInstallation=' + profile, '--headless', '--convert-to', 'ppt:MS PowerPoint 97', '--outdir', str(tmp_path), str(modern), env=environment(), stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
    try:
        await asyncio.wait_for(process.wait(), 60)
    finally:
        if process.returncode is None:
            process.kill()
        await process.wait()
    source = tmp_path / 'large.ppt'
    assert process.returncode == 0 and 30 * 1024 * 1024 < source.stat().st_size < 100 * 1024 * 1024
    result = await parse_office(source, '.ppt')
    assert result['status'] in ('ready', 'partial'), result
    assert 'SLIDE-2: amount 101' in str(result['chunks'])
