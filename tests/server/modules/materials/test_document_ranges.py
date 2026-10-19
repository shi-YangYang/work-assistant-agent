"""Original-source boundaries, not enlarged output caps or model claims."""
import io
import zipfile
from datetime import date, time, timedelta
from pathlib import Path

import pytest
from _support.materials.spec072_samples import fixture_values
from app.integrations.parsing.document_ranges import extract_batch
from app.integrations.parsing.process import parse_process


@pytest.fixture(scope='module')
def source_files(tmp_path_factory):
    root = tmp_path_factory.mktemp('spec072-source')
    for name, data in fixture_values().items():
        (root / name).write_bytes(data)
    return root


@pytest.mark.asyncio
@pytest.mark.parametrize('name', ['records.txt', 'records.docx', 'pages.pdf', 'pages.pptx'])
async def test_real_process_continues_past_old_source_caps_without_changing_ordinals(source_files, name):
    path = source_files / name
    first = await parse_process(path, path.suffix)
    assert first['status'] == 'partial' and not first['info']['sourceComplete'], first
    assert 'Z9P4T7' not in str(first['chunks'])
    cursor = first['info']['nextExtractCursor']
    second = await parse_process(path, path.suffix, start=cursor)
    assert second['chunks'][0]['ordinal'] == cursor
    assert second['info']['sourceComplete'] and second['info']['nextExtractCursor'] is None
    assert 'Z9P4T7' in str(second['chunks']) and '6842' in str(second['chunks'])
    assert max(c['ordinal'] for c in first['chunks']) + 1 == min(c['ordinal'] for c in second['chunks'])
    again = await parse_process(path, path.suffix, start=cursor)
    assert again['chunks'] == second['chunks']


def test_long_single_paragraph_cell_and_more_than_2000_paragraphs(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    long = ''.join(f'{i:07d}abcdefgh' for i in range(16000)) + 'FINAL-SINGLE-PARAGRAPH'
    path = tmp_path / 'long.txt'; path.write_text(long)
    first = extract_batch(path, '.txt')
    second = extract_batch(path, '.txt', first['info']['nextExtractCursor'])
    assert ''.join(c['text'] for c in first['chunks'] + second['chunks']) == long
    word = Document()
    for i in range(2002): word.add_paragraph(f'P{i}')
    path = tmp_path / 'many.docx'; word.save(path)
    first = extract_batch(path, '.docx'); second = extract_batch(path, '.docx', 2000)
    assert len(first['chunks']) == 2000 and [c['text'] for c in second['chunks']] == ['P2000', 'P2001']
    book = Workbook(); book.active['A1'] = 'x' * 9000 + 'CELL-TAIL'
    path = tmp_path / 'cell.xlsx'; book.save(path)
    assert 'CELL-TAIL' in str(extract_batch(path, '.xlsx')['chunks'])


def test_visible_sheet_cell_and_formula_saved_value_boundaries(tmp_path, source_files):
    from openpyxl import Workbook
    from lxml import etree
    book = Workbook(); sheet = book.active
    sheet['A1'] = date(2026, 11, 23); sheet['B1'] = '=1+2'; sheet['C1'] = '=4+5'
    sheet['A2'] = 'HIDDEN-ROW'; sheet.row_dimensions[2].hidden = True
    sheet['D1'] = 'HIDDEN-COL'; sheet.column_dimensions['D'].hidden = True
    sheet['KW22050'] = 'FAR-CELL'; sheet.merge_cells('A4:B4'); sheet['A4'] = 'MERGED'
    for i in range(21): book.create_sheet('Sheet' + str(i))['A1'] = 'LAST-SHEET' if i == 20 else 'middle'
    book.create_sheet('Hidden').sheet_state = 'hidden'; book['Hidden']['A1'] = 'HIDDEN-SHEET'
    raw = io.BytesIO(); book.save(raw)
    path = tmp_path / 'ranges.xlsx'
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as output:
        for entry in source.infolist():
            content = source.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                root = etree.fromstring(content)
                root.find("{*}sheetData/{*}row/{*}c[@r='B1']/{*}v").text = '3'
                content = etree.tostring(root)
            output.writestr(entry, content)
    value = extract_batch(path, '.xlsx'); text = str(value['chunks'])
    assert all(fact in text for fact in ('2026-11-23', 'FAR-CELL', 'LAST-SHEET', 'MERGED', '保存值（非实时计算）：3', '未保存计算结果'))
    assert all(fact not in text for fact in ('HIDDEN-ROW', 'HIDDEN-COL', 'HIDDEN-SHEET'))
    assert any('隐藏' in warning for warning in value['info']['warnings'])
    sparse = extract_batch(source_files / 'register.xlsx', '.xlsx')
    assert '22050' in str(sparse['chunks']) and '6842' in str(sparse['chunks'])


def test_word_other_regions_nested_tables_and_textboxes(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    word = Document(); word.add_paragraph('BODY')
    word.sections[0].header.paragraphs[0].text = 'HEADER'
    word.sections[0].footer.paragraphs[0].text = 'FOOTER'
    cell = word.add_table(rows=1, cols=1).cell(0, 0); cell.text = 'TABLE'
    cell.add_table(rows=1, cols=1).cell(0, 0).text = 'NESTED'
    box = OxmlElement('w:txbxContent'); para = OxmlElement('w:p'); run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'TEXTBOX'
    run.append(text); para.append(run); box.append(para); word.paragraphs[0]._p.append(box)
    path = tmp_path / 'regions.docx'; word.save(path)
    raw = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as output:
        for entry in source.infolist(): output.writestr(entry, source.read(entry))
        output.writestr('word/footnotes.xml', '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>FOOTNOTE-83</w:t></w:r></w:p></w:footnote></w:footnotes>')
    value = extract_batch(path, '.docx')
    text = '\n'.join(c['text'] for c in value['chunks'])
    assert all(text.count(fact) == 1 for fact in ('BODY', 'HEADER', 'FOOTER', 'TABLE', 'NESTED', 'TEXTBOX', 'FOOTNOTE-83'))
    assert value['info']['hasVisuals']


def test_scan_blank_pages_native_shapes_and_ppt_notes_are_discoverable(source_files, tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches
    scan = extract_batch(source_files / 'scan.pdf', '.pdf')
    assert scan['status'] == 'partial' and scan['info']['hasVisuals'] and scan['info']['totalPages'] == 1
    assert scan['chunks'][0]['text'] == ''
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(1), Inches(1), Inches(3), Inches(1))
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1), Inches(4), Inches(4))
    slide.notes_slide.notes_text_frame.text = 'RETAIN-47'
    path = tmp_path / 'native.pptx'; deck.save(path)
    value = extract_batch(path, '.pptx')
    assert value['info']['hasVisuals'] and 'RETAIN-47' in str(value['chunks'])
    empty = tmp_path / 'empty.txt'; empty.write_text('')
    assert extract_batch(empty, '.txt')['status'] == 'empty'


def test_dense_sheet_past_old_nonempty_boundary_and_rotated_blank_pdf(tmp_path):
    from openpyxl import Workbook
    from app.integrations.parsing.document_parser import MAX_NONEMPTY_CELLS
    book = Workbook(write_only=True); sheet = book.create_sheet()
    for i in range(MAX_NONEMPTY_CELLS + 1): sheet.append([f'VALUE-{i}'])
    path = tmp_path / 'dense.xlsx'; book.save(path)
    boundary = extract_batch(path, '.xlsx', MAX_NONEMPTY_CELLS - 1)
    assert [row['text'].split('：')[1] for row in boundary['chunks']] == [f'VALUE-{MAX_NONEMPTY_CELLS - 1}', f'VALUE-{MAX_NONEMPTY_CELLS}']
    assert boundary['info']['sourceComplete']
    from pypdf import PdfWriter
    writer = PdfWriter(); writer.add_blank_page(width=600, height=800).rotate(90)
    path = tmp_path / 'blank.pdf'
    with path.open('wb') as output: writer.write(output)
    blank = extract_batch(path, '.pdf')
    assert blank['status'] == 'partial' and blank['info']['totalPages'] == 1
    assert blank['info']['hasVisuals'] and blank['chunks'][0]['text'] == ''


def test_plain_word_theme_is_not_visual_and_run_separators_are_preserved(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_BREAK
    word = Document()
    paragraph = word.add_paragraph('SKU\tA12\nQuantity\t34')
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run('After page')
    carriage = OxmlElement('w:cr'); paragraph.add_run()._r.append(carriage)
    paragraph.add_run('Next line')
    expected = paragraph.text
    word.sections[0].header.paragraphs[0].text = 'Header\tvalue'
    path = tmp_path / 'plain.docx'; word.save(path)
    result = extract_batch(path, '.docx')
    assert not result['info']['hasVisuals']
    assert [item['text'] for item in result['chunks']] == [expected, 'Header\tvalue']
    assert expected == 'SKU\tA12\nQuantity\t34After page\nNext line'
    # A real drawing in the header remains visual even without embedded media.
    word.sections[0].header.paragraphs[0]._p.append(OxmlElement('w:pict'))
    word.save(path)
    assert extract_batch(path, '.docx')['info']['hasVisuals']


def test_sheet_saved_durations_dates_times_and_formula_caches_keep_their_meaning(tmp_path):
    from openpyxl import Workbook
    from lxml import etree
    book = Workbook(); sheet = book.active
    sheet['A1'] = timedelta(hours=49, minutes=30); sheet['A1'].number_format = '[h]:mm:ss'
    sheet['A2'] = date(2026, 11, 23)
    sheet['A3'] = time(13, 45, 30)
    sheet['A4'] = '=SUM(A1,A1)'; sheet['A4'].number_format = '[h]:mm:ss'
    sheet['A5'] = '=DATE(2026,11,23)'; sheet['A5'].number_format = 'yyyy-mm-dd'
    raw = io.BytesIO(); book.save(raw)
    path = tmp_path / 'typed-values.xlsx'
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(path, 'w') as output:
        for entry in source.infolist():
            data = source.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                root = etree.fromstring(data)
                for address, value in [('A4', '4.125'), ('A5', '46349')]:
                    root.find(f"{{*}}sheetData/{{*}}row/{{*}}c[@r='{address}']/{{*}}v").text = value
                data = etree.tostring(root)
            output.writestr(entry, data)
    cells = {row['source']['cell']: row['text'] for row in extract_batch(path, '.xlsx')['chunks']}
    assert cells['A1'] == 'A1：2 days, 1:30:00'
    assert cells['A2'] == 'A2：2026-11-23T00:00:00'
    assert cells['A3'] == 'A3：13:45:30'
    assert cells['A4'].endswith('文件保存值（非实时计算）：4 days, 3:00:00')
    assert cells['A5'].endswith('文件保存值（非实时计算）：2026-11-23T00:00:00')


@pytest.mark.parametrize('kind', ['simple', 'split', 'split-bounded', 'header-simple'])
def test_active_word_fields_are_rejected_before_native_rendering(tmp_path, kind):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from app.integrations.parsing.document_parser import ParseFailure
    from app.integrations.parsing.office_visual_validation import validate
    word = Document()
    paragraph = word.sections[0].header.paragraphs[0] if kind == 'header-simple' else word.add_paragraph('Material')
    if kind.endswith('simple'):
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'INCLUDEPICTURE "https://never-access.invalid/image"')
        paragraph._p.append(field)
    else:
        if kind == 'split-bounded':
            begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin'); paragraph.add_run()._r.append(begin)
        for piece in ('INCLUDE', 'PICTURE "https://never-access.invalid/image"'):
            node = OxmlElement('w:instrText'); node.text = piece; paragraph.add_run()._r.append(node)
        if kind == 'split-bounded':
            end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); paragraph.add_run()._r.append(end)
    path = tmp_path / 'active.docx'; word.save(path)
    with pytest.raises(ParseFailure, match='活动字段'):
        validate(path)
    # PAGE is a local display field; preserve safe common document formatting.
    for node in paragraph._p.iter(qn('w:instrText')): node.text = ' PAGE '
    for node in paragraph._p.iter(qn('w:fldSimple')): node.set(qn('w:instr'), ' PAGE ')
    word.save(path)
    validate(path)


@pytest.mark.parametrize('kind', ['comment', 'processing-instruction'])
def test_word_xml_annotations_preserve_text_and_active_field_checks(tmp_path, kind):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree
    from app.integrations.parsing.document_parser import ParseFailure
    from app.integrations.parsing.office_visual_validation import validate
    word = Document()
    paragraph = word.add_paragraph('SKU A12; quantity 34.')
    header = word.sections[0].header.paragraphs[0]
    header.text = 'Header value'
    for parent in (word._element, word._element.body, paragraph._p, header._p):
        node = etree.Comment('Ordinary annotation') if kind == 'comment' else etree.ProcessingInstruction('audit-note', 'local annotation')
        parent.insert(0, node)
    path = tmp_path / 'annotations.docx'
    word.save(path)
    result = extract_batch(path, '.docx')
    assert result['status'] == 'ready' and result['info']['sourceComplete']
    assert not result['info']['hasVisuals']
    assert [item['text'] for item in result['chunks']] == ['SKU A12; quantity 34.', 'Header value']
    validate(path)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'INCLUDEPICTURE "https://never-access.invalid/image"')
    paragraph._p.append(field)
    word.save(path)
    with pytest.raises(ParseFailure, match='活动字段'):
        validate(path)


@pytest.mark.parametrize('row', [1, 2, 12050])
def test_sheet_structure_counts_visible_row_gaps_without_styled_extent(tmp_path, row):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    book = Workbook(); sheet = book.active
    sheet['A1'] = 'Header'
    sheet.cell(row, 2, 'Tail'); sheet.cell(row, 3, 6842)
    sheet['Z50000'].fill = PatternFill('solid', fgColor='FF0000')
    path = tmp_path / 'sparse.xlsx'; book.save(path)
    result = extract_batch(path, '.xlsx')
    stats = result['info']['tableStructure']['sheets'][0]
    assert stats['firstContentRow'] == 1 and stats['lastContentRow'] == row
    assert stats['visibleContentCells'] == 3
    assert stats['visibleContentRows'] == (1 if row == 1 else 2)
    assert stats['visibleRowsWithoutValuesBetween'] == max(0, row - 2)
    assert not stats['mergedCellsPresent']
    assert result['info']['sourceComplete']
    assert len(result['chunks']) == 3


def test_sheet_structure_excludes_hidden_content_and_does_not_invent_merged_blank_rows(tmp_path):
    from openpyxl import Workbook
    book = Workbook(); sheet = book.active
    sheet['A1'] = 'Header'; sheet['A10'] = 'Tail'
    sheet['A3'] = 'HIDDEN'; sheet.row_dimensions[3].hidden = True
    sheet['C4'] = 'HIDDEN'; sheet.column_dimensions['C'].hidden = True
    book.create_sheet('Secret').sheet_state = 'hidden'
    path = tmp_path / 'scoped.xlsx'; book.save(path)
    result = extract_batch(path, '.xlsx')
    stats = result['info']['tableStructure']['sheets']
    assert len(stats) == 1 and stats[0]['visibleContentCells'] == 2
    assert stats[0]['visibleRowsWithoutValuesBetween'] == 7
    assert stats[0]['hiddenRowsExcluded'] and stats[0]['hiddenColumnsExcluded']
    assert 'HIDDEN' not in str(result['chunks'])
    sheet.merge_cells('A6:A8'); sheet['A6'] = 'Merged'
    book.save(path)
    merged = extract_batch(path, '.xlsx')['info']['tableStructure']['sheets'][0]
    assert merged['mergedCellsPresent'] and merged['visibleRowsWithoutValuesBetween'] is None


def test_sheet_structure_does_not_describe_unfinished_sheet_as_complete(tmp_path, monkeypatch):
    from openpyxl import Workbook
    from app.integrations.parsing import document_ranges
    book = Workbook(); book.active['A1'] = 'Small'
    dense = book.create_sheet('Dense')
    for row in range(1, 6): dense.cell(row, 1, str(row))
    path = tmp_path / 'partial.xlsx'; book.save(path)
    monkeypatch.setattr(document_ranges, 'MAX_CHUNKS', 3)
    first = extract_batch(path, '.xlsx')
    assert not first['info']['sourceComplete']
    assert [s['sheet'] for s in first['info']['tableStructure']['sheets']] == ['Sheet']
    final = extract_batch(path, '.xlsx', first['info']['nextExtractCursor'])
    assert final['info']['sourceComplete']
    assert final['info']['tableStructure']['sheets'][1]['visibleContentRows'] == 5
    assert [c['ordinal'] for c in first['chunks'] + final['chunks']] == list(range(6))
