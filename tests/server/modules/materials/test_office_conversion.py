import asyncio
import io
import os
from pathlib import Path
import signal
import struct
import sys
import time
import zipfile

import pytest
from app.integrations.parsing import office_conversion as office
from _support.materials.legacy_samples import formula_boundary_samples, legacy_samples, ods_boolean_cache, ods_boolean_saved_text, ods_repeated_tail, stale_xls, TEXT

pytestmark = pytest.mark.asyncio


@pytest.fixture(scope='module')
def real_samples():
    if not office.converter()[0]:
        pytest.skip('Real LibreOffice is not installed; conversion compatibility not verified')
    return legacy_samples()


@pytest.mark.parametrize('suffix', ['doc', 'xls', 'ppt', 'rtf', 'odt', 'ods', 'odp'])
async def test_real_conversion_content_locations_and_original_immutable(tmp_path, real_samples, suffix):
    original = real_samples['sample.' + suffix]
    path = tmp_path / 'original'; path.write_bytes(original)
    stages = []
    async def stage(value): stages.append(value)
    result = await office.parse_office(path, '.' + suffix, on_stage=stage)
    assert result['status'] in ('ready', 'partial'), result
    body = '\n'.join(row['text'] for row in result['chunks'])
    assert TEXT in body
    assert path.read_bytes() == original
    assert stages == ['queued', 'converting', 'parsing']
    assert result['info']['originalFormat'] == suffix
    assert all(row['location'] for row in result['chunks'])
    if suffix in ('xls', 'ods'):
        assert all(value in body for value in ('1234.5', '2026-10-16T', '2469', '#DIV/0!', '缓存中文', '第二工作表标记', 'A4:B4'))
        boolean = ods_boolean_saved_text(original) if suffix == 'ods' else 'True'
        assert 'E2：=1=1；文件保存值（非实时计算）：' + boolean + '\n' in body
        assert '不可见标记' not in body
        assert any('隐藏' in value for value in result['info']['warnings'])
    if suffix in ('ppt', 'odp'):
        assert '备注标记' in body
        assert any('演讲备注' in row['location'] for row in result['chunks'])


async def test_stale_formula_is_rejected_instead_of_silently_recomputed(tmp_path, real_samples):
    path = tmp_path / 'stale.xls'; path.write_bytes(stale_xls(real_samples['sample.xls']))
    result = await office.parse_office(path, '.xls')
    assert result['status'] == 'failed' and not result['chunks']
    assert '公式结果在转换时发生变化' in result['info']['error']
    # Next real conversion still works, proving the failed path released its slot.
    path.write_bytes(real_samples['sample.xls'])
    assert (await office.parse_office(path, '.xls'))['status'] == 'partial'


@pytest.mark.parametrize('name', ['now.xls', 'now.ods', 'rand.xls', 'rand.ods', 'unknown.xls', 'unknown.ods', 'missing.xls', 'missing.ods', 'stale.ods'])
async def test_formula_boundaries_never_invent_saved_values(tmp_path, real_samples, name):
    original, expected = formula_boundary_samples()[name]
    path = tmp_path / name; path.write_bytes(original)
    result = await office.parse_office(path, path.suffix)
    assert path.read_bytes() == original
    if result['status'] == 'failed':
        assert not result['chunks'] and result['info']['error']
        assert any(word in result['info']['error'] for word in ('公式', '读取失败', '活动嵌入')), result
    else:
        # A converter that really keeps a saved value/error is allowed. It may
        # not turn missing caches into zero or substitute a recalculated value.
        assert expected is not None, result
        assert result['status'] in ('ready', 'partial'), result
        body = '\n'.join(row['text'] for row in result['chunks'])
        assert str(expected) in body, (name, expected, result)
        assert all(row['location'] for row in result['chunks'])


@pytest.mark.parametrize('formula', [False, True])
async def test_odf_blank_padding_preserves_cells_but_formula_expansion_is_bounded(tmp_path, real_samples, formula):
    original = ods_repeated_tail(real_samples['sample.ods'], formula=formula)
    path = tmp_path / 'padding.ods'; path.write_bytes(original)
    result = await office.parse_office(path, '.ods')
    assert path.read_bytes() == original
    if formula:
        assert result['status'] == 'failed' and not result['chunks']
        assert '表格展开后过大' in result['info']['error'], result
    else:
        assert result['status'] == 'partial', result
        row = next(row for row in result['chunks'] if row['location'] == '工作表「金额」!A2:G2')
        assert 'B2：1234.5' in row['text'] and '文件保存值（非实时计算）：2469' in row['text']


@pytest.mark.parametrize('numeric', [False, True])
async def test_odf_boolean_formula_keeps_verified_original_cache_type(tmp_path, real_samples, numeric):
    original = ods_boolean_cache(real_samples['sample.ods'], numeric=numeric)
    path = tmp_path / 'boolean.ods'; path.write_bytes(original)
    result = await office.parse_office(path, '.ods')
    assert result['status'] == 'partial', result
    assert path.read_bytes() == original
    row = next(row for row in result['chunks'] if row['location'] == '工作表「金额」!A2:G2')
    expected = '1' if numeric else 'True'
    assert 'E2：=1=1；文件保存值（非实时计算）：' + expected + '\n' in row['text']


@pytest.mark.parametrize('saved', ['1', '0', None, 'invalid'])
async def test_odf_boolean_cache_requires_valid_value_and_keeps_xml_booleans(tmp_path, real_samples, saved):
    original = ods_boolean_cache(real_samples['sample.ods'], boolean_value=saved)
    path = tmp_path / 'boolean.ods'; path.write_bytes(original)
    result = await office.parse_office(path, '.ods')
    assert path.read_bytes() == original
    if saved in ('1', '0'):
        assert result['status'] == 'partial', result
        row = next(row for row in result['chunks'] if row['location'] == '工作表「金额」!A2:G2')
        formula, value = ('1=1', 'True') if saved == '1' else ('1=0', 'False')
        assert 'E2：=' + formula + '；文件保存值（非实时计算）：' + value + '\n' in row['text']
    else:
        assert result['status'] == 'failed' and not result['chunks']
        assert '公式保存结果无法可靠保留' in result['info']['error']


@pytest.mark.parametrize('original_value', [0, 0.5])
async def test_boolean_type_normalization_never_accepts_unequal_numeric_cache(tmp_path, monkeypatch, original_value):
    import importlib
    from openpyxl import Workbook
    from lxml import etree
    monkeypatch.syspath_prepend(str(Path(office.__file__).parent))
    validation = importlib.import_module('office_validation')
    path = tmp_path / 'converted.xlsx'
    book = Workbook(); book.active['A1'] = '=1=1'; book.save(path)
    def saved_true(name, raw):
        if name != 'xl/worksheets/sheet1.xml': return raw
        tree = etree.fromstring(raw)
        namespace = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        cell = tree.find('.//s:c', namespace); cell.set('t', 'b')
        cell.find('s:v', namespace).text = '1'
        return etree.tostring(tree)
    path.write_bytes(rewrite_zip(path.read_bytes(), saved_true))
    with pytest.raises(validation.ParseFailure, match='公式结果在转换时发生变化'):
        validation.verify_formula_caches(path, {(0, 1, 1): ('n', original_value)})


async def test_compatible_doc_disguises_use_only_recognized_routes(tmp_path, real_samples):
    path = tmp_path / 'misnamed.doc'; path.write_bytes(real_samples['sample.rtf'])
    assert (await office.parse_office(path, '.doc'))['status'] == 'ready'
    from docx import Document
    document = Document(); document.add_paragraph(TEXT); document.save(path)
    stages = []
    async def stage(value): stages.append(value)
    result = await office.parse_office(path, '.doc', on_stage=stage)
    assert result['status'] == 'ready' and stages == ['queued', 'parsing']
    for data in (b'MZbad executable', b'PK\x03\x04broken', b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1broken'):
        path.write_bytes(data)
        assert (await office.parse_office(path, '.doc'))['status'] == 'failed'


def rewrite_zip(data, change):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, 'w') as target:
        for name in source.namelist(): target.writestr(name, change(name, source.read(name)))
    return out.getvalue()


async def test_active_content_encryption_and_xml_entities_fail_closed(tmp_path, real_samples):
    samples = [('.rtf', b'{\\rtf1 {\\object\\objdata deadbeef}}')]
    samples.append(('.odt', rewrite_zip(real_samples['sample.odt'], lambda name, raw: raw.replace(b'<office:body>', b'<office:scripts><office:script/></office:scripts><office:body>') if name == 'content.xml' else raw)))
    samples.append(('.odt', rewrite_zip(real_samples['sample.odt'], lambda name, raw: raw.replace(b'<office:document-content ', b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><office:document-content ', 1) if name == 'content.xml' else raw)))
    # An actual PPT container with VBAInfoAtom's fHasMacros set, no code is run.
    import olefile
    raw = io.BytesIO(real_samples['sample.ppt'])
    with olefile.OleFileIO(raw, write_mode=True) as ole:
        stream = bytearray(ole.openstream('PowerPoint Document').read())
        marker = b'\x02\x00\x00\x04\x0c\x00\x00\x00'
        at = stream.index(marker); struct.pack_into('<I', stream, at + 12, 1)
        ole.write_stream('PowerPoint Document', bytes(stream))
    samples.append(('.ppt', raw.getvalue()))
    for index, (suffix, data) in enumerate(samples):
        path = tmp_path / str(index); path.write_bytes(data)
        result = await office.parse_office(path, suffix)
        assert result['status'] == 'failed' and not result['chunks'], result


async def test_missing_converter_is_local_failure(tmp_path, monkeypatch):
    monkeypatch.setattr(office, 'converter', lambda: ('', ''))
    result = await office.parse_office(tmp_path / 'unused.doc', '.doc')
    assert result['status'] == 'failed' and '联系管理员' in result['info']['error']


def fake_converter(tmp_path, body):
    script = tmp_path / 'soffice'
    script.write_text('#!' + sys.executable + '\n' + body)
    script.chmod(0o700)
    return script


async def test_exit_zero_without_output_and_crash_are_not_success(tmp_path, monkeypatch):
    path = tmp_path / 'input.rtf'; path.write_bytes(b'{\\rtf1 safe}')
    for code in (0, 9):
        executable = fake_converter(tmp_path, 'import sys\nsys.exit(' + str(code) + ')\n')
        monkeypatch.setattr(office, 'converter', lambda: (str(executable), 'LibreOffice controlled'))
        result = await office.parse_office(path, '.rtf')
        assert result['status'] == 'failed' and not result['chunks']


@pytest.mark.skipif(os.name != 'posix', reason='POSIX process-tree fault injection')
@pytest.mark.parametrize('cancel', [False, True])
async def test_timeout_and_cancel_reap_grandchildren_and_release_slot(tmp_path, monkeypatch, cancel):
    pidfile = tmp_path / 'pids'
    executable = fake_converter(tmp_path, 'import os, subprocess, sys, time\nfrom pathlib import Path\np=subprocess.Popen([sys.executable,"-c","import time; time.sleep(120)"])\nPath(' + repr(str(pidfile)) + ').write_text(str(os.getpid())+" "+str(p.pid)+" "+os.getcwd())\ntime.sleep(120)\n')
    monkeypatch.setattr(office, 'converter', lambda: (str(executable), 'LibreOffice controlled'))
    source = tmp_path / 'input.rtf'; source.write_bytes(b'{\\rtf1 safe}')
    task = asyncio.create_task(office.parse_office(source, '.rtf', timeout=4 if cancel else 1))
    deadline = time.monotonic() + 3
    while not pidfile.exists() and time.monotonic() < deadline: await asyncio.sleep(.02)
    assert pidfile.exists()
    pids = pidfile.read_text().split(' ', 2)
    if cancel:
        task.cancel()
        with pytest.raises(asyncio.CancelledError): await task
    else:
        assert (await task)['status'] == 'failed'
    for raw in pids[:2]:
        # A killed orphan may briefly remain a zombie until the host init reaps it.
        import subprocess
        probe = subprocess.run(['/bin/ps', '-o', 'stat=', '-p', raw], capture_output=True, text=True)
        assert not probe.stdout.strip() or probe.stdout.strip().startswith('Z')
    assert not Path(pids[2]).exists()
    assert not office._limits[asyncio.get_running_loop()].locked()


@pytest.mark.skipif(os.name != 'posix', reason='POSIX supervisor crash injection')
async def test_supervisor_sigkill_still_reaps_descendants_and_temporary_files(tmp_path, monkeypatch):
    pidfile = tmp_path / 'pids'
    executable = fake_converter(tmp_path, 'import os, subprocess, sys, time\nfrom pathlib import Path\np=subprocess.Popen([sys.executable,"-c","import time; time.sleep(120)"])\nPath(' + repr(str(pidfile)) + ').write_text(f"{os.getppid()} {os.getpid()} {p.pid} {os.getcwd()}")\ntime.sleep(120)\n')
    monkeypatch.setattr(office, 'converter', lambda: (str(executable), 'LibreOffice controlled'))
    source = tmp_path / 'input.rtf'; source.write_bytes(b'{\\rtf1 safe}')
    task = asyncio.create_task(office.parse_office(source, '.rtf', timeout=3))
    deadline = time.monotonic() + 3
    while not pidfile.exists() and time.monotonic() < deadline: await asyncio.sleep(.02)
    assert pidfile.exists()
    supervisor, child, grandchild, workspace = pidfile.read_text().split(' ', 3)
    os.kill(int(supervisor), signal.SIGKILL)
    assert (await task)['status'] == 'failed'
    import subprocess
    for pid in (child, grandchild):
        state = subprocess.run(['/bin/ps', '-o', 'stat=', '-p', pid], capture_output=True, text=True).stdout.strip()
        assert not state or state.startswith('Z')
    assert not Path(workspace).exists()
    assert not office._limits[asyncio.get_running_loop()].locked()


@pytest.mark.skipif(os.name != 'posix', reason='POSIX parent crash injection')
async def test_worker_sigkill_supervisor_cleans_its_group(tmp_path):
    import subprocess
    pidfile = tmp_path / 'pids'
    executable = fake_converter(tmp_path, 'import os, subprocess, sys, time\nfrom pathlib import Path\np=subprocess.Popen([sys.executable,"-c","import time; time.sleep(120)"])\nPath(' + repr(str(pidfile)) + ').write_text(f"{os.getppid()} {os.getpid()} {p.pid} {os.getcwd()}")\ntime.sleep(120)\n')
    source = tmp_path / 'input.rtf'; source.write_bytes(b'{\\rtf1 safe}')
    code = 'import asyncio\nfrom app.integrations.parsing import office_conversion as m\nm.converter=lambda:('+repr(str(executable))+',"LibreOffice controlled")\nasyncio.run(m.parse_office('+repr(str(source))+',".rtf",timeout=20))'
    worker = subprocess.Popen([sys.executable, '-c', code], env={**os.environ, 'PYTHONPATH': str(Path('apps/server').resolve())}, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        deadline = time.monotonic() + 3
        while not pidfile.exists() and time.monotonic() < deadline: await asyncio.sleep(.02)
        assert pidfile.exists()
        supervisor, child, grandchild, workspace = pidfile.read_text().split(' ', 3)
        worker.kill(); worker.wait()
        deadline = time.monotonic() + 4
        while Path(workspace).exists() and time.monotonic() < deadline: await asyncio.sleep(.05)
        assert not Path(workspace).exists()
        for pid in (supervisor, child, grandchild):
            state = subprocess.run(['/bin/ps', '-o', 'stat=', '-p', pid], capture_output=True, text=True).stdout.strip()
            assert not state or state.startswith('Z')
    finally:
        if worker.poll() is None: worker.kill(); worker.wait()


async def test_odf_external_relative_resources_and_formula_links_never_reach_converter(tmp_path, monkeypatch, real_samples):
    marker = tmp_path / 'started'
    executable = fake_converter(tmp_path, 'from pathlib import Path\nPath('+repr(str(marker))+').write_text("started")\n')
    monkeypatch.setattr(office, 'converter', lambda: (str(executable), 'LibreOffice controlled'))
    for href in ('../outside', '%2e%2e/outside', 'file:///tmp/private', 'https://invalid.example/x', 'missing.png', '%252e%252e/outside'):
        data = rewrite_zip(real_samples['sample.odt'], lambda name, raw: raw.replace(b'<office:body>', ('<office:body><draw:image xlink:href="'+href+'"/>').encode()) if name == 'content.xml' else raw)
        source = tmp_path / 'external.odt'; source.write_bytes(data)
        result = await office.parse_office(source, '.odt')
        assert result['status'] == 'failed' and '外部数据' in result['info']['error'], result
        assert not marker.exists()
    data = rewrite_zip(real_samples['sample.ods'], lambda name, raw: raw.replace(b'of:=SUM(', b'of:=WEBSERVICE(', 1) if name == 'content.xml' else raw)
    # Ensure the fixture actually contains the injected formula, not a no-op.
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if b'WEBSERVICE' not in z.read('content.xml'):
            data = rewrite_zip(real_samples['sample.ods'], lambda name, raw: raw.replace(b'table:formula="', b'table:formula="of:=WEBSERVICE(1)+', 1) if name == 'content.xml' else raw)
    source.write_bytes(data)
    result = await office.parse_office(source, '.ods')
    assert result['status'] == 'failed' and '外部数据' in result['info']['error']
    assert not marker.exists()


async def test_actual_libreoffice_loads_locked_down_user_profile(tmp_path, monkeypatch):
    executable, _ = office.converter()
    if not executable: pytest.skip('Real LibreOffice required')
    saved = tmp_path / 'effective.xcu'
    wrapper = fake_converter(tmp_path, 'import subprocess, sys\nfrom pathlib import Path\nfrom urllib.parse import urlsplit, unquote\nprofile=Path(unquote(urlsplit(next(a.split("=",1)[1] for a in sys.argv if a.startswith("-env:UserInstallation="))).path))\nr=subprocess.run(['+repr(executable)+',*sys.argv[1:]])\nPath('+repr(str(saved))+').write_bytes((profile/"user"/"registrymodifications.xcu").read_bytes())\nsys.exit(r.returncode)\n')
    monkeypatch.setattr(office, 'converter', lambda: (str(wrapper), 'LibreOffice controlled'))
    source = tmp_path / 'source.rtf'; source.write_bytes(b'{\\rtf1 safe text}')
    result = await office.parse_office(source, '.rtf')
    assert result['status'] == 'ready', result
    from lxml import etree
    tree = etree.fromstring(saved.read_bytes())
    ns = {'oor': 'http://openoffice.org/2001/registry'}
    for property_name, expected in [('DisableMacrosExecution','true'),('MacroSecurityLevel','3'),('ODFRecalcMode','0'),('OOXMLRecalcMode','0')]:
        values = tree.xpath('//prop[@oor:name="'+property_name+'"]/value/text()', namespaces=ns)
        assert values == [expected], (property_name, values)
    assert tree.xpath('//prop[@oor:name="Link"]/value/text()', namespaces=ns) == ['0', '0']


async def test_temp_permission_failure_and_oversized_output_are_local(tmp_path, monkeypatch):
    source = tmp_path / 'source.rtf'; source.write_bytes(b'{\\rtf1 safe}')
    executable = fake_converter(tmp_path, 'import os, sys\nfrom pathlib import Path\nout=Path(sys.argv[sys.argv.index("--outdir")+1])/"source.docx"\nwith out.open("wb") as f: f.truncate(33*1024*1024)\n')
    monkeypatch.setattr(office, 'converter', lambda: (str(executable), 'LibreOffice controlled'))
    assert (await office.parse_office(source, '.rtf'))['status'] == 'failed'
    with monkeypatch.context() as scoped:
        def denied(*args, **kwargs): raise PermissionError('controlled permission fault')
        scoped.setattr(office.tempfile, 'mkdtemp', denied)
        assert (await office.parse_office(source, '.rtf'))['status'] == 'failed'
    assert not office._limits[asyncio.get_running_loop()].locked()


async def test_startup_only_removes_labelled_dead_owned_workspaces(tmp_path, monkeypatch):
    import json
    monkeypatch.setattr(office.tempfile, 'gettempdir', lambda: str(tmp_path))
    dead = tmp_path / 'noria-office-dead'; dead.mkdir()
    (dead / 'owner.json').write_text(json.dumps({'worker': 99999999, 'supervisor': 99999998}))
    active = tmp_path / 'noria-office-active'; active.mkdir()
    (active / 'owner.json').write_text(json.dumps({'worker': os.getpid(), 'supervisor': 99999998}))
    unrelated = tmp_path / 'noria-office-unknown'; unrelated.mkdir()
    for path in (dead, active, unrelated): os.utime(path, (time.time()-120, time.time()-120))
    office.cleanup_abandoned()
    assert not dead.exists() and active.exists() and unrelated.exists()
