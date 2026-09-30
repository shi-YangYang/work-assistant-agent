from pathlib import Path
from pptx import Presentation
from support import Workspace
from noria_tools import create_slides


class Slides(Workspace):
    def test_short_blocks_share_requested_page_and_remain_editable(self):
        value = create_slides(filename='计划.pptx', title='文件元信息', pages=[{'title': '计划第一页', 'blocks': [
            {'type': 'paragraph', 'text': '这是第一段。'}, {'type': 'paragraph', 'text': '这是第二段。'},
            {'type': 'table', 'table': {'columns': ['项目', '数量'], 'rows': [['甲', 2]]}}]}])
        deck = Presentation('output/计划.pptx')
        self.assertEqual(len(deck.slides), 1)
        self.assertEqual(value['data']['pageCount'], 1)
        self.assertEqual(deck.core_properties.title, '文件元信息')
        self.assertTrue(any(shape.has_table for shape in deck.slides[0].shapes))
        text = '\n'.join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
        self.assertIn('第一段', text); self.assertIn('第二段', text)
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt
        first = next(shape.text_frame.paragraphs[0] for shape in deck.slides[0].shapes if shape.has_text_frame and '第一段' in shape.text)
        self.assertEqual(first.alignment, PP_ALIGN.LEFT)
        # OOXML stores line spacing in hundredths of a point.
        self.assertLessEqual(abs(first.line_spacing - Pt(24 * 1.45)), Pt(.02))

    def test_long_body_paginated_without_losing_end(self):
        text = '中文内容需要换行并分页。' * 120 + '终点'
        value = create_slides(filename='长文.pptx', title='长文', pages=[{'title': '具有真实宽度测量的中文长标题保持在书页安全范围之内', 'blocks': [{'type': 'paragraph', 'text': text}]}])
        deck = Presentation('output/长文.pptx')
        self.assertGreater(len(deck.slides), 1)
        all_text = ''.join(shape.text.replace('\n', '') for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn(text, all_text.replace('具有真实宽度测量的中文长标题保持在书页安全范围之内', '').translate(str.maketrans('', '', '0123456789')))
        self.assertTrue(value['warnings'])
        for slide in deck.slides:
            for shape in slide.shapes:
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)

    def test_oversized_title_or_row_rejected_not_clipped(self):
        for page in [{'title': '长' * 150}, {'title': '表格', 'blocks': [{'type': 'table', 'table': {'columns': ['a', 'b'], 'rows': [['长' * 1900, 'x']]}}]}]:
            with self.assertRaises(ValueError):
                create_slides(filename='bad.pptx', title='bad', pages=[page])
        self.assertFalse(list(Path('output').iterdir()))
