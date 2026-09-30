from pathlib import Path
from docx import Document
from pypdf import PdfReader
from PIL import Image
from support import Workspace
from noria_tools import create_document


class Documents(Workspace):
    def test_word_and_pdf_keep_text_table_images_and_chinese_font(self):
        Image.new('RGB', (80, 240), 'white').save('inputs/image.png')
        blocks = [{'type': 'heading', 'text': '第一阶段'}, {'type': 'paragraph', 'text': '中文实施计划。'},
                  {'type': 'list', 'items': ['明确目标', '验证结果']},
                  {'type': 'table', 'table': {'columns': ['任务', '数量'], 'rows': [['检查', 3]]}},
                  {'type': 'image', 'input': 'image.png', 'caption': '参考图'}]
        create_document(filename='计划.docx', title='项目计划', blocks=blocks)
        document = Document('output/计划.docx')
        self.assertIn('中文实施计划。', [p.text for p in document.paragraphs])
        self.assertEqual(document.tables[0].cell(1, 1).text, '3')
        from docx.oxml.ns import qn
        self.assertIsNone(document.styles['Title'].element.get_or_add_pPr().find(qn('w:pBdr')))
        value = create_document(filename='计划.pdf', format='pdf', title='项目计划', blocks=blocks)
        reader = PdfReader('output/计划.pdf')
        self.assertIn('中文实施计划', ''.join(page.extract_text() for page in reader.pages))
        self.assertEqual(value['data']['pageCount'], len(reader.pages))
        self.assertTrue(value['data']['fontEmbedded'])

    def test_long_pdf_paginates_and_keeps_end_marker(self):
        text = '这是需要完整保留的中文长段落。' * 600 + '终点标记'
        value = create_document(filename='长文.pdf', format='pdf', title='多页文档', blocks=[{'type': 'paragraph', 'text': text}])
        reader = PdfReader('output/长文.pdf')
        self.assertGreater(value['data']['pageCount'], 2)
        self.assertIn('终点标记', reader.pages[-1].extract_text())

    def test_unsupported_width_and_missing_image_do_not_publish(self):
        with self.assertRaises(ValueError):
            create_document(filename='bad.pdf', format='pdf', title='宽表', blocks=[{'type': 'table', 'table': {'columns': [str(i) for i in range(9)], 'rows': []}}])
        with self.assertRaises(ValueError):
            create_document(filename='bad.docx', title='缺图', blocks=[{'type': 'image', 'input': 'missing.png'}])
        self.assertFalse(list(Path('output').iterdir()))
