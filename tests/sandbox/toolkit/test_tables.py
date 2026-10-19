import csv
import json
from pathlib import Path
from openpyxl import load_workbook, Workbook
from support import Workspace
from noria_tools import inspect_table, export_table


class Tables(Workspace):
    def test_inspection_preserves_zeroes_missing_duplicate_and_sample_bounds(self):
        Path('inputs/数据.csv').write_text('账号,数量\n001,2\n001,2\n002,\n', encoding='utf-8')
        result = inspect_table(source={'input': '数据.csv', 'column_types': {'数量': 'number'}}, sample_rows=1)
        data = result['data']
        self.assertEqual(data['rowCount'], 3)
        self.assertEqual(data['sample'], [['001', 2]])
        self.assertEqual(data['duplicateRows'], 1)
        self.assertEqual(data['columns'][1]['missing'], 1)
        self.assertTrue(data['sampleTruncated'])
        self.assertFalse(list(Path('output').iterdir()))

    def test_exports_literal_formula_strings_and_empty_table(self):
        data = {'columns': ['编号', '说明'], 'rows': [['001', '=SUM(1,2)'], ['002', '中文']]}
        export_table(filename='结果.xlsx', data=data)
        book = load_workbook('output/结果.xlsx')
        self.assertEqual(book.active['A2'].value, '001')
        self.assertEqual(book.active['B2'].data_type, 's')
        self.assertEqual(book.active['B2'].value, '=SUM(1,2)')
        self.assertEqual(book.active.freeze_panes, 'A2'); book.close()
        value = export_table(filename='结果.csv', format='csv', data=data)
        with Path('output/结果.csv').open(encoding='utf-8-sig', newline='') as file:
            rows = list(csv.reader(file))
        self.assertEqual(rows[1][1], "'=SUM(1,2)")
        self.assertTrue(value['warnings'])
        empty = export_table(filename='空表.csv', format='csv', data={'columns': ['名称'], 'rows': []})
        self.assertEqual(empty['data']['rowCount'], 0)

    def test_bad_encoding_type_column_shape_are_not_silently_repaired(self):
        Path('inputs/bad.csv').write_bytes(b'name\n\xff')
        with self.assertRaises(UnicodeError):
            inspect_table(source={'input': 'bad.csv'})
        Path('inputs/bad.csv').write_text('a,b\n1\n')
        with self.assertRaises(ValueError):
            inspect_table(source={'input': 'bad.csv'})
        Path('inputs/bad.csv').write_text('a\nhello\n')
        for types in ({'a': 'number'}, {'missing': 'string'}):
            with self.assertRaises(ValueError):
                inspect_table(source={'input': 'bad.csv', 'column_types': types})

    def test_xlsx_sheet_selection_and_formula_not_calculated(self):
        book = Workbook(); book.active.title = '忽略'
        sheet = book.create_sheet('选中'); sheet.append(['编号', '公式']); sheet.append(['001', '=1+2'])
        book.save('inputs/data.xlsx')
        value = inspect_table(source={'input': 'data.xlsx', 'sheet': '选中'})
        self.assertEqual(value['data']['sample'], [['001', '=1+2']])
        self.assertTrue(value['warnings'])
        with self.assertRaises(ValueError):
            inspect_table(source={'input': 'data.xlsx', 'sheet': '不存在'})

    def test_large_samples_have_explicit_truncation_and_valid_json(self):
        with Path('inputs/wide.csv').open('w', newline='') as file:
            writer = csv.writer(file); writer.writerow([f'列{i}' for i in range(40)])
            writer.writerows([['字' * 1900] * 40 for _ in range(3)])
        value = inspect_table(source={'input': 'wide.csv'}, sample_rows=20)
        self.assertTrue(value['data']['sampleTruncated'])
        self.assertLessEqual(len(json.dumps(value, ensure_ascii=False).encode()), 65536)

    def test_explicit_numeric_conversion_and_excel_keep_large_integers(self):
        Path('inputs/large.csv').write_text('number\n9007199254740993\n')
        inspected = inspect_table(source={'input': 'large.csv', 'column_types': {'number': 'number'}})
        self.assertEqual(inspected['data']['sample'], [[9007199254740993]])
        exported = export_table(filename='large.xlsx', data={'columns': ['number'], 'rows': [[9007199254740993]]})
        book = load_workbook('output/large.xlsx')
        self.assertEqual(book.active['A2'].value, '9007199254740993')
        self.assertEqual(book.active['A2'].data_type, 's')
        self.assertTrue(exported['warnings']); book.close()
        with self.assertRaisesRegex(ValueError, 'CSV没有工作表'):
            inspect_table(source={'input': 'large.csv', 'sheet': 'ignored'})

    def test_sparse_workbook_inspection_keeps_first_row_blanks_positions_and_formula(self):
        book = Workbook(); sheet = book.active
        sheet['A1'] = 'Registry'; sheet['B12050'] = 'Tail'; sheet['C12050'] = 6842
        sheet['D2'] = '=1+2'
        book.save('inputs/sparse.xlsx')
        inspected = inspect_table(source={'input': 'sparse.xlsx'}, sample_rows=3)
        data = inspected['data']
        self.assertEqual(data['headerStatus'], 'unavailable')
        self.assertEqual(data['rowCount'], 12050)
        self.assertEqual(data['columnCount'], 4)
        self.assertEqual(data['firstContentRow'], 1)
        self.assertEqual(data['lastContentRow'], 12050)
        self.assertEqual(data['emptyRows'], 12047)
        self.assertEqual(data['sample'], [['Registry', None, None, None], [None, None, None, '=1+2']])
        self.assertEqual(data['contentSamples'], [{'rowNumber': 12050, 'values': [None, 'Tail', 6842, None]}])
        self.assertFalse(data['contentSamplesTruncated'])
        self.assertEqual(data['sampleStartRow'], 1)
        self.assertEqual(data['columns'][2], {'position': 3, 'absent': 0, 'types': {'null': 12049, 'number': 1}, 'missing': 12049})
        self.assertTrue(any('公式' in warning for warning in inspected['warnings']))
        # Inspection does not turn an arbitrary spreadsheet into a named table.
        with self.assertRaisesRegex(ValueError, '表头'):
            export_table(filename='invalid.xlsx', source={'input': 'sparse.xlsx'})
        with self.assertRaisesRegex(ValueError, 'column_types'):
            inspect_table(source={'input': 'sparse.xlsx', 'column_types': {'Registry': 'string'}})
        self.assertFalse(list(Path('output').iterdir()))

    def test_raw_csv_keeps_blank_first_record_ragged_rows_duplicates_and_zeroes(self):
        Path('inputs/raw.csv').write_text('\n001,2\n\n001,2\n003\n', encoding='utf-8')
        data = inspect_table(source={'input': 'raw.csv'})['data']
        self.assertEqual(data['headerStatus'], 'unavailable')
        self.assertEqual(data['rowCount'], 5)
        self.assertEqual(data['sample'], [[], ['001', '2'], [], ['001', '2'], ['003']])
        self.assertEqual(data['emptyRows'], 2)
        self.assertEqual(data['duplicateRows'], 2)
        self.assertEqual(data['columns'][1]['absent'], 3)
        self.assertEqual(data['columns'][1]['missing'], 3)
        for text, count in [('', 0), ('\n', 1), ('\n\n', 2)]:
            Path('inputs/empty.csv').write_text(text)
            empty = inspect_table(source={'input': 'empty.csv'})['data']
            self.assertEqual(empty['rowCount'], count)
            self.assertEqual(empty['emptyRows'], count)
            self.assertIsNone(empty['firstContentRow'])

    def test_raw_inspection_retains_limits_and_never_converts_numeric_header(self):
        book = Workbook(); sheet = book.active
        sheet.append([12, 34]); sheet.append([56, 78])
        book.save('inputs/numeric.xlsx')
        data = inspect_table(source={'input': 'numeric.xlsx'})['data']
        self.assertEqual(data['sample'], [[12, 34], [56, 78]])
        self.assertEqual(data['rowCount'], 2)
        self.assertEqual(data['columns'][0]['types'], {'number': 2})
        for text in (',' * 40 + '\n', ',\n' + '1,2\n' * 20000, ',\n' + 'x' * 2001 + ',2\n'):
            Path('inputs/invalid.csv').write_text(text)
            with self.assertRaises(ValueError):
                inspect_table(source={'input': 'invalid.csv'})

    def test_first_last_content_samples_use_source_rows_and_share_requested_budget(self):
        for header in (False, True):
            with self.subTest(header=header):
                Path('inputs/boundaries.csv').write_text(
                    ('name,value\n' if header else '') + ',\n,\nfirst,001\n,\nlast,009\n,\n')
                data = inspect_table(source={'input': 'boundaries.csv'}, sample_rows=3)['data']
                offset = int(header)
                self.assertEqual(data['sample'], [['', '']])
                self.assertEqual(data['sampleStartRow'], 1 + offset)
                self.assertEqual(data['contentSamples'], [
                    {'rowNumber': 3 + offset, 'values': ['first', '001']},
                    {'rowNumber': 5 + offset, 'values': ['last', '009']},
                ])
                self.assertEqual(data['rowCount'], 6)
                self.assertEqual(data['emptyRows'], 4)
                self.assertEqual(data['firstContentRow'], 3 + offset)
                self.assertEqual(data['lastContentRow'], 5 + offset)
                self.assertTrue(data['sampleTruncated'])
                self.assertFalse(data['contentSamplesTruncated'])
                for limit in (0, 1, 2, 3, 8, 20):
                    bounded = inspect_table(source={'input': 'boundaries.csv'}, sample_rows=limit)['data']
                    positions = [*range(1 + offset, 1 + offset + len(bounded['sample'])),
                                 *(item['rowNumber'] for item in bounded['contentSamples'])]
                    self.assertEqual(len(positions), len(set(positions)))
                    self.assertLessEqual(len(positions), limit)
                    self.assertEqual(bounded['rowCount'], 6)
                    if limit == 0:
                        self.assertEqual(bounded['sample'], [])
                        self.assertEqual(bounded['contentSamples'], [])
                        self.assertTrue(bounded['contentSamplesTruncated'])
                    elif limit == 1:
                        self.assertEqual(bounded['sample'], [['', '']])
                        self.assertEqual(bounded['contentSamples'], [])

    def test_one_content_row_is_not_repeated_as_first_and_last(self):
        Path('inputs/one.csv').write_text('\n\n\n001\n\n\n')
        data = inspect_table(source={'input': 'one.csv'}, sample_rows=3)['data']
        self.assertEqual(data['sample'], [[], []])
        self.assertEqual(data['contentSamples'], [{'rowNumber': 4, 'values': ['001']}])
        self.assertEqual(data['firstContentRow'], data['lastContentRow'])
        self.assertFalse(data['contentSamplesTruncated'])

    def test_oversized_boundary_is_omitted_without_losing_smaller_tail_or_statistics(self):
        with Path('inputs/oversized.csv').open('w', newline='') as file:
            writer = csv.writer(file)
            writer.writerow([f'column{i}' for i in range(40)])
            writer.writerows([[''] * 40, ['字' * 1900] * 40, [''] * 40, ['尾部', '001'] + [''] * 38])
        value = inspect_table(source={'input': 'oversized.csv'}, sample_rows=3)
        data = value['data']
        self.assertEqual(data['rowCount'], 4)
        self.assertEqual(data['emptyRows'], 2)
        self.assertEqual(data['firstContentRow'], 3)
        self.assertEqual(data['lastContentRow'], 5)
        self.assertEqual(data['sample'], [[''] * 40])
        self.assertEqual(data['contentSamples'], [{'rowNumber': 5, 'values': ['尾部', '001'] + [''] * 38}])
        self.assertTrue(data['sampleTruncated'])
        self.assertTrue(data['contentSamplesTruncated'])
        self.assertLessEqual(len(json.dumps(value, ensure_ascii=False).encode()), 65536)
