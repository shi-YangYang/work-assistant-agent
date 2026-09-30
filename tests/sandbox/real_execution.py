"""Opt-in destructive tests against an isolated test control service only.

Requires SPEC042_SANDBOX_TEST=1, SANDBOX_TEST_URL and SANDBOX_TEST_TOKEN.
Never run against the application/production service. Generated code stays in
resource-limited runsc containers; the outer script has bounded HTTP timeouts.
"""
import asyncio
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import time
from uuid import uuid4
import zipfile
import httpx


class Runner:
    def __init__(self):
        if os.environ.get('SPEC042_SANDBOX_TEST') != '1':
            raise RuntimeError('Explicit isolated-test opt-in required')
        self.url = os.environ['SANDBOX_TEST_URL']
        if not self.url.startswith('http://127.0.0.1:'):
            raise RuntimeError('Tests must target a local isolated control service')
        self.client = httpx.AsyncClient(base_url=self.url, headers={'Authorization': 'Bearer ' + os.environ['SANDBOX_TEST_TOKEN']}, timeout=25, trust_env=False)
        self.results = []
        self.ids = []

    async def submit(self, code=None, owner='a', inputs=None, task=None):
        identifier = hashlib.sha256(uuid4().bytes).hexdigest()
        body = {'id': identifier, 'owner': hashlib.sha256(owner.encode()).hexdigest(), 'inputs': inputs or []}
        body.update({'task': task} if task is not None else {'code': code})
        response = await self.client.post('/executions', json=body)
        assert response.status_code == 200, response.text
        self.ids.append(identifier)
        return body

    async def wait(self, identifier):
        async with asyncio.timeout(190):
            while True:
                response = await self.client.get('/executions/' + identifier)
                assert response.status_code == 200, response.text
                value = response.json()
                if value['state'] not in ('queued', 'running'):
                    return value
                await asyncio.sleep(.2)

    async def case(self, name, callback):
        started = time.monotonic()
        try:
            evidence = await callback()
            row = {'name': name, 'pass': True, 'seconds': round(time.monotonic() - started, 2), 'evidence': evidence}
        except Exception as error:
            row = {'name': name, 'pass': False, 'seconds': round(time.monotonic() - started, 2), 'error': type(error).__name__ + ': ' + str(error)[:1200]}
        self.results.append(row)
        print(json.dumps(row, ensure_ascii=False), flush=True)

    async def files(self):
        body = await self.submit('''from pathlib import Path
import json, csv
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from pptx import Presentation
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
out=Path('/work/output')
rows=[{'name':'甲','sales':120,'refund':False},{'name':'甲','sales':120,'refund':False},{'name':'乙','sales':80,'refund':True},{'name':'丙','sales':None,'refund':False}]
df=pd.DataFrame(rows).drop_duplicates(); net=df.loc[~df.refund,'sales'].sum()
assert net==120
(out/'结果.txt').write_text('净销售额：120',encoding='utf-8')
(out/'结果.md').write_text('# 净销售额\\n120',encoding='utf-8')
(out/'结果.json').write_text(json.dumps({'net':int(net)}),encoding='utf-8')
df.to_csv(out/'结果.csv',index=False)
df.to_excel(out/'结果.xlsx',index=False)
plt.bar(['net'],[net]); plt.savefig(out/'图表.png'); plt.close()
doc=Document();doc.add_heading('数据分析',0);doc.add_paragraph('净销售额：120');doc.save(out/'结果.docx')
p=Presentation()
for title,text in [('分析结果','净销售额 120'),('下一步','核实缺失数据')]:
 s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text=title;s.placeholders[1].text=text
p.save(out/'结果.pptx')
pdfmetrics.registerFont(TTFont('CJK','/usr/share/fonts/truetype/wqy/wqy-microhei.ttc',subfontIndex=0));c=canvas.Canvas(str(out/'结果.pdf'));c.setFont('CJK',18);c.drawString(60,720,'净销售额：120');c.save()
assert len(Presentation(out/'结果.pptx').slides)==2
assert Document(out/'结果.docx').paragraphs[1].text=='净销售额：120'
print(json.dumps({'net':int(net),'rows':len(df)}))''')
        result = await self.wait(body['id'])
        assert result['state'] == 'succeeded', result
        assert json.loads(result['stdout']) == {'net': 120, 'rows': 3}
        assert len(result['files']) == 9, result
        for item in result['files']:
            response = await self.client.get('/executions/' + body['id'] + '/files/' + item['id'])
            assert response.status_code == 200
            data = response.content
            generated = Path(os.environ.get('SANDBOX_GENERATED_DIR', 'artifacts/spec042/generated'))
            generated.mkdir(parents=True, exist_ok=True)
            (generated / item['name']).write_bytes(data)
            assert len(data) == item['size'] and hashlib.sha256(data).hexdigest() == item['sha256']
            if item['name'].endswith('.pptx'):
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    slides = [name for name in archive.namelist() if name.startswith('ppt/slides/slide') and name.endswith('.xml')]
                    assert len(slides) == 2
        self.file_receipt, self.file_body = result, body
        return {'formats': [item['name'] for item in result['files']], 'knownNet': 120}

    async def builtins(self):
        table = {'columns': ['月份', '销售额'], 'rows': [['一月', 10], ['二月', 20]]}
        blocks = [{'type': 'heading', 'text': '第一阶段'}, {'type': 'paragraph', 'text': '中文实施计划：保留实际内容。'}]
        inputs = [{'name': '表.csv', 'data': base64.b64encode('账号,数量\n001,2\n002,3\n'.encode()).decode()}]
        cases = [('inspect_table', {'source': {'input': '表.csv'}}, inputs),
                 ('export_table', {'filename': '导出.xlsx', 'data': table}, []),
                 ('create_chart', {'filename': '中文图表.png', 'x': '月份', 'y': ['销售额'], 'data': table, 'title': '中文销售趋势'}, []),
                 ('create_document', {'filename': '中文计划.pdf', 'format': 'pdf', 'title': '实施计划', 'blocks': blocks}, []),
                 ('create_document', {'filename': '中文计划.docx', 'title': '实施计划', 'blocks': blocks}, []),
                 ('create_slides', {'filename': '计划.pptx', 'title': '实施计划', 'pages': [{'title': '阶段计划', 'blocks': blocks}]}, [])]
        receipts = []
        for name, arguments, inputs in cases:
            body = await self.submit(task={'kind': 'builtin', 'name': name, 'version': 1, 'arguments': arguments}, inputs=inputs)
            value = await self.wait(body['id'])
            assert value['state'] == 'succeeded', value
            assert isinstance(value['data'], dict) and isinstance(value['warnings'], list)
            if name == 'inspect_table':
                assert value['data']['sample'][0] == ['001', '2'] and not value['files'], value
            else:
                assert len(value['files']) == 1, value
                item = value['files'][0]
                response = await self.client.get('/executions/' + body['id'] + '/files/' + item['id'])
                assert response.status_code == 200 and hashlib.sha256(response.content).hexdigest() == item['sha256']
                generated = Path(os.environ.get('SANDBOX_GENERATED_DIR', 'artifacts/spec042/generated'))
                generated.mkdir(parents=True, exist_ok=True)
                (generated / item['name']).write_bytes(response.content)
            receipts.append({'tool': name, 'data': value['data'], 'warnings': value['warnings']})
            replay = await self.client.post('/executions', json=body)
            assert replay.json()['data'] == value['data']
        bad = await self.submit(task={'kind': 'builtin', 'name': 'inspect_table', 'version': 1, 'arguments': {'source': {'input': '../secret.csv'}}})
        rejected = await self.wait(bad['id'])
        assert rejected['state'] == 'failed' and not rejected['files']
        helper = await self.submit("from noria_tools import export_table\nprint(export_table(filename='helper.csv',format='csv',data={'columns':['值'],'rows':[[3]]})['data']['rowCount'])")
        value = await self.wait(helper['id'])
        assert value['state'] == 'succeeded' and value['stdout'].strip() == '1', value
        return receipts

    async def same_name_parallel(self):
        async def one(owner):
            body = await self.submit("from pathlib import Path\nimport time\ntime.sleep(2)\nPath('/work/output/result.txt').write_text(" + repr(owner) + ")\nprint(" + repr(owner) + ")", owner=owner)
            value = await self.wait(body['id'])
            assert value['state'] == 'succeeded', value
            data = (await self.client.get('/executions/' + body['id'] + '/files/' + value['files'][0]['id'])).content
            assert data.decode() == owner
            return body, value
        pairs = await asyncio.gather(one('张三'), one('李四'))
        first, second = pairs
        cross = await self.client.get('/executions/' + first[0]['id'] + '/files/' + second[1]['files'][0]['id'])
        assert cross.status_code == 404
        duplicate = await self.client.post('/executions', json=first[0])
        assert duplicate.json()['stdout'] == first[1]['stdout']
        changed = await self.client.post('/executions', json={**first[0], 'code': 'print("changed")'})
        assert changed.status_code == 409
        return {'distinctOutputs': True, 'crossExecutionRead': cross.status_code, 'duplicateStable': True}

    async def boundary(self):
        body = await self.submit('''import os,socket
assert os.uname().release.endswith('gvisor')
assert os.getuid()==65532
assert not os.path.exists('/var/run/docker.sock')
assert not os.environ.get('SANDBOX_TOKEN')
assert not os.path.exists('/app/apps/server/.env.web')
for path in ['/etc/blocked','/runner/blocked']:
 try: open(path,'w').write('x')
 except OSError: pass
 else: raise AssertionError('root filesystem writable')
for host in ['1.1.1.1','169.254.169.254','127.0.0.1']:
 s=socket.socket();s.settimeout(1)
 try: s.connect((host,80))
 except OSError: pass
 else: raise AssertionError('network reachable')
 finally: s.close()
print('gvisor; nonroot; readonly; no host credentials; no network')''')
        result = await self.wait(body['id'])
        assert result['state'] == 'succeeded', result
        return result['stdout']

    async def reject_output(self):
        codes = ["from pathlib import Path\nPath('/work/output/bad.txt').symlink_to('/etc/passwd')", "from pathlib import Path\nPath('/work/output/bad.pdf').write_bytes(b'%PDF-1.7')", "from pathlib import Path\nPath('/work/output/bad.html').write_text('<script>alert(1)</script>')"]
        for code in codes:
            body = await self.submit(code)
            result = await self.wait(body['id'])
            assert result['state'] == 'failed' and not result['files'], result
        return {'symlink': 'blocked', 'corruptPDF': 'blocked', 'activeHTML': 'blocked'}

    async def resource_limits(self):
        cases = [('output', "print('x'*1000000)"), ('memory', "x=bytearray(2*1024*1024*1024)"), ('timeout', 'while True: pass')]
        results = {}
        for name, code in cases:
            body = await self.submit(code)
            result = await self.wait(body['id'])
            assert result['state'] == 'failed', result
            assert (await self.client.get('/health')).status_code == 200
            results[name] = 'failed and control remains healthy'
        body = await self.submit('''import os
try:
 with open('/work/tmp/full','wb') as f:
  for _ in range(32): f.write(b'x'*8388608)
except OSError: print('disk limit enforced')
else: raise AssertionError('disk limit missing')''')
        result = await self.wait(body['id'])
        assert result['state'] == 'succeeded' and 'disk limit enforced' in result['stdout'], result
        results['disk'] = 'limited'
        body = await self.submit('''import os,time,signal
children=[]
try:
 for _ in range(100):
  pid=os.fork()
  if pid==0: time.sleep(10);os._exit(0)
  children.append(pid)
except OSError: print('process limit enforced')
finally:
 for pid in children:
  try: os.kill(pid,signal.SIGKILL);os.waitpid(pid,0)
  except OSError: pass
assert len(children)<100''')
        result = await self.wait(body['id'])
        # gVisor's own threads share the cgroup PID budget. On some kernels the
        # sandbox is killed before Python can catch EAGAIN; both paths are bounded.
        assert (result['state'] == 'succeeded' and 'process limit enforced' in result['stdout']) or (result['state'] == 'failed' and result.get('exitCode') != 0), result
        assert (await self.client.get('/health')).status_code == 200
        after = await self.submit("print('healthy after process exhaustion')")
        assert (await self.wait(after['id']))['stdout'].strip() == 'healthy after process exhaustion'
        results['pids'] = {'limited': True, 'state': result['state'], 'exitCode': result.get('exitCode'), 'followingTask': 'succeeded'}
        return results

    async def cancel(self):
        body = await self.submit("import subprocess,time\nsubprocess.Popen(['python','-c','import time;time.sleep(60)'])\ntime.sleep(60)")
        await asyncio.sleep(1)
        response = await self.client.post('/executions/' + body['id'] + '/cancel')
        assert response.json()['state'] == 'cancelled', response.text
        after = await self.submit("print('recovered')")
        assert (await self.wait(after['id']))['stdout'].strip() == 'recovered'
        return {'cancelled': True, 'followingTask': 'succeeded'}

    async def authentication(self):
        response = await httpx.AsyncClient(timeout=5, trust_env=False).get(self.url + '/health')
        assert response.status_code == 401
        bad = await self.client.post('/executions', json={'id': 'x', 'owner': 'x', 'code': '1'})
        assert bad.status_code == 422
        return {'unauthorized': 401, 'invalidId': 422}

    async def main(self):
        try:
            for name, callback in [('authentication', self.authentication), ('nine-formats-known-calculation', self.files), ('structured-builtins-and-python-helper', self.builtins), ('parallel-isolation-and-idempotency', self.same_name_parallel), ('runtime-boundary', self.boundary), ('malicious-outputs', self.reject_output), ('cancel-child-and-recover', self.cancel), ('resource-exhaustion', self.resource_limits)]:
                await self.case(name, callback)
        finally:
            for identifier in self.ids:
                try: await self.client.delete('/executions/' + identifier)
                except httpx.HTTPError: pass
            await self.client.aclose()
        path = Path(os.environ.get('SANDBOX_TEST_REPORT', 'artifacts/spec042/sandbox-real.json'))
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(self.results, ensure_ascii=False, indent=2))
        return all(row['pass'] for row in self.results)


if __name__ == '__main__':
    raise SystemExit(0 if asyncio.run(Runner().main()) else 1)
