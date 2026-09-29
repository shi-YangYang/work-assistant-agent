"""No untrusted code is executed in these protocol boundary tests."""
import asyncio
import base64
import io
import json
import os
from pathlib import Path
import sys
import tarfile
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'apps/sandbox'))
from app.files import export_archive, validate
from app.schemas import ExecutionRequest
from app.service import Service


def archive_file(name, content=b'hello', kind=tarfile.REGTYPE):
    stream = io.BytesIO()
    with tarfile.open(fileobj=stream, mode='w') as archive:
        info = tarfile.TarInfo(name)
        info.type = kind
        info.linkname = '/etc/passwd' if kind != tarfile.REGTYPE else ''
        info.size = len(content) if kind == tarfile.REGTYPE else 0
        archive.addfile(info, io.BytesIO(content))
    return stream.getvalue()


def word_document_field(instructions, *, simple=False):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    paragraph = document.add_paragraph('Local field validation fixture')
    if simple:
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), ''.join(instructions))
        paragraph._p.append(field)
    else:
        begin = OxmlElement('w:fldChar')
        begin.set(qn('w:fldCharType'), 'begin')
        paragraph.add_run()._r.append(begin)
        for instruction in instructions:
            field = OxmlElement('w:instrText')
            field.set(qn('xml:space'), 'preserve')
            field.text = instruction
            paragraph.add_run()._r.append(field)
        separate = OxmlElement('w:fldChar')
        separate.set(qn('w:fldCharType'), 'separate')
        paragraph.add_run()._r.append(separate)
        paragraph.add_run('1')
        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')
        paragraph.add_run()._r.append(end)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class Files(unittest.TestCase):
    def test_safe_archive_and_path_link_special_types(self):
        with tempfile.TemporaryDirectory() as path:
            root = Path(path)
            good = export_archive(archive_file('output/结果.txt'), root)
            self.assertEqual((root / good[0]['id']).read_bytes(), b'hello')
            for name, kind in [('../leak.txt', tarfile.REGTYPE), ('output/nested/no.txt', tarfile.REGTYPE),
                               ('output/link.txt', tarfile.SYMTYPE), ('output/hard.txt', tarfile.LNKTYPE),
                               ('output/fifo.txt', tarfile.FIFOTYPE), ('output/code.html', tarfile.REGTYPE)]:
                with self.subTest(name=name), self.assertRaises(ValueError):
                    export_archive(archive_file(name, kind=kind), root)

    def test_corrupt_formats_and_spreadsheet_injection(self):
        for suffix, data in [('.png', b'\x89PNG\r\n\x1a\n'), ('.pdf', b'%PDF-1.7'), ('.xlsx', b'PK'), ('.json', b'{no'), ('.csv', b'a,b\n=HYPERLINK("bad"),1')]:
            with self.subTest(suffix=suffix), self.assertRaises(Exception):
                validate(data, suffix)
        validate(b'a,b\n-12,3\n', '.csv')

    def test_escaped_pdf_action_names_are_rejected(self):
        objects = [b'<< /Type /Catalog /Pages 2 0 R /O#70enAction 4 0 R >>',
                   b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
                   b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] >>',
                   b'<< /S /J#61vaScript /J#53 (app.alert\\(1\\)) >>']
        data, offsets = b'%PDF-1.7\n', [0]
        for index, value in enumerate(objects, 1):
            offsets.append(len(data)); data += str(index).encode() + b' 0 obj\n' + value + b'\nendobj\n'
        position = len(data)
        data += b'xref\n0 5\n0000000000 65535 f \n' + b''.join(f'{offset:010} 00000 n \n'.encode() for offset in offsets[1:])
        data += f'trailer\n<< /Size 5 /Root 1 0 R >>\nstartxref\n{position}\n%%EOF'.encode()
        from pypdf import PdfReader
        self.assertEqual(PdfReader(io.BytesIO(data), strict=True).trailer['/Root']['/OpenAction']['/S'], '/JavaScript')
        with self.assertRaises(ValueError):
            validate(data, '.pdf')

    def test_pdf_cjk_fonts_must_be_embedded(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject
        writer = PdfWriter(); page = writer.add_blank_page(100, 100)
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): DictionaryObject({
            NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/CIDFontType0'), NameObject('/BaseFont'): NameObject('/STSong-Light')})})})
        raw = io.BytesIO(); writer.write(raw)
        with self.assertRaisesRegex(ValueError, 'embedded'):
            validate(raw.getvalue(), '.pdf')

    def test_document_macro_and_external_resources(self):
        for malicious in ('<x><f>WEBSERVICE("http://evil")</f></x>', '<x><Relationship TargetMode="External" Target="file:///secret"/></x>', '<!DOCTYPE foo [<!ENTITY a "bad">]><x>&a;</x>'):
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, 'w') as archive:
                archive.writestr('[Content_Types].xml', '<Types/>')
                archive.writestr('xl/workbook.xml', malicious)
            with self.assertRaises(Exception):
                validate(stream.getvalue(), '.xlsx')

    def test_word_external_field_instructions_are_rejected(self):
        for instruction in ('DDE cmd " /c calc"', ' dDeAuto cmd " /c calc"',
                            'INCLUDETEXT "file:///secret"', 'INCLUDEPICTURE "https://invalid.example/image.png"',
                            'LINK Excel.Sheet.12 "file:///secret"', 'DATABASE \\d "file:///secret"'):
            for variant in ('simple', 'complex', 'split'):
                fragments = [instruction[:3], instruction[3:]] if variant == 'split' else [instruction]
                data = word_document_field(fragments, simple=variant == 'simple')
                with self.subTest(instruction=instruction, variant=variant), self.assertRaisesRegex(ValueError, 'Unsafe formula or field'):
                    validate(data, '.docx')

    def test_word_normal_field_instructions_remain_valid(self):
        for instruction in ('PAGE', 'NUMPAGES', 'DATE \\@ "yyyy-MM-dd"', 'REF bookmark', 'REF INCLUDETEXT', 'QUOTE "INCLUDETEXT"'):
            for simple in (True, False):
                with self.subTest(instruction=instruction, simple=simple):
                    validate(word_document_field([instruction[:4], instruction[4:]], simple=simple), '.docx')

    def test_fake_office_files_cannot_be_published(self):
        for suffix, entry in (('.pptx', 'ppt/presentation.xml'), ('.docx', 'word/document.xml'), ('.xlsx', 'xl/workbook.xml')):
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, 'w') as archive:
                archive.writestr('[Content_Types].xml', '<Types/>')
                archive.writestr(entry, '<x/>')
            with self.subTest(suffix=suffix), self.assertRaises(Exception):
                validate(stream.getvalue(), suffix)

    def test_request_rejects_name_collisions_and_traversal(self):
        for name in ('../a', '/root/a', 'a\\b', '..'):
            with self.assertRaises(ValueError):
                ExecutionRequest(id='a'*64, owner='b'*64, code='1', inputs=[{'name': name, 'data': ''}])
        request = ExecutionRequest(id='a'*64, owner='b'*64, code='1', inputs=[{'name': 'a', 'data': ''}, {'name': 'a', 'data': ''}])
        with self.assertRaises(ValueError):
            request.validate_inputs()


class Queue(unittest.IsolatedAsyncioTestCase):
    async def test_idempotency_limits_fair_queue_cancel_and_restart(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {'SANDBOX_STATE_DIR': directory, 'SANDBOX_CONCURRENCY': '1'}):
            service = Service()
            started = []
            async def check(): pass
            async def remove(identifier): pass
            async def run(request, target):
                started.append(request.owner)
                await asyncio.sleep(.05)
                return {'state': 'succeeded', 'files': []}
            service.runtime.check, service.runtime.verify, service.runtime.remove, service.runtime.run = check, check, remove, run
            await service.start()
            try:
                for index, owner in enumerate(['a', 'a', 'a', 'b', 'b', 'c', 'c']):
                    request = ExecutionRequest(id=f'{index:064x}', owner=owner * 64, code='print(1)')
                    service.submit(request)
                    self.assertEqual(service.submit(request)['state'], 'queued')
                with self.assertRaises(ValueError):
                    service.submit(ExecutionRequest(id='0'*64, owner='a'*64, code='print(2)'))
                await asyncio.sleep(1)
                self.assertEqual(started[:3], ['a'*64, 'b'*64, 'c'*64])
                self.assertEqual(len(started), 7)
                await service.release('0'*64)
                self.assertTrue(service.get('0'*64)['released'])
                await service.stop()
                restored = Service()
                restored.runtime.check, restored.runtime.verify, restored.runtime.remove = check, check, remove
                await restored.start()
                try:
                    self.assertEqual(restored.get('0'*64)['state'], 'succeeded')
                finally:
                    await restored.stop()
            finally:
                await service.stop()


if __name__ == '__main__':
    unittest.main()
