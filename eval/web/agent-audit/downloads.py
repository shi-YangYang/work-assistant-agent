"""Parse already-downloaded synthetic outputs without executing generated code."""
import csv
import hashlib
import json
from pathlib import Path


def inspect_file(item):
    path = Path(item['path'])
    result = {'path': str(path), 'name': item['name'], 'status': 'not_checked'}
    try:
        data = path.read_bytes()
        result.update(bytes=len(data), sha256=hashlib.sha256(data).hexdigest())
        if len(data) != item['size'] or len(data) != item['declaredSize'] or result['sha256'] != item['sha256']:
            raise ValueError('Downloaded length or digest differs from recorded delivery')
        suffix = path.suffix.lower()
        if suffix in ('.txt', '.md'):
            text = data.decode('utf-8-sig')
            result.update(characters=len(text), sample=text[:4000])
        elif suffix == '.json':
            value = json.loads(data)
            result.update(jsonType=type(value).__name__, sample=json.dumps(value, ensure_ascii=False)[:4000])
        elif suffix == '.csv':
            rows = list(csv.reader(data.decode('utf-8-sig').splitlines()))
            result.update(rows=len(rows), sample=rows[:12])
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                result['sheets'] = [{'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column, 'sample': list(sheet.iter_rows(max_row=min(12, sheet.max_row), values_only=True))} for sheet in workbook]
            finally:
                workbook.close()
        elif suffix == '.docx':
            from docx import Document
            document = Document(path)
            result.update(paragraphs=len(document.paragraphs), tables=len(document.tables), sample='\n'.join(p.text for p in document.paragraphs)[:4000])
        elif suffix == '.pptx':
            from pptx import Presentation
            presentation = Presentation(path)
            result.update(slides=len(presentation.slides), sample='\n'.join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)[:4000])
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            result.update(pages=len(reader.pages), sample='\n'.join(page.extract_text() or '' for page in reader.pages)[:4000])
        elif suffix in ('.png', '.jpg', '.jpeg'):
            from PIL import Image
            with Image.open(path) as image:
                result.update(dimensions=list(image.size), format=image.format)
                image.verify()
        else:
            result.update(status='unsupported_format', extension=suffix)
            return result
        result['status'] = 'parsed'
    except Exception as error:
        result.update(status='failed', errorType=type(error).__name__, detail=str(error)[:300])
    return result
