"""Create immutable synthetic binary inputs; expected facts never enter prompts.

This is only fixture preparation, not a model server or an answer generator.
Run once per comparison and use the exact same bytes for both code versions.
"""
import argparse
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import zipfile


def stable_archive(raw):
    """Normalize ZIP timestamps; preserve all source entries and their bytes."""
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        return raw
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(result, 'w', zipfile.ZIP_DEFLATED) as target:
        for old in sorted(source.infolist(), key=lambda entry: entry.filename):
            info = zipfile.ZipInfo(old.filename, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(old.filename))
    return result.getvalue()


def prepare(folder):
    from docx import Document
    from docx.shared import Inches
    from openpyxl import Workbook
    from pptx import Presentation
    from PIL import Image, ImageDraw, ImageFont
    folder.mkdir(parents=True, exist_ok=True)
    if (folder / 'manifest.json').exists():
        manifest = json.loads((folder / 'manifest.json').read_text())
        for name, checksum in manifest['sha256'].items():
            assert hashlib.sha256((folder / name).read_bytes()).hexdigest() == checksum
        return manifest
    created = []
    def save(name, document):
        stream = io.BytesIO()
        document.save(stream)
        raw = stable_archive(stream.getvalue())
        (folder / name).write_bytes(raw)
        created.append(name)
    for missing in (False, True):
        deck = Presentation()
        deck.core_properties.created = deck.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
        for title, body in ([('智能仓储试点项目', '范围：仓库扫码入库、库存查询和使用培训。'), ('实施安排', '需求确认后开展安装、联调和验收，其他报价见合同。')] + ([] if missing else [('预算与交付', '项目总预算：12,800元。\n最终交付日期：2026年11月6日。')])):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        save('项目方案-缺关键页.pptx' if missing else '项目方案.pptx', deck)
    contract = Document()
    contract.core_properties.created = contract.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    contract.add_heading('设备运维服务合同', 0)
    sections = {
        3: ('付款与验收', '按季度验收支付服务费。终止时已验收但未支付的服务费仍须结清；未验收项目不得直接视为已完成。'),
        12: ('专用材料与库存', '经甲方书面批准采购且不能退货、转用的专用材料，在提前终止时按凭证结算实际不可回收净成本，扣除残值和供应商退款，累计上限20,000元。未采购、可退货或可转用材料不收费。该成本不计入第18条违约金基数。'),
        18: ('提前终止', '甲方为自身便利提前终止，须提前30日书面通知。除第3条已验收未付款项及第12条合格材料成本外，按尚未履行服务价款的8%支付提前终止违约金；基数不含税费、材料费及已履行价款。同一损失不重复收取赔偿和违约金。第26、29条例外优先适用。'),
        26: ('乙方违约', '乙方发生实质违约，经书面催告15日仍未改正，甲方可解除合同，免除第18条提前终止违约金。已验收未付款仍结清；材料是否结算须按第12条核对，不能把免违约金理解为所有费用免除。'),
        29: ('不可抗力与协商解除', '不可抗力持续超过30日使主要合同目的无法实现时，任一方可书面解除，免除第18条提前终止违约金；已履行款项及合格材料依第3、12条结算。双方协商解除时以书面解除协议确定费用。')
    }
    topics = ['服务范围', '人员资质', '设备交接', '排班安排', '巡检记录', '故障响应', '备件管理', '验收流程', '数据保密', '成果归档', '质量回访', '培训安排', '联系方式', '安全操作', '信息管理']
    for number in range(1, 33):
        title, body = sections.get(number, (topics[(number - 1) % len(topics)], '双方按经确认的服务清单履行本节约定，具体设备、联系人、记录编号和交付状态以附件登记表为准。'))
        contract.add_heading(f'第{number}条 {title}', 1)
        contract.add_paragraph(body)
        for item in range(8):
            contract.add_paragraph(f'{number}.{item + 1} 执行记录：负责人员应如实登记设备编码、现场状态、处理过程、交接时间和验收结论；出现资料不齐时及时补充，与联络人核对，保留书面记录。常规巡检按批准的时间表安排，交接双方签字后归档。记录不得代替双方另行约定的费用、付款、责任或解除条款。')
    save('设备服务合同.docx', contract)
    image = Image.new('RGB', (1500, 650), '#ffffff')
    draw = ImageDraw.Draw(image)
    fonts = [Path('/System/Library/Fonts/PingFang.ttc'), Path('/System/Library/Fonts/STHeiti Medium.ttc'), Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'), Path('/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc')]
    font_path = next((path for path in fonts if path.exists()), None)
    if font_path is None:
        raise RuntimeError('A real CJK font is required to create the readable progress image')
    font = ImageFont.truetype(str(font_path), 34)
    draw.text((35, 30), '项目进度看板（2026-11-09更新）', font=font, fill='#1f2937')
    rows = [('任务', '计划完成', '当前完成/预计', '状态'), ('需求确认', '2026-10-20', '2026-10-20', '已完成'), ('设备安装', '2026-10-28', '2026-10-27', '已完成'), ('集成测试', '2026-11-06', '2026-11-09', '延期3天'), ('用户培训', '2026-11-12', '2026-11-12', '按计划')]
    for row, values in enumerate(rows):
        y = 120 + row * 95
        draw.rectangle((25, y - 8, 1470, y + 73), fill='#fff0ed' if row == 3 else '#f4f6f8')
        for x, value in zip((40, 395, 760, 1180), values):
            draw.text((x, y), value, font=font, fill='#a63128' if row == 3 else '#1f2937')
    diagram = io.BytesIO()
    image.save(diagram, format='PNG')
    (folder / '项目进度图.png').write_bytes(diagram.getvalue())
    created.append('项目进度图.png')
    document = Document()
    document.core_properties.created = document.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    document.add_heading('项目进度简报', 0)
    document.add_paragraph('下图为当前各项工作的计划和实际进展，请根据图中记录核对。')
    document.add_picture(io.BytesIO(diagram.getvalue()), width=Inches(6.3))
    save('项目进度图.docx', document)
    workbook = Workbook()
    workbook.properties.created = workbook.properties.modified = datetime(2026, 1, 1)
    sheet = workbook.active
    sheet.title = '费用明细'
    for row in [('部门', '项目', '金额'), ('研发部', '设备', 12000), ('销售部', '差旅', 4500), ('行政部', '办公', 2000), ('研发部', '工具', 6500), ('销售部', '活动', 2500), ('行政部', '维修', 1500)]:
        sheet.append(row)
    save('部门费用.xlsx', workbook)
    manifest = {'revision': 1, 'boundary': 'Synthetic input documents only, never assistant responses or mocked tool results.', 'sha256': {name: hashlib.sha256((folder / name).read_bytes()).hexdigest() for name in created}}
    (folder / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
    return manifest


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    print(json.dumps(prepare(parser.parse_args().output), ensure_ascii=False, indent=2))
