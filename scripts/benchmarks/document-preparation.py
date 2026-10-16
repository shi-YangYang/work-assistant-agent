"""Controlled API -> worker -> real parsing benchmark in an owned test database.

Run separately with --source pointing at a git-archive baseline or the checkout.
No live provider, sandbox, user database, or original media is used. Results are
append-only; browser paint and provider reasoning latency are deliberately absent.
"""
import argparse
import asyncio
from contextvars import ContextVar
import hashlib
import importlib.util
import io
import json
import logging
import math
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import time


def options():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--fixtures', type=Path, required=True)
    parser.add_argument('--samples', type=int, default=30)
    parser.add_argument('--cases', default='none,txt,py,docx,xlsx,pdf,duplicate,prepared')
    parser.add_argument('--label', required=True)
    parser.add_argument('--make-fixtures', action='store_true')
    parser.add_argument('--load-case', help='One real slow format plus three TXT tasks through three actual worker slots')
    return parser.parse_args()


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def create_fixtures(directory):
    from docx import Document
    from openpyxl import Workbook
    from _support.materials.document_samples import pdf_bytes
    directory.mkdir(parents=True, exist_ok=True)
    values = {
        'sample.txt': b'SPEC071\nAmount: 1234.50\nDue date: 2026-10-16\n',
        'sample.py': b'# SPEC071\namount = 1234.50\ndue_date = "2026-10-16"\n',
        'sample.pdf': pdf_bytes(('SPEC071 Amount 1234.50 Due date 2026-10-16',)),
    }
    document = Document()
    document.add_paragraph('SPEC071 金额 1234.50 到期 2026-10-16')
    word = io.BytesIO()
    document.save(word)
    values['sample.docx'] = word.getvalue()
    book = Workbook()
    sheet = book.active
    sheet.append(['SPEC071', '金额', '到期'])
    sheet.append(['采购', 1234.50, '2026-10-16'])
    excel = io.BytesIO()
    book.save(excel)
    values['sample.xlsx'] = excel.getvalue()
    for name, value in values.items():
        target = directory / name
        if target.exists():
            raise RuntimeError('Refusing to replace frozen fixtures: ' + str(target))
        target.write_bytes(value)


class Timings(logging.Handler):
    def __init__(self):
        super().__init__()
        self.jobs = {}

    def emit(self, record):
        message = record.getMessage()
        if message.startswith('latency '):
            value = json.loads(message[len('latency '):])
            self.jobs.setdefault(value['job'], []).append(value)


def summary(rows):
    groups = {}
    for row in rows:
        if row['sample'] < 0:
            continue
        groups.setdefault(row['case'], []).append(row)
    result = {}
    for case, values in groups.items():
        metrics = {}
        for field in ('upload_ms', 'send_ms', 'prepare_ms', 'first_model_ms', 'total_ms', 'parse_ms', 'office_ms', 'conversion_ms'):
            ordered = sorted(value[field] for value in values if value.get(field) is not None)
            if ordered:
                metrics[field] = {'p50': round(statistics.median(ordered), 3),
                                  'p95': round(ordered[math.ceil(len(ordered) * .95) - 1], 3)}
        result[case] = {'n': len(values), 'metrics': metrics,
                        'model_calls': sorted({value['model_calls'] for value in values}),
                        'parse_calls': sorted({value['parse_calls'] for value in values}),
                        'office_calls': sorted({value.get('office_calls', 0) for value in values}),
                        'states': sorted({value['state'] for value in values})}
    return result


async def resource_samples(stop, rows):
    """Sample this benchmark process and descendants; never inspect file content."""
    while not stop.is_set():
        process = await asyncio.create_subprocess_exec('/bin/ps', '-axo', 'pid=,ppid=,rss=,%cpu=', stdout=asyncio.subprocess.PIPE)
        output, _ = await process.communicate()
        entries = []
        for line in output.splitlines():
            fields = line.split()
            if len(fields) == 4:
                entries.append((int(fields[0]), int(fields[1]), int(fields[2]), float(fields[3])))
        owned = {os.getpid()}
        while True:
            descendants = owned | {pid for pid, parent, _, _ in entries if parent in owned}
            if descendants == owned:
                break
            owned = descendants
        included = [entry for entry in entries if entry[0] in owned and entry[0] != process.pid]
        rows.append({'at': time.perf_counter(), 'rss_kb': sum(row[2] for row in included),
                     'ps_cpu_percent': sum(row[3] for row in included), 'process_count': len(included)})
        try:
            await asyncio.wait_for(stop.wait(), .25)
        except TimeoutError:
            pass


async def execute(args):
    root = args.source.resolve()
    sys.path[:0] = [str(root / 'apps/server'), str(root / 'tests/server'), str(root / 'packages/voiceprint-engine/src')]
    from sqlalchemy.engine import make_url
    database_url = os.environ.get('DATABASE_TEST_URL', '')
    url = make_url(database_url)
    if url.host != '127.0.0.1' or url.port != 15471 or url.database != 'paa_company_test' or url.username != 'spec071':
        raise RuntimeError('Only the explicitly owned spec071 localhost:15471 test database is accepted')
    args.output.mkdir(parents=True, exist_ok=True)
    result_path = args.output / 'results.jsonl'
    if result_path.exists():
        raise RuntimeError('Refusing to overwrite or silently retry existing benchmark results')
    # Resolve helpers from the selected tree, so the baseline cannot accidentally
    # import modified application or test modules from the working checkout.
    spec = importlib.util.spec_from_file_location('benchmark_fixtures', root / 'tests/server/conftest.py')
    fixture_module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fixture_module)
    from _support.agent.fakes import ReviewedFixtureModel, completion
    from _support.http.management import conversation, message
    from _support.materials.documents import upload, run, running_context
    from app.tasks.processing import documents
    assert Path(documents.__file__).resolve().is_relative_to(root), 'Wrong application tree imported'
    from langchain_core.messages import AIMessage, ToolMessage
    from langchain_core.outputs import ChatGeneration, ChatResult
    from pydantic import Field
    from pwdlib import PasswordHash
    if args.make_fixtures:
        create_fixtures(args.fixtures)
    cases = args.cases.split(',')
    if args.samples < 1 or not cases:
        raise ValueError('At least one sample and case is required')
    files = {path.name: path.read_bytes() for path in args.fixtures.iterdir() if path.is_file()}
    write(args.output / 'environment.json', {
        'label': args.label, 'source': str(root), 'python': sys.version,
        'platform': platform.platform(), 'sample_count': args.samples,
        'database': {'host': url.host, 'port': url.port, 'name': url.database},
        'fixtures': {name: {'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()} for name, data in files.items()},
        'parsing_sources': {str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
                            for directory in ('apps/server/app/integrations/parsing', 'apps/server/app/tasks/processing')
                            for path in (root / directory).glob('*.py')},
        'controlled_model': True, 'real_provider': False, 'sandbox_enabled': False,
        'timing_scope': 'In-process ASGI upload/send, leased real worker process_job, real parser subprocess, PostgreSQL tools/checkpoints; no worker polling, HTTP network, SSE observation or browser paint',
        'first_round_excluded_from_quantiles': -1, 'warmups': 1,
        'startup_scope': 'Process/app startup excluded; first request reported separately, OS file cache is not flushed; each parser/converter still starts its own real child processes',
        'prepared_case': 'A previous attachment preparation is primed before timing; models still execute through the normal worker path',
    })

    class Model(ReviewedFixtureModel):
        has_documents: bool = False
        observed: list = Field(default_factory=list, exclude=True)
        model_times: list = Field(default_factory=list, exclude=True)

        async def _agenerate(self, messages, stop=None, run_manager=None, **kwargs):
            self.model_times.append(time.perf_counter())
            last = messages[-1]
            def call(name, values):
                return AIMessage(content='', tool_calls=[{'id': 'bench-' + name + '-' + str(len(self.model_times)), 'name': name, 'args': values}])
            if not self.has_documents:
                answer = completion('SPEC071：普通问题正常回复。')
            elif isinstance(last, ToolMessage) and last.name == 'find_documents':
                values = json.loads(last.content)
                if not values.get('items'):
                    raise AssertionError('No document was returned by the real tool')
                answer = call('read_document', {'attachment_id': values['items'][0]['id'], 'start': 0})
            elif isinstance(last, ToolMessage) and last.name == 'read_document':
                values = json.loads(last.content)
                self.observed.append(values)
                text = '\n'.join(row['text'] + row['citation'] for row in values['items'])
                if 'SPEC071' not in text:
                    raise AssertionError('Known fixture content was not read')
                answer = completion(('已读取的部分片段：\n' if values.get('nextCursor') is not None else '读取结果：\n') + text)
            else:
                answer = call('find_documents', {'query': ''})
            return ChatResult(generations=[ChatGeneration(message=answer)])

    timing = Timings()
    logger = logging.getLogger('paa.latency')
    logger.handlers[:] = [timing]
    logger.propagate = False
    active = ContextVar('spec071_active_sample', default=None)
    actual_parse = documents.parse_process
    async def measured_parse(*values, **kwargs):
        started = time.perf_counter()
        row = active.get()
        if row is not None:
            row['parse_calls'] += 1
        try:
            return await actual_parse(*values, **kwargs)
        finally:
            if row is not None:
                row['parse_ms'] += (time.perf_counter() - started) * 1000
    documents.parse_process = measured_parse
    actual_office = getattr(documents, 'parse_office', None)
    if actual_office:
        async def measured_office(*values, **kwargs):
            row = active.get()
            started = time.perf_counter()
            stages = []
            original_stage = kwargs.pop('on_stage', None)
            async def stage(name):
                stages.append((name, time.perf_counter()))
                if original_stage:
                    await original_stage(name)
            if row is not None:
                row['office_calls'] = row.get('office_calls', 0) + 1
            try:
                return await actual_office(*values, **kwargs, on_stage=stage)
            finally:
                if row is not None:
                    row['office_ms'] = row.get('office_ms', 0) + (time.perf_counter() - started) * 1000
                    clock = dict(stages)
                    if 'converting' in clock and 'parsing' in clock:
                        row['conversion_ms'] = row.get('conversion_ms', 0) + (clock['parsing'] - clock['converting']) * 1000
                    row.setdefault('office_stages', []).extend({'name': name, 'elapsed_ms': (at - started) * 1000} for name, at in stages)
        documents.parse_office = measured_office
    rows = []
    with tempfile.TemporaryDirectory(prefix='spec071-benchmark-') as temporary:
        fixtures = fixture_module.setup.__wrapped__(Path(temporary), PasswordHash.recommended().hash('controlled-test-password'))
        settings, sessions, users, clients = await anext(fixtures)
        try:
            client = clients['employee']
            if args.load_case:
                if not url.query.get('options', '').startswith('-csearch_path=spec071_load_'):
                    raise RuntimeError('Concurrent worker claims require a separate owned load schema')
                from dataclasses import replace
                from app.tasks.runtime.runner import run_slots
                from app.tasks.processing.handlers import process_job
                models, receipts, observed = {}, {}, {}
                stop_worker, done = asyncio.Event(), asyncio.Event()
                async def measured_runner(job, worker_sessions, worker_settings, saver):
                    row = observed[job.id]
                    token = active.set(row)
                    row['queue_ms'] = (time.perf_counter() - row['sent_clock']) * 1000
                    try:
                        await process_job(job, worker_sessions, worker_settings, saver, model=models[job.id])
                        row['completed_clock'] = time.perf_counter()
                        row['total_ms'] = (row['completed_clock'] - row['sent_clock']) * 1000
                    finally:
                        active.reset(token)
                        if all('completed_clock' in value for value in observed.values()):
                            done.set()
                for index in range(args.samples):
                    models.clear(); receipts.clear(); observed.clear(); done.clear()
                    for role, kind in zip(('admin', 'employee', 'peer', 'outsider'), (args.load_case, 'txt', 'txt', 'txt')):
                        actor_client = clients[role]
                        conv = await conversation(actor_client, 'SPEC071 concurrent files')
                        content = files['sample.' + kind]
                        if kind == 'txt':
                            content += ('\nSynthetic actor: ' + role).encode()
                        item = await upload(actor_client, 'sample.' + kind, content)
                        sent_clock = time.perf_counter()
                        receipt = await message(actor_client, conv, '读取附件并列出实际内容。', [item['id']])
                        jid = receipt['jobId']
                        receipts[jid] = (actor_client, receipt, item['id'])
                        models[jid] = Model(model='controlled-spec071', api_key='test-no-network', base_url='http://127.0.0.1:1', max_retries=0, has_documents=True)
                        observed[jid] = {'sample': index, 'case': kind, 'actor': role, 'job_id': jid, 'sent_clock': sent_clock, 'parse_calls': 0, 'parse_ms': 0.0}
                    stop_worker.clear()
                    stop_monitor, resources = asyncio.Event(), []
                    monitor = asyncio.create_task(resource_samples(stop_monitor, resources))
                    workers = asyncio.create_task(run_slots(sessions, replace(settings, worker_concurrency=3), stop_worker, runner=measured_runner))
                    try:
                        await asyncio.wait_for(done.wait(), timeout=90)
                    finally:
                        stop_worker.set()
                        await workers
                        stop_monitor.set()
                        await monitor
                    write(args.output / f'load-resources-{index}.json', {'samples': resources, 'peak_rss_kb': max((r['rss_kb'] for r in resources), default=None), 'peak_ps_cpu_percent': max((r['ps_cpu_percent'] for r in resources), default=None), 'peak_process_count': max((r['process_count'] for r in resources), default=None), 'scope': '250ms sampled benchmark worker and child processes; excludes database/container/browser; ps CPU is the operating-system reported average'})
                    for jid, row in observed.items():
                        actor_client, receipt, attachment_id = receipts[jid]
                        data = (await actor_client.get('/api/v1/messages/' + receipt['messageId'])).json()
                        assert data['job']['state'] in ('succeeded', 'awaiting_input') and not data['job']['error'], data
                        assert 'SPEC071' in data['reply'] and data.get('citations'), data
                        assert all(citation['attachmentId'] == attachment_id for citation in data['citations']), data
                        if row['case'] == 'txt':
                            assert 'Synthetic actor: ' + row['actor'] in data['reply'], data
                        events = timing.jobs.pop(jid, [])
                        row.update(events=events, state=data['job']['state'], model_calls=len(models[jid].model_times),
                                   first_model_ms=(models[jid].model_times[0] - row['sent_clock']) * 1000)
                        with result_path.open('a') as output:
                            output.write(json.dumps(row, ensure_ascii=False) + '\n')
                        rows.append(row)
                    slow = next(value for value in observed.values() if value['case'] == args.load_case)
                    fast = [value for value in observed.values() if value['case'] == 'txt']
                    assert all(value['completed_clock'] < slow['completed_clock'] for value in fast), observed
                    print(json.dumps({'load_sample': index, 'slow_ms': slow['total_ms'], 'fast_ms': [value['total_ms'] for value in fast]}), flush=True)
                write(args.output / 'summary.json', summary(rows))
                return
            for index in range(-1, args.samples):
                for case in cases:
                    row = {'sample': index, 'case': case, 'label': args.label, 'parse_calls': 0, 'parse_ms': 0.0}
                    active.set(row)
                    conv = await conversation(client, 'SPEC071 performance')
                    file_kind = 'docx' if case in ('duplicate', 'prepared') else case.split(':', 1)[-1]
                    selected = [] if case == 'none' else [('sample.' + file_kind, files['sample.' + file_kind])]
                    if case == 'duplicate' or case.startswith('duplicate:'):
                        selected *= 2
                    started = time.perf_counter()
                    attachments = [await upload(client, name, data) for name, data in selected]
                    row['upload_ms'] = (time.perf_counter() - started) * 1000
                    assert all(item['extraction']['status'] == 'unsent' for item in attachments)
                    assert row['parse_calls'] == 0, 'Uploading started parsing before send'
                    sent_at = time.perf_counter()
                    sent = await message(client, conv, '读取附件并给出材料中金额与日期。' if attachments else '正常回复一句即可。', [item['id'] for item in attachments])
                    row['send_ms'] = (time.perf_counter() - sent_at) * 1000
                    if case == 'prepared' or case.startswith('prepared:'):
                        context = await running_context(settings, sessions, users['employee'], sent)
                        for item in attachments:
                            await documents.prepare_document(context, item['id'])
                        row['priming_ms'] = (time.perf_counter() - sent_at) * 1000
                        row['parse_calls'], row['parse_ms'] = 0, 0.0
                        row['office_calls'], row['office_ms'], row['conversion_ms'], row['office_stages'] = 0, 0.0, 0.0, []
                        sent_at = time.perf_counter()
                    model = Model(model='controlled-spec071', api_key='test-no-network', base_url='http://127.0.0.1:1', max_retries=0, has_documents=bool(attachments))
                    await run(settings, sessions, users, sent, model=model)
                    data = (await client.get('/api/v1/messages/' + sent['messageId'])).json()
                    row['total_ms'] = (time.perf_counter() - sent_at) * 1000
                    row['model_calls'] = len(model.model_times)
                    row['first_model_ms'] = (model.model_times[0] - sent_at) * 1000 if model.model_times else None
                    row['state'] = data['job']['state']
                    row['job_id'] = sent['jobId']
                    events = timing.jobs.pop(sent['jobId'], [])
                    row['prepare_ms'] = next((item['at_ms'] for item in events if item['event'] == 'prepare.end'), None)
                    row['events'] = events
                    row['extraction'] = [item['extraction'] for item in data['attachments']]
                    row['citation_count'] = len(data.get('citations', []))
                    assert row['state'] in ('succeeded', 'awaiting_input') and not data['job']['error'], data['job']
                    assert data['job'].get('taskOutcome', {}).get('state') == 'completed', data['job']
                    assert 'SPEC071' in data['reply'], data['reply']
                    assert not attachments or model.observed and data.get('citations'), data
                    assert all(item['extraction']['status'] in ('ready', 'partial') for item in data['attachments']), data
                    # Read back again to cover durable final content, not merely a
                    # callback or a successful status with an empty answer.
                    refreshed = (await client.get('/api/v1/messages/' + sent['messageId'])).json()
                    assert refreshed['reply'] == data['reply'] and refreshed['citations'] == data['citations']
                    with result_path.open('a') as output:
                        output.write(json.dumps(row, ensure_ascii=False) + '\n')
                    rows.append(row)
                print(json.dumps({'label': args.label, 'sample': index, 'completed': len(rows)}), flush=True)
            write(args.output / 'summary.json', summary(rows))
        finally:
            active.set(None)
            documents.parse_process = actual_parse
            if actual_office:
                documents.parse_office = actual_office
            await fixtures.aclose()
    print(json.dumps(summary(rows), ensure_ascii=False), flush=True)


if __name__ == '__main__':
    asyncio.run(execute(options()))
