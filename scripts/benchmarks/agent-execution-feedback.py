"""Opt-in Spec075 real-model API/worker evaluation; never part of default CI.

Separate immutable source trees and an owned local database keep baseline and
modified runs independent. No model outputs/judges are mocked. Historical
production evidence must already be anonymized and stays outside Git. Failure
branches not reached by a live model are reported as unexercised, never forced.
"""
import argparse
import asyncio
import copy
from contextvars import ContextVar
import hashlib
import importlib.util
import io
import json
import logging
import os
from pathlib import Path
import sys
import tempfile
import time
from types import SimpleNamespace
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[2]
DATABASE = 'postgresql://spec075:spec075-local-only@127.0.0.1:55437/paa_company_test'
CASE_FILE = Path(__file__).with_name('agent-execution-feedback-cases.json')


def arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--app-root', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--label', choices=('before', 'after'), required=True)
    parser.add_argument('--source-user', default='admin')
    parser.add_argument('--evidence', type=Path)
    parser.add_argument('--ids', default='E14-Q')
    parser.add_argument('--efforts', default='low,high')
    parser.add_argument('--samples', default='0')
    parser.add_argument('--resume', action='store_true')
    parser.add_argument('--prepare-only', action='store_true')
    parser.add_argument('--serve', action='store_true', help='Serve isolated accounts for --ids/--efforts; record every submitted UI job')
    parser.add_argument('--port', type=int, default=8075)
    parser.add_argument('--concurrent-live-processes', type=int, default=1)
    parser.add_argument('--capture-request-policy', action='store_true', help='Record actual system instructions and offered tool schemas, never model reasoning')
    return parser.parse_args()


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + '\n')


def append(path, value):
    with path.open('a') as stream:
        stream.write(json.dumps(value, ensure_ascii=False, default=str) + '\n')


def failed_tool_arguments(result, tool_index, metadata, fragments):
    """Only the failed tool payload; never copy response text or reasoning."""
    try:
        function = result['choices'][0]['message']['tool_calls'][tool_index]['function']
        name, arguments = function['name'], function['arguments']
    except (KeyError, TypeError, IndexError):
        return None
    if not isinstance(name, str) or not isinstance(arguments, str):
        return None
    selected = [item for item in fragments if item['index'] == tool_index] if fragments is not None else None
    fields = ('json_scope', 'json_line', 'json_column', 'json_offset', 'json_input_chars',
              'http_status', 'content_type', 'provider_request_id', 'stream_complete')
    return {'toolIndex': tool_index, 'functionName': name, 'arguments': arguments,
            'diagnostic': {key: metadata[key] for key in fields if key in metadata},
            'toolArgumentFragments': selected,
            'fragmentAssemblyMatches': ''.join(item['arguments'] for item in selected) == arguments if selected is not None else None,
            'fragmentNameMatches': ''.join(item['name'] for item in selected) == name if selected is not None else None,
            'boundary': 'Only existing on_delta tool fragments; no raw SSE bytes, response body or hidden reasoning.'}


def replace_loaded_aliases(original, replacement):
    """Observe direct imports too, without changing arguments or return values."""
    patched = []
    for name, module in tuple(sys.modules.items()):
        if not name.startswith('app.') or module is None:
            continue
        for attribute, value in tuple(vars(module).items()):
            if value is original:
                setattr(module, attribute, replacement)
                patched.append((module, attribute, original, replacement))
    return patched


def restore_aliases(patched):
    for module, attribute, original, replacement in reversed(patched):
        if getattr(module, attribute, None) is replacement:
            setattr(module, attribute, original)
        # A module first imported after installation may have copied the wrapper.
        replace_loaded_aliases(replacement, original)


def observe_questions(module, active_row):
    original_ask, original_finish = module.ask, module.finish_waiting

    async def ask(*values, **kwargs):
        caller = sys._getframe(1)
        callsite = caller.f_globals.get('__name__', '') + '.' + caller.f_code.co_name
        del caller
        result = await original_ask(*values, **kwargs)
        row = active_row()
        if row is not None and result.get('state') == 'waiting':
            at = (time.perf_counter() - row['_clock']) * 1000
            origin = ('program:' + callsite.removeprefix('app.agent.actions.') if callsite.startswith('app.agent.actions.material_questions.')
                      else 'program:clarification.ask_missing' if callsite == 'app.agent.actions.clarification.ask_missing'
                      else 'model_tool:request_user_input' if callsite == 'app.agent.tools.control.questions.request_user_input'
                      else 'unclassified')
            row.setdefault('questionAvailableMs', at)
            row.setdefault('questionEvents', []).append({'at_ms': at, 'interactionId': result.get('interactionId'),
                'origin': origin, 'callsite': callsite, 'boundary': 'ask returned after transaction commit; excludes browser paint'})
        return result

    async def finish(*values, **kwargs):
        result = await original_finish(*values, **kwargs)
        row = active_row()
        if row is not None and result:
            row.setdefault('deliveryEvents', []).append({'at_ms': (time.perf_counter()-row['_clock'])*1000,
                'origin': 'program:interactions.finish_waiting'})
        return result

    return replace_loaded_aliases(original_ask, ask) + replace_loaded_aliases(original_finish, finish)


def captured_request_policy(messages, tools, tool_choice):
    """Copy system inputs/schemas, omitting user/assistant/tool messages."""
    return {'systemMessages': [copy.deepcopy({'role': item['role'], 'content': item.get('content')})
                               for item in messages if item.get('role') in ('system', 'developer')],
            'tools': copy.deepcopy(tools or []), 'toolChoice': tool_choice,
            'boundary': 'Actual chat-call inputs before provider transport; excludes user/assistant/tool messages and response reasoning.'}


def complete_public_observation(row):
    """Join already-collected API data; never add a query to the timed path."""
    panels = {item['id']: item for item in row.get('interactions', [])}
    for event in row.get('questionEvents', []):
        panel = panels.get(event['interactionId'])
        if panel:
            event['questions'] = panel.get('questions', [])
            event['persistedState'] = panel.get('state')
            event['contentSource'] = 'final interactions API snapshot; timestamp captured separately at ask return'
    if any(event['origin'] == 'program:interactions.finish_waiting' for event in row.get('deliveryEvents', [])):
        phase = row.get('message', {}).get('job', {}).get('phase')
        row['answer_origin'] = {'awaiting_answer': 'program:interaction_question',
                               'awaiting_confirmation': 'program:confirmation_receipt'}.get(phase, 'program:finish_waiting_unclassified')
    elif row.get('jobResult', {}).get('responseMode') == 'receipt':
        row['answer_origin'] = 'program:business_receipt'
    else:
        row['answer_origin'] = 'unclassified'
        if row.get('answer'):
            for request in row.get('requests', []):
                if any(item['text'] == row['answer'] for item in request.get('publicContent', [])):
                    row['answer_origin'] = 'model:public_content_exact_match'
                for tool in request.get('toolCalls', []):
                    if tool.get('name') != 'finish_task':
                        continue
                    try:
                        answer = json.loads(tool.get('arguments', '{}')).get('answer')
                    except (TypeError, ValueError):
                        continue
                    if answer == row['answer']:
                        row['answer_origin'] = 'model:finish_task_answer_exact_match'
    candidates = [(row.get('first_answer_ms'), row.get('first_answer_source'), row['answer_origin'])]
    candidates += [(event['at_ms'], 'persisted interaction', event['origin']) for event in row.get('questionEvents', [])]
    available = [item for item in candidates if item[0] is not None]
    if available:
        at, boundary, origin = min(available, key=lambda item: item[0])
        row.update(first_effective_ms=at, first_effective_source=boundary, first_effective_origin=origin)
    else:
        row['first_effective_ms'] = None


def checked(response, status=200):
    if response.status_code != status:
        raise RuntimeError(f'API {response.request.url.path}: expected {status}, got {response.status_code}; {response.text[:300]}')
    return response.json()


def keyed():
    return {'Idempotency-Key': str(uuid4())}


def source_digest(root):
    digest = hashlib.sha256()
    for path in sorted((root / 'apps/server/app').rglob('*.py')):
        digest.update(str(path.relative_to(root)).encode())
        digest.update(path.read_bytes())
    return digest.hexdigest()


class Timings(logging.Handler):
    def __init__(self):
        super().__init__()
        self.events = []

    def emit(self, record):
        value = record.getMessage()
        if value.startswith('latency '):
            self.events.append(json.loads(value[8:]))


async def configuration(args):
    from dataclasses import replace
    from dotenv import dotenv_values
    from sqlalchemy import select
    from sqlalchemy.engine import make_url
    from app.core.config import Settings
    from app.db.session import database
    from app.modules.members.models import Member
    from app.modules.model_services.models import ModelServiceRevision
    from app.modules.model_services.bindings import bind_job
    from app.security.secrets import decrypt
    actual = dotenv_values(ROOT / 'apps/server/.env.web')
    settings = replace(Settings(), database_url=actual.get('DATABASE_URL', Settings().database_url).replace('postgresql://', 'postgresql+psycopg://', 1), model_key_file=(ROOT / actual.get('PAA_MODEL_KEY_FILE', 'data/company/model-master.key')).resolve())
    assert make_url(settings.database_url).host in ('127.0.0.1', 'localhost'), 'Source configuration must be local'
    engine, sessions = database(settings)
    try:
        async with sessions() as db:
            users = list((await db.scalars(select(Member).where(Member.username == args.source_user, Member.active.is_(True)))).all())
            if len(users) != 1:
                raise RuntimeError('Source user must identify exactly one active local account')
            job = SimpleNamespace(kind='message', result={}, model_binding=None, config_attempt=0, company_id=users[0].company_id, owner_id=users[0].id)
            selected = (await bind_job(db, job, settings))['assistant']
            revision = await db.get(ModelServiceRevision, selected['revisionId'])
            model = copy.deepcopy(next(item for item in revision.models if item['id'] == selected['modelId']))
            return {'model': model, 'base': revision.base_url, 'choice': selected, 'key': decrypt(settings.model_key_file, revision.credential, users[0].company_id, selected['serviceId'], revision.revision)}
    finally:
        await engine.dispose()


async def install(config, settings, sessions, company):
    from app.modules.model_services.models import ModelService, ModelServiceRevision, ModelRouting
    from app.security.secrets import encrypt
    async with sessions.begin() as db:
        service = ModelService(company_id=company, name='Spec075 isolated evaluation', base_url=config['base'], models=[config['model']])
        db.add(service)
        await db.flush()
        db.add(ModelServiceRevision(company_id=company, service_id=service.id, revision=1, name=service.name, base_url=service.base_url, models=service.models, credential=encrypt(settings.model_key_file, config['key'], company, service.id, 1)))
        db.add(ModelRouting(company_id=company, choices={'assistant': {'serviceId': service.id, 'modelId': config['model']['id'], 'streaming': config['choice'].get('streaming', True)}, 'asr': None, 'image': None}))


def sample_file(name, text):
    stream = io.BytesIO()
    suffix = Path(name).suffix.lower()
    if suffix == '.docx':
        from docx import Document
        document = Document()
        document.add_heading(Path(name).stem, 0)
        document.add_paragraph(text)
        document.save(stream)
    elif suffix == '.xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(['事项', '说明'])
        workbook.active.append([Path(name).stem, text])
        workbook.save(stream)
    elif suffix == '.pptx':
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = Path(name).stem
        slide.placeholders[1].text = text
        deck.save(stream)
    elif suffix in ('.png', '.jpg', '.jpeg'):
        from PIL import Image
        Image.new('RGB', (640, 360), color=(242, 244, 248)).save(stream, format='PNG' if suffix == '.png' else 'JPEG')
    elif suffix == '.pdf':
        from pypdf import PdfWriter
        document = PdfWriter()
        document.add_blank_page(width=595, height=842)
        document.write(stream)
    else:
        stream.write(text.encode())
    return stream.getvalue()


async def snapshot(sessions, actor):
    from sqlalchemy import select
    from app.modules.work.models import WorkItem, WorkMaterial
    from app.modules.reports.models import Report
    from app.modules.operations.models import BusinessAction
    async with sessions() as db:
        works = (await db.scalars(select(WorkItem).where(WorkItem.owner_id == actor.id).order_by(WorkItem.created_at))).all()
        materials = (await db.scalars(select(WorkMaterial).where(WorkMaterial.owner_id == actor.id).order_by(WorkMaterial.created_at))).all()
        reports = (await db.scalars(select(Report).where(Report.owner_id == actor.id).order_by(Report.created_at))).all()
        actions = (await db.scalars(select(BusinessAction).where(BusinessAction.owner_id == actor.id).order_by(BusinessAction.created_at))).all()
    return {'works': [{'id': item.id, 'content': item.content, 'revision': item.revision, 'deleted': item.deleted} for item in works], 'materials': [{'id': item.id, 'workId': item.work_id, 'source': item.source, 'removed': item.removed} for item in materials], 'reports': [{'id': item.id, 'content': item.content, 'revision': item.revision, 'publishedRevision': item.published_revision} for item in reports], 'actions': [{'id': item.id, 'action': item.action, 'state': item.state, 'result': item.result} for item in actions]}


async def seed(case, settings, sessions, actor, client, clients, evidence):
    from sqlalchemy import select
    from app.modules.members.models import Member
    from app.modules.conversations.models import Conversation
    async with sessions.begin() as db:
        member = await db.get(Member, actor.id)
        member.execution_mode = 'full'
        member.full_access_confirmed = True
    conv = checked(await client.post('/api/v1/conversations', json={'title': 'Spec075 ' + case['id'], 'personaId': 'dabao'}), 201)
    seed_kind = case['seed']
    state = {'conversationId': conv['id'], 'attachmentIds': [], 'differences': []}
    if seed_kind.startswith('production-'):
        if evidence is None:
            raise RuntimeError('Anonymized production evidence is required')
        await seed_production(case, evidence, state, settings, sessions, actor)
        return state
    if seed_kind == 'private':
        work = checked(await clients['peer'].post('/api/v1/work-items', json={'title': '同事私密资料整理', 'summary': '仅限归属账号'}, headers=keyed()), 201)
        state['privateWorkId'] = work['id']
    if seed_kind in ('one-work', 'two', 'four', 'eleven', 'same-title'):
        work = checked(await client.post('/api/v1/work-items', json={'title': '资料整理', 'body': {'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': '准备合同台账与内部资料，保留原始说明。'}]}]}, 'status': 'in_progress'}, headers=keyed()), 201)
        state['workId'] = work['id']
        state['workRevision'] = work['revision']
        if seed_kind == 'same-title':
            checked(await client.post('/api/v1/work-items', json={'title': '资料整理', 'summary': '另一项工作，请勿修改。'}, headers=keyed()), 201)
    if seed_kind in ('two', 'four', 'eleven', 'same-title'):
        await seed_artifacts(11 if seed_kind == 'eleven' else 4 if seed_kind == 'four' else 2, state, settings, sessions, actor)
    if seed_kind in ('text-file', 'bad-file'):
        content = sample_file('损坏前的项目说明.docx', '项目说明') if seed_kind == 'bad-file' else '项目预算：12800元。交付日期：2026-11-06。'.encode()
        name = '项目说明.docx' if seed_kind == 'bad-file' else '项目说明.txt'
        uploaded = checked(await client.post('/api/v1/uploads', files={'file': (name, content, 'application/octet-stream')}), 201)
        state['attachmentIds'].append(uploaded['id'])
        if seed_kind == 'bad-file':
            (settings.media_dir / uploaded['id']).write_bytes(b'controlled post-upload storage corruption')
            state['differences'].append('Controlled disk corruption after a valid upload; live model response is not mocked.')
    return state


async def seed_artifacts(count, state, settings, sessions, actor):
    from datetime import datetime, timedelta, timezone
    from app.modules.messages.models import Message
    from app.modules.deliverables.models import Deliverable, DeliverableRevision
    from app.modules.executions.files import file_path
    titles = ['资料清单', '访谈答复', '内控制度', '内控制度', '采购制度', '定价办法', '运营说明', '研发进展', '合同台账', '产品介绍', '产品介绍'][:count]
    state['artifacts'] = []
    async with sessions.begin() as db:
        for index, title in enumerate(titles):
            at = datetime.now(timezone.utc) - timedelta(minutes=count-index)
            message = Message(company_id=actor.company_id, owner_id=actor.id, conversation_id=state['conversationId'], text='请生成' + title, reply='', created_at=at, persona_id='dabao')
            db.add(message)
            await db.flush()
            item = Deliverable(company_id=actor.company_id, owner_id=actor.id, conversation_id=state['conversationId'], title=title, revision=1, created_at=at, updated_at=at)
            db.add(item)
            await db.flush()
            suffix = '.xlsx' if index == 8 else '.pptx' if index >= 9 else '.docx'
            name = title + suffix
            raw = sample_file(name, '隔离测试资料，第' + str(index + 1) + '份。预算12800元，交付日期2026-11-06。')
            file_id = str(uuid4())
            storage = hashlib.sha256((actor.id + file_id + name).encode()).hexdigest()
            path = file_path(settings, storage)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(raw)
            mime = {'.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation'}[suffix]
            file = {'id': file_id, 'name': name, 'mimeType': mime, 'size': len(raw), 'sha256': hashlib.sha256(raw).hexdigest(), 'storageKey': storage, 'sources': []}
            record = DeliverableRevision(company_id=actor.company_id, owner_id=actor.id, deliverable_id=item.id, revision=1, message_id=message.id, step=1, title=title, body='用于隔离评测的' + title, items=[], files=[file], digest=hashlib.sha256(raw).hexdigest(), created_at=at)
            db.add(record)
            state['artifacts'].append({'deliverableId': item.id, 'revision': 1, 'fileId': file_id, 'name': name})


async def seed_production(case, evidence, state, settings, sessions, actor):
    from datetime import datetime
    import re
    from app.core.digests import digest
    from app.modules.messages.models import Message
    from app.modules.attachments.models import Attachment, DocumentChunk
    from app.modules.deliverables.models import Deliverable, DeliverableRevision
    from app.modules.executions.files import file_path
    from app.modules.work.models import WorkItem, WorkRevision, WorkMaterial
    from app.modules.conversations.models import ConversationTaskState
    from app.tasks.models import Job
    # Every run gets independent database identities, while relative references,
    # ordering and the eleven independent source artifacts stay unchanged.
    identifiers = {value: str(uuid4()) for value in set(re.findall(r'[0-9a-f]{8}(?:-[0-9a-f]{4}){3}-[0-9a-f]{12}', json.dumps(evidence)))}
    for deliverable in evidence['deliverables']:
        for file in deliverable['files']:
            for source in file.get('sources', []):
                access = source.get('access', {})
                if access.get('actorId'):
                    identifiers[access['actorId']] = actor.id
                if access.get('companyId'):
                    identifiers[access['companyId']] = actor.company_id
    def remap(value):
        if isinstance(value, dict):
            return {remap(key): remap(item) for key, item in value.items()}
        if isinstance(value, list):
            return [remap(item) for item in value]
        if isinstance(value, str):
            return re.sub(r'[0-9a-f]{8}(?:-[0-9a-f]{4}){3}-[0-9a-f]{12}', lambda match: identifiers.get(match[0], match[0]), value)
        return value
    data = remap(copy.deepcopy(evidence))
    cutoff = 32 if case['seed'] == 'production-first' else 33
    history = data['messages'][:cutoff]
    by_message = {item['id']: item for item in history}
    all_messages = {item['id']: item for item in data['messages']}
    by_job = {item['target_id']: item for item in data['jobs']}
    source_revisions = {}
    for deliverable in data['deliverables']:
        for file in deliverable['files']:
            for ref in file.get('sources', []):
                if ref.get('kind') == 'attachment':
                    source_revisions[ref['id']] = ref.get('revision', 0)
    owned = {'company_id': actor.company_id, 'owner_id': actor.id}
    def date(value):
        return datetime.fromisoformat(value)
    def repair_sources(value):
        if isinstance(value, dict):
            if value.get('messageId') in all_messages and 'quote' in value and 'digest' in value:
                origin = all_messages[value['messageId']]
                # The fixture keeps text and transcript unchanged after redaction;
                # digests must authenticate those new fixture bytes, not originals.
                value['digest'] = digest(origin['text'] or origin.get('transcript', ''))
            for child in value.values():
                repair_sources(child)
        elif isinstance(value, list):
            for child in value:
                repair_sources(child)
    repair_sources(data)
    attachment_checksums = {}
    async with sessions.begin() as db:
        for original in history:
            row = Message(**owned, id=original['id'], conversation_id=state['conversationId'], created_at=date(original['created_at']), text=original['text'], reply=original['reply'], transcript=original.get('transcript', ''), persona_id=original.get('persona_id', 'dabao'), work_reference=original.get('work_reference', {}), deliverable_reference=original.get('deliverable_reference', {}), reply_to=original.get('reply_to'))
            db.add(row)
        await db.flush()
        for item in data['attachments']:
            if item.get('message_id') and item['message_id'] not in by_message:
                continue
            name = item['name']
            raw = sample_file(name, '隔离测试原始材料：' + Path(name).stem + '。项目预算12800元，交付日期2026-11-06。')
            checksum = hashlib.sha256(raw).hexdigest()
            attachment_checksums[item['id']] = checksum
            path = settings.media_dir / item['id']
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(raw)
            revision = source_revisions.get(item['id'], 1)
            db.add(Attachment(**owned, id=item['id'], message_id=item.get('message_id'), name=name, kind=item['kind'], mime=item['mime'], size=len(raw), sha256=checksum, extraction_status='ready' if revision else 'pending', extraction_revision=revision, created_at=date(item['created_at'])))
            if revision:
                await db.flush()
                db.add(DocumentChunk(**owned, attachment_id=item['id'], revision=revision, ordinal=0, location='隔离测试文档', text='隔离测试材料：' + Path(name).stem))
        await db.flush()
        work = data['work']
        db.add(WorkItem(**owned, id=work['id'], created_at=date(work['created_at']), title=work['title'], content=work['content'], revision=work['revision'], origin='manual'))
        await db.flush()
        for rev in data['workRevisions']:
            db.add(WorkRevision(**owned, work_id=work['id'], revision=rev['revision'], content=rev['content'], source_ids=rev['source_ids'], created_at=date(rev['created_at']), origin='manual'))
        for material in data['materials']:
            db.add(WorkMaterial(**owned, work_id=work['id'], id=material['id'], attachment_id=material['attachment_id'], source=material['source'], removed=material['removed'], revision=material['revision'], created_at=date(material['created_at'])))
        state['workId'], state['workRevision'] = work['id'], work['revision']
        state['artifacts'] = []
        for item in data['deliverables']:
            db.add(Deliverable(**owned, id=item['id'], conversation_id=state['conversationId'], title=item['title'], revision=item['revision'], created_at=date(item['created_at']), updated_at=date(item['created_at'])))
            await db.flush()
            files = []
            for file in item['files']:
                raw = sample_file(file['name'], '隔离测试成果：' + item['title'] + '。不使用真实业务内容。')
                entry = {**file, 'sha256': hashlib.sha256(raw).hexdigest(), 'size': len(raw), 'storageKey': hashlib.sha256((actor.id + file['id']).encode()).hexdigest()}
                for ref in entry.get('sources', []):
                    if ref.get('kind') == 'attachment' and ref['id'] in attachment_checksums:
                        ref['sha256'] = attachment_checksums[ref['id']]
                path = file_path(settings, entry['storageKey'])
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(raw)
                files.append(entry)
                state['artifacts'].append({'deliverableId': item['id'], 'revision': item['source_revision'], 'fileId': file['id'], 'name': file['name']})
            db.add(DeliverableRevision(**owned, deliverable_id=item['id'], revision=item['source_revision'], message_id=item['message_id'], step=item['step'], title=item['revision_title'], body=item['body'], items=item['items'], files=files, digest=hashlib.sha256(json.dumps(files, sort_keys=True).encode()).hexdigest(), created_at=date(item['created_at'])))
        for message in history:
            original = by_job[message['id']]
            db.add(Job(**owned, id=original['id'], target_id=message['id'], kind='message', state=original['state'], phase=original['phase'], created_at=date(original['created_at']), updated_at=date(original['created_at']), attempt=original['attempt'], result=original['result']))
        next_job = by_job[data['messages'][cutoff]['id']]
        snapshot = next_job['result'].get('taskSnapshot', {})
        db.add(ConversationTaskState(**owned, conversation_id=state['conversationId'], revision=max(1, snapshot.get('version', 2)-1), payload={'latestMessageId': history[-1]['id'], 'directives': snapshot.get('directives', []), 'task': snapshot.get('previousTask', {})}))
    state['differences'] = data['fixtureDisclosure']['changes'] + ['Source bytes replaced by valid short synthetic documents; historical conversation text remains anonymized original. IDs are per-run remapped with stable relationships; authenticated digest/SHA values recomputed. No browser or production network queue timing is simulated.']
    state['historyCount'] = len(history)
    state['historyChars'] = sum(len(item['text']) + len(item['reply']) for item in history)
    state['fixtureSha256'] = hashlib.sha256(json.dumps(evidence, sort_keys=True).encode()).hexdigest()



async def run(args):
    from dotenv import load_dotenv
    load_dotenv(ROOT / 'apps/server/.env.web')
    source = args.app_root.resolve()
    sys.path[:0] = [str(source / 'apps/server'), str(source / 'tests/server'), str(source / 'packages/voiceprint-engine/src')]
    from app.core.config import Settings
    from sqlalchemy import select
    from app.tasks.models import Job
    from app.modules.model_services.models import ModelUsage
    from app.tasks.processing.handlers import process_job
    from app.tasks.runtime.queue import claim
    from app.tasks.streaming.redis_store import RedisReplyStore
    from app.integrations.models import chat as chat_module
    from langgraph.checkpoint.postgres.aio import AsyncPostgresSaver
    from pwdlib import PasswordHash
    import psycopg
    manifest = json.loads(CASE_FILE.read_text())
    cases = {item['id']: item for item in manifest['cases']}
    wanted = args.ids.split(',')
    efforts = args.efforts.split(',')
    samples = [int(value) for value in args.samples.split(',')]
    assert set(wanted) <= set(cases) and set(efforts) <= {'low', 'high'}
    assert all(0 <= sample < cases[identifier]['samples'] for identifier in wanted for sample in samples)
    evidence = json.loads(args.evidence.read_text()) if args.evidence else None
    args.output.mkdir(parents=True, exist_ok=True)
    output = args.output / 'results.jsonl'
    if output.exists() and not args.resume:
        raise RuntimeError('Results exist; never overwrite or rerun failures silently')
    if args.serve and (args.output / 'ui-fixture.json').exists():
        raise RuntimeError('UI fixture output already exists; use a new run directory')
    done = {(row['case'], row['effort'], row['sample']) for row in map(json.loads, output.read_text().splitlines())} if output.exists() else set()
    config = await configuration(args)
    if not set(efforts) <= set(config['model'].get('reasoning', {}).get('efforts', [])):
        raise RuntimeError('Configured model does not expose requested efforts')
    digest = source_digest(source)
    schema = 'spec075_' + args.label + '_' + uuid4().hex[:10]
    with psycopg.connect(DATABASE, autocommit=True) as db:
        db.execute('CREATE SCHEMA ' + schema)
    test_url = DATABASE + '?options=-csearch_path%3D' + schema
    env = {**os.environ, 'DATABASE_URL': test_url, 'PYTHONPATH': str(source / 'apps/server')}
    migration = await asyncio.create_subprocess_exec(sys.executable, '-m', 'app.cli', 'migrate', cwd=source, env=env, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
    _, error = await migration.communicate()
    if migration.returncode:
        raise RuntimeError('Owned migration failed: ' + error.decode()[-1000:])
    os.environ['DATABASE_TEST_URL'] = test_url
    os.environ['PAA_REPLY_STREAM_REDIS_URL'] = 'redis://127.0.0.1:6379/0'
    os.environ['PAA_REPLY_STREAM_NAMESPACE'] = schema
    spec = importlib.util.spec_from_file_location('spec075_fixture', source / 'tests/server/conftest.py')
    fixture_module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fixture_module)
    environment = {'source': str(source), 'digest': digest, 'label': args.label, 'manifestSha256': hashlib.sha256(CASE_FILE.read_bytes()).hexdigest(), 'canonicalManifestSha256': hashlib.sha256(json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()).hexdigest(), 'started': time.time(), 'declaredConcurrentLiveProcesses': args.concurrent_live_processes, 'model': config['model'], 'selectedParameters': config['choice'].get('parameters', {}), 'providerHost': __import__('urllib.parse', fromlist=['urlparse']).urlparse(config['base']).hostname, 'schema': schema, 'loadedIntent': __import__('app.agent.actions.intent', fromlist=['authorize_intent']).__file__, 'timingScope': 'real in-process HTTP admission and worker; first public Redis event, excludes browser/network paint and periodic queue polling', 'planned': [{'case': identifier, 'effort': effort, 'sample': sample} for sample in samples for effort in efforts for identifier in wanted]}
    environment['failureDiagnostics'] = {'version': 1, 'trigger': 'tool_arguments_json',
        'capture': 'failed tool name/arguments and existing on_delta tool fragments only',
        'changesModelInput': False, 'changesParsingOrRetry': False,
        'oldSamplesUnchanged': True}
    environment['publicObservation'] = {'version': 2, 'questionBoundary': 'ask returned after transaction commit',
        'captures': 'persisted public questions, direct caller identity and provider public content only',
        'changesModelInput': False, 'extraTimedQueries': 0, 'oldSamplesUnchanged': True}
    environment['captureRequestPolicy'] = args.capture_request_policy
    write(args.output / ('environment-' + schema + '.json'), environment)
    timings = Timings()
    logger = logging.getLogger('paa.latency')
    logger.handlers[:] = [timings]
    logger.setLevel(logging.INFO)
    logger.propagate = False
    active = None
    original_chat = chat_module.chat
    original_invalid_diagnostic = chat_module.invalid_response_diagnostic
    diagnostic_request = ContextVar('spec075_failed_tool_arguments', default=None)
    def observed_invalid_diagnostic(result, reason, tool_index, on_timing, **metadata):
        original_invalid_diagnostic(result, reason, tool_index, on_timing, **metadata)
        context = diagnostic_request.get()
        if reason != 'tool_arguments_json' or context is None or tool_index is None:
            return
        row, observation, fragments, ordinal = context
        began = time.perf_counter()
        try:
            detail = failed_tool_arguments(result, tool_index, metadata, fragments)
            if detail is None:
                return
            folder = args.output / 'failed-tool-arguments'
            folder.mkdir(mode=0o700, exist_ok=True)
            path = folder / f"{row['case']}-{row['effort']}-{row['sample']}-{ordinal}.json"
            with os.fdopen(os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600), 'w') as stream:
                json.dump(detail, stream, ensure_ascii=False, indent=2)
                stream.write('\n')
            observation['failedToolArgumentsFile'] = str(path.resolve())
        except Exception as error:
            # Observation failures must not change parsing, retries or execution.
            observation['failedToolArgumentsObservationError'] = type(error).__name__
        finally:
            observation['failedToolArgumentsObservationMs'] = (time.perf_counter()-began)*1000
    chat_module.invalid_response_diagnostic = observed_invalid_diagnostic
    from app.tasks.timing import call_kind
    from app.agent.actions import interactions as interactions_module
    question_aliases = observe_questions(interactions_module, lambda: active)
    async def observed_chat(settings, wire, key, messages, **kwargs):
        row = active
        if row is None:
            return await original_chat(settings, wire, key, messages, **kwargs)
        at = time.perf_counter()
        input_text = json.dumps(messages, ensure_ascii=False, default=str)
        observation = {'start_ms': (at-row['_clock'])*1000, 'model': wire['model'], 'parameters': wire.get('parameters', {}), 'kind': call_kind.get(), 'inputChars': len(input_text), 'inputSha256': hashlib.sha256(input_text.encode()).hexdigest()}
        if args.capture_request_policy:
            observation['requestPolicy'] = captured_request_policy(messages, kwargs.get('tools'), kwargs.get('tool_choice'))
        row['requests'].append(observation)
        original_delta = kwargs.get('on_delta')
        fragments = [] if original_delta is not None else None
        if original_delta is not None:
            async def observed_delta(delta):
                try:
                    fragments.extend({'index': item.index, 'name': item.name, 'arguments': item.arguments} for item in delta.tools)
                except Exception as error:
                    observation['toolFragmentObservationError'] = type(error).__name__
                return await original_delta(delta)
            kwargs['on_delta'] = observed_delta
        diagnostic_token = diagnostic_request.set((row, observation, fragments, len(row['requests'])))
        try:
            result = await original_chat(settings, wire, key, messages, **kwargs)
            observation['usage'] = result.get('usage', {})
            observation['toolCalls'] = [call.get('function', {}) for choice in result.get('choices', []) for call in choice.get('message', {}).get('tool_calls', [])]
            observation['publicContent'] = [{'choiceIndex': index, 'text': choice['message']['content']}
                for index, choice in enumerate(result.get('choices', []))
                if isinstance(choice.get('message', {}).get('content'), str) and choice['message']['content'].strip()]
            return result
        except BaseException as error:
            observation['error'] = {'type': type(error).__name__, 'code': getattr(error, 'code', None), 'status': getattr(error, 'status', None)}
            raise
        finally:
            observation['elapsed_ms'] = (time.perf_counter()-at)*1000
            diagnostic_request.reset(diagnostic_token)
    chat_module.chat = observed_chat
    original_publish = RedisReplyStore.publish
    async def observed_publish(store, identity, snapshot, event):
        accepted = await original_publish(store, identity, snapshot, event)
        if accepted and active is not None and snapshot.get('text', '').strip():
            at = (time.perf_counter()-active['_clock'])*1000
            active['stream'].append({'at_ms': at, 'text': snapshot['text'], 'state': snapshot['state']})
            if active['first_answer_ms'] is None:
                active['first_answer_ms'] = at
        return accepted
    RedisReplyStore.publish = observed_publish
    password = PasswordHash.recommended().hash('controlled-test-password')
    if args.serve:
        spec = importlib.util.spec_from_file_location('spec075_ui', Path(__file__).with_name('agent-execution-feedback-ui.py'))
        ui_module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(ui_module)
        def activate(row):
            nonlocal active
            active = row
        try:
            await ui_module.serve(args, sys.modules[__name__], fixture_module, config, cases,
                                  evidence, password, schema, digest, timings, activate)
        finally:
            active = None
            chat_module.chat = original_chat
            chat_module.invalid_response_diagnostic = original_invalid_diagnostic
            RedisReplyStore.publish = original_publish
            restore_aliases(question_aliases)
        return
    try:
        for sample in samples:
            for effort in efforts:
                for identifier in wanted:
                    if (identifier, effort, sample) in done:
                        continue
                    assert source_digest(source) == digest, 'Source changed during benchmark; freeze it before comparing'
                    case = cases[identifier]
                    row = {'case': identifier, 'effort': effort, 'sample': sample, 'label': args.label, 'prompt': case['prompt'], 'expect': case['expect'], 'requests': [], 'stream': [], 'first_answer_ms': None, 'publicObservationVersion': 2, 'started': time.time()}
                    fixture = None
                    with tempfile.TemporaryDirectory(prefix='noria-spec075-') as temporary:
                        try:
                            fixture = fixture_module.setup.__wrapped__(Path(temporary), password)
                            settings, sessions, users, clients = await anext(fixture)
                            actor, client = users['employee'], clients['employee']
                            await install(config, settings, sessions, actor.company_id)
                            seeded = await seed(case, settings, sessions, actor, client, clients, evidence)
                            row['seed'] = seeded
                            row['initial'] = await snapshot(sessions, actor)
                            if case['seed'] == 'private':
                                row['peerWorkInitial'] = checked(await clients['peer'].get('/api/v1/work-items/' + seeded['privateWorkId']))
                            if args.prepare_only:
                                row['preparedOnly'] = True
                                continue
                            preference = checked(await client.get('/api/v1/assistant/reasoning-settings'))
                            selected = checked(await client.patch('/api/v1/assistant/reasoning-settings', json={'effort': effort, 'preferenceRevision': preference['preferenceRevision'], 'configVersion': preference['configVersion']}))
                            body = {'text': case['prompt'].replace('{{privateWorkId}}', seeded.get('privateWorkId', '')), 'conversationId': seeded['conversationId'], 'attachmentIds': seeded['attachmentIds'], 'reasoning': {'effort': effort, 'preferenceRevision': selected['preferenceRevision'], 'configVersion': selected['configVersion']}}
                            if case.get('workReference'):
                                body['workReference'] = {'workId': seeded['workId']}
                            row['_clock'] = time.perf_counter()
                            active = row
                            receipt = checked(await client.post('/api/v1/messages', json=body, headers=(headers := keyed())), 202)
                            row['send_ms'] = (time.perf_counter()-row['_clock'])*1000
                            if case.get('duplicateSubmission'):
                                duplicate = checked(await client.post('/api/v1/messages', json=body, headers=headers), 202)
                                row['duplicateSameReceipt'] = duplicate == receipt
                            job = await claim(sessions, actor.id)
                            if not job or job.id != receipt['jobId']:
                                raise RuntimeError('Worker did not claim exact submitted message')
                            row['queue_ms'] = (time.perf_counter()-row['_clock'])*1000-row['send_ms']
                            async with AsyncPostgresSaver.from_conn_string(settings.checkpoint_url) as saver:
                                await process_job(job, sessions, settings, saver)
                            dto = checked(await client.get('/api/v1/messages/' + receipt['messageId']))
                            row['total_ms'] = (time.perf_counter()-row['_clock'])*1000
                            row['answer'] = dto['reply']
                            row['interactions'] = checked(await client.get('/api/v1/conversations/' + seeded['conversationId'] + '/interactions'))['items']
                            row['message'] = dto
                            row['final'] = await snapshot(sessions, actor)
                            if case['seed'] == 'private':
                                row['peerWorkFinal'] = checked(await clients['peer'].get('/api/v1/work-items/' + seeded['privateWorkId']))
                            async with sessions() as db:
                                current = await db.get(Job, job.id)
                                row['jobResult'] = current.result
                                usages = (await db.scalars(select(ModelUsage).where(ModelUsage.job_id == job.id).order_by(ModelUsage.created_at))).all()
                                row['usage'] = [{'kind': item.kind, 'state': item.status, 'ms': item.elapsed_ms, 'input': item.actual_input_tokens, 'output': item.actual_output_tokens, 'error': item.error_code} for item in usages]
                            row['events'] = [event for event in timings.events if event.get('job') == job.id]
                            row['refreshSameReply'] = checked(await client.get('/api/v1/messages/' + receipt['messageId']))['reply'] == dto['reply']
                            if row['first_answer_ms'] is None and dto['reply']:
                                row['first_answer_ms'] = row['total_ms']
                                row['first_answer_source'] = 'final API snapshot only'
                            else:
                                row['first_answer_source'] = 'public Redis event'
                            complete_public_observation(row)
                        except Exception as error:
                            row['error'] = {'type': type(error).__name__, 'detail': str(error)[:600]}
                            row['total_ms'] = (time.perf_counter()-row['_clock'])*1000 if '_clock' in row else None
                        finally:
                            active = None
                            row.pop('_clock', None)
                            append(output, row)
                            print(json.dumps({'case': identifier, 'effort': effort, 'sample': sample, 'label': args.label, 'state': row.get('message', {}).get('job', {}).get('state'), 'total_ms': row.get('total_ms'), 'calls': len(row['requests']), 'error': row.get('error')}, ensure_ascii=False), flush=True)
                            if fixture:
                                await fixture.aclose()
    finally:
        chat_module.chat = original_chat
        chat_module.invalid_response_diagnostic = original_invalid_diagnostic
        restore_aliases(question_aliases)
        RedisReplyStore.publish = original_publish


if __name__ == '__main__':
    logging.basicConfig(level=logging.WARNING)
    asyncio.run(run(arguments()))
