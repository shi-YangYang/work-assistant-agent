"""One real file task while another synthetic member deletes an unrelated work."""
import argparse
import asyncio
import io
import json
import time
from pathlib import Path
from uuid import uuid4
import httpx
from pptx import Presentation


async def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--state', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    private = Path(args.state)
    if private.stat().st_mode & 0o077:
        raise RuntimeError('Private credentials required')
    state = json.loads(private.read_text())
    if not state['schema'].startswith('spec044_web_') or state['apiOrigin'] != 'http://127.0.0.1:8016':
        raise RuntimeError('Isolated API only')
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=False)
    result = {'scope': 'C05 repair; one real user message, one unrelated synthetic fixture deletion', 'status': 'running'}

    def save():
        (output / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')

    async def login(client, role):
        user = state['users'][role]
        response = await client.post('/api/v1/auth/login', json={'username': user['username'], 'password': user['password']})
        response.raise_for_status()
        client.headers['X-CSRF-Token'] = response.json()['csrf']

    options = {'base_url': state['apiOrigin'], 'headers': {'Origin': state['origin']}, 'timeout': 30}
    async with httpx.AsyncClient(**options) as client, httpx.AsyncClient(**options) as peer:
        await login(client, 'realEmployee')
        await login(peer, 'realApiEmployee')
        try:
            fixture = await peer.post('/api/v1/work-items', headers={'Idempotency-Key': str(uuid4())}, json={
                'title': 'C05专用无关删除夹具', 'summary': '测试准备，不算Agent创建成果',
                'status': 'in_progress', 'blocker': '', 'nextStep': '',
            })
            fixture.raise_for_status()
            result['fixture'] = {'work': fixture.json(), 'countsAsAgentSuccess': False}
            case = next(row for row in json.loads(Path(__file__).with_name('cases.json').read_text()) if row['id'] == 'C05')
            prompt = case['turns'][0]
            sent = await client.post('/api/v1/messages', headers={'Idempotency-Key': str(uuid4())}, json={
                'text': prompt, 'attachmentIds': [], 'newConversation': True,
                'personaId': 'professional', 'executionMode': 'full', 'fullAccessConfirmed': True,
            })
            sent.raise_for_status()
            result.update(prompt=prompt, receipt=sent.json())
            save()
            started = time.monotonic()
            deleted = False
            async with asyncio.timeout(300):
                while True:
                    response = await client.get('/api/v1/messages/' + result['receipt']['messageId'])
                    response.raise_for_status()
                    message = response.json()
                    job = message['job']
                    if not deleted and job['state'] == 'running' and job.get('nodes'):
                        work = result['fixture']['work']
                        removed = await peer.request('DELETE', '/api/v1/work-items/' + work['id'], json={'expectedRevision': work['revision']})
                        removed.raise_for_status()
                        result['deletion'] = {'status': removed.status_code, 'duringJobState': job['state'], 'nodes': job['nodes']}
                        deleted = True
                        save()
                    if job['state'] in ('succeeded', 'awaiting_input', 'awaiting_retry', 'cancelled', 'failed'):
                        break
                    await asyncio.sleep(.5)
            result.update(message=message, seconds=round(time.monotonic() - started, 2))
            files = []
            for delivery in message.get('deliverables', []):
                for item in delivery.get('files', []):
                    if not item['url'].startswith('/api/v1/'):
                        raise RuntimeError('Unexpected download URL')
                    response = await client.get(item['url'])
                    response.raise_for_status()
                    name = Path(item['name']).name
                    (output / name).write_bytes(response.content)
                    if name.lower().endswith('.pptx'):
                        slides = Presentation(io.BytesIO(response.content)).slides
                        files.append({'name': name, 'slides': [
                            '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
                            for slide in slides
                        ]})
            result['files'] = files
            assert deleted, 'Deletion was not injected during the real task'
            assert job['state'] in ('succeeded', 'awaiting_input'), job
            assert len(files) == 1 and len(files[0]['slides']) == 2
            assert '上线' in files[0]['slides'][0]
            assert all(word in files[0]['slides'][1] for word in ('测试', '培训', '反馈'))
            assert not message.get('actions'), 'A file request must not create business actions'
            result['status'] = 'PASS'
        except Exception as error:
            result.update(status='FAIL', error=f'{type(error).__name__}: {error}')
            raise
        finally:
            save()
            await client.post('/api/v1/auth/logout', json={})
            await peer.post('/api/v1/auth/logout', json={})
    print(json.dumps({'status': result['status'], 'seconds': result['seconds'], 'files': files}, ensure_ascii=False))


if __name__ == '__main__':
    asyncio.run(main())
