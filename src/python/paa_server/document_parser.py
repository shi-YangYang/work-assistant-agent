"""Standalone, resource-bounded child entry point. Never imports service configuration."""
import csv
import io
import json
from pathlib import Path
import sys
import zipfile

PARSER_VERSION = 'text-v1:pypdf6.18.1:docx1.2.0:pptx1.0.2'
MAX_TEXT = 200_000
MAX_CHUNKS = 2000
MAX_PAGES = 100
MAX_EXPANDED = 64 * 1024 * 1024
MEMORY_BYTES = 512 * 1024 * 1024


class ParseFailure(ValueError):
    pass


class Full(Exception):
    pass


_WINDOWS_JOB = None


def windows_resource_limits():
    """Windows has no resource module. Bind this child to a private Job Object
    before loading parsers; the kernel enforces process memory/CPU/process count.
    The handle remains open for this process's lifetime (kill on close).
    """
    import ctypes
    from ctypes import wintypes
    global _WINDOWS_JOB
    class BasicLimits(ctypes.Structure):
        _fields_ = [
            ('PerProcessUserTimeLimit', ctypes.c_int64), ('PerJobUserTimeLimit', ctypes.c_int64),
            ('LimitFlags', ctypes.c_uint32), ('MinimumWorkingSetSize', ctypes.c_size_t),
            ('MaximumWorkingSetSize', ctypes.c_size_t), ('ActiveProcessLimit', ctypes.c_uint32),
            ('Affinity', ctypes.c_size_t), ('PriorityClass', ctypes.c_uint32), ('SchedulingClass', ctypes.c_uint32),
        ]
    class IOCounters(ctypes.Structure):
        _fields_ = [(name, ctypes.c_uint64) for name in ('ReadOperationCount', 'WriteOperationCount', 'OtherOperationCount', 'ReadTransferCount', 'WriteTransferCount', 'OtherTransferCount')]
    class ExtendedLimits(ctypes.Structure):
        _fields_ = [('BasicLimitInformation', BasicLimits), ('IoInfo', IOCounters),
            ('ProcessMemoryLimit', ctypes.c_size_t), ('JobMemoryLimit', ctypes.c_size_t),
            ('PeakProcessMemoryUsed', ctypes.c_size_t), ('PeakJobMemoryUsed', ctypes.c_size_t)]
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.CreateJobObjectW.argtypes, kernel.CreateJobObjectW.restype = [ctypes.c_void_p, wintypes.LPCWSTR], wintypes.HANDLE
    kernel.SetInformationJobObject.argtypes, kernel.SetInformationJobObject.restype = [wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p, ctypes.c_uint32], wintypes.BOOL
    kernel.AssignProcessToJobObject.argtypes, kernel.AssignProcessToJobObject.restype = [wintypes.HANDLE, wintypes.HANDLE], wintypes.BOOL
    kernel.GetCurrentProcess.argtypes, kernel.GetCurrentProcess.restype = [], wintypes.HANDLE
    kernel.CloseHandle.argtypes, kernel.CloseHandle.restype = [wintypes.HANDLE], wintypes.BOOL
    handle = kernel.CreateJobObjectW(None, None)
    if not handle:
        raise OSError('无法创建文档解析资源限制')
    limits = ExtendedLimits()
    limits.BasicLimitInformation.LimitFlags = 0x2 | 0x8 | 0x100 | 0x2000
    limits.BasicLimitInformation.PerProcessUserTimeLimit = 45 * 10_000_000
    limits.BasicLimitInformation.ActiveProcessLimit = 1
    limits.ProcessMemoryLimit = MEMORY_BYTES
    if not kernel.SetInformationJobObject(handle, 9, ctypes.byref(limits), ctypes.sizeof(limits)) or not kernel.AssignProcessToJobObject(handle, kernel.GetCurrentProcess()):
        kernel.CloseHandle(handle)
        raise OSError('无法启用文档解析资源限制')
    _WINDOWS_JOB = handle


def resource_limits():
    if sys.platform == 'win32':
        import os
        import threading
        windows_resource_limits()
        timer = threading.Timer(60, lambda: os._exit(124))
        timer.daemon = True
        timer.start()
    else:
        import resource
        import signal
        signal.setitimer(signal.ITIMER_REAL, 60)
        if sys.platform != 'darwin':
            resource.setrlimit(resource.RLIMIT_AS, (MEMORY_BYTES, MEMORY_BYTES))
        # macOS rejects RLIMIT_AS/DATA/RSS; the parent enforces an RSS watchdog.
        resource.setrlimit(resource.RLIMIT_CPU, (45, 46))
        resource.setrlimit(resource.RLIMIT_FSIZE, (4 * 1024 * 1024, 4 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
    # Parsers never fetch external relationships. Deny Python network/process
    # APIs too; -I and the minimal child environment keep keys/config out.
    def audit(event, args):
        if event.startswith(('socket.', 'subprocess.')) or event in ('os.system', 'os.exec', 'os.posix_spawn', 'os.fork'):
            raise PermissionError('document parser cannot access network or execute programs')
    sys.addaudithook(audit)


def office_check(path, suffix):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        names = {entry.filename for entry in entries}
        required = 'word/document.xml' if suffix == '.docx' else 'ppt/presentation.xml'
        if required not in names or '[Content_Types].xml' not in names:
            raise ParseFailure('文件结构与扩展名不符，请保存为正确的 DOCX／PPTX')
        if len(entries) > 2048 or len(names) != len(entries) or sum(e.file_size for e in entries) > MAX_EXPANDED:
            raise ParseFailure('文档解压后过大或条目过多，请拆分文件')
        for entry in entries:
            name = entry.filename.lower()
            if entry.flag_bits & 1:
                raise ParseFailure('不支持加密文档，请先解密')
            if entry.file_size > 16 * 1024 * 1024 or entry.file_size > max(1024 * 1024, entry.compress_size * 200):
                raise ParseFailure('文档压缩比例或单个内容过大，请拆分文件')
            if 'vbaproject' in name or name.endswith('.bin') and ('embeddings/' in name or 'activex/' in name):
                raise ParseFailure('文档包含宏或活动嵌入对象，请移除后重新保存')
            if name.endswith(('.xml', '.rels')):
                from lxml import etree
                raw = archive.read(entry)
                tree = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False))
                if tree.getroottree().docinfo.doctype:
                    raise ParseFailure('文档包含不支持的 XML 实体声明')


def text_decode(data):
    try:
        text = data.decode('utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')
    except UnicodeError:
        raise ParseFailure('无法识别文本编码，请保存为 UTF-8 或带 BOM 的 UTF-16') from None
    if any(ord(c) < 32 and c not in '\n\r\t' for c in text):
        raise ParseFailure('文件含有二进制或不支持的控制字符，请保存为纯文本')
    return text


def extract(path, suffix):
    chunks, warnings = [], []
    count = 0
    def add(location, text):
        nonlocal count
        if not text.strip():
            return
        for start in range(0, len(text), 3000):
            if count >= MAX_TEXT or len(chunks) >= MAX_CHUNKS:
                raise Full()
            part = text[start:start + min(3000, MAX_TEXT - count)]
            chunks.append({'location': (location + (f' · 字符 {start + 1}–{start + len(part)}' if len(text) > 3000 else ''))[:300], 'text': part})
            count += len(part)
            if start + len(part) < len(text) and count >= MAX_TEXT:
                raise Full()
    try:
        if suffix in ('.docx', '.pptx'):
            office_check(path, suffix)
        if suffix == '.pdf':
            from pypdf import PdfReader, overwrite_configuration
            from pypdf.errors import LimitReachedError
            overwrite_configuration(zlib_maximum_output_length=MAX_EXPANDED, lzw_maximum_output_length=MAX_EXPANDED, run_length_maximum_output_length=MAX_EXPANDED, array_based_stream_maximum_output_length=MAX_EXPANDED, disable_legacy_handling=True)
            reader = PdfReader(path, strict=True)
            if reader.is_encrypted:
                raise ParseFailure('不支持加密 PDF，请先解密')
            if len(reader.pages) > MAX_PAGES:
                warnings.append('仅解析前 100 页，后续页面未读取')
            missing = []
            for i, page in enumerate(reader.pages[:MAX_PAGES], 1):
                try:
                    text = page.extract_text() or ''
                except LimitReachedError:
                    raise ParseFailure('PDF 内容解压后超过限制，请拆分文件') from None
                if not text.strip():
                    missing.append(str(i))
                add(f'第 {i} 页', text)
            if missing:
                warnings.append('以下页面无可提取文字（可能为扫描件）：' + '、'.join(missing))
        elif suffix == '.docx':
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            document = Document(path)
            paragraphs = tables = 0
            for element in document.iter_inner_content():
                if isinstance(element, Paragraph):
                    paragraphs += 1
                    add(f'段落 {paragraphs}', element.text)
                elif isinstance(element, Table):
                    tables += 1
                    for rownum, row in enumerate(element.rows, 1):
                        add(f'表格 {tables} · 行 {rownum}', '\n'.join(f'列 {i}：{cell.text}' for i, cell in enumerate(row.cells, 1)))
        elif suffix == '.pptx':
            from pptx import Presentation
            slides = Presentation(path).slides
            if len(slides) > MAX_PAGES:
                warnings.append('仅解析前 100 张幻灯片，后续未读取')
            def shape_text(shapes):
                for shape in shapes:
                    if hasattr(shape, 'shapes'):
                        yield from shape_text(shape.shapes)
                    if shape.has_text_frame:
                        yield shape.text
                    if shape.has_table:
                        for n, row in enumerate(shape.table.rows, 1):
                            yield f'表格行 {n}：' + ' | '.join(f'列 {i}：{cell.text}' for i, cell in enumerate(row.cells, 1))
            for i, slide in enumerate(slides, 1):
                if i > MAX_PAGES:
                    break
                add(f'幻灯片 {i}', '\n'.join(shape_text(slide.shapes)))
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame
                    if notes:
                        add(f'幻灯片 {i} · 演讲备注', notes.text)
        else:
            text = text_decode(Path(path).read_bytes())
            if suffix == '.json':
                def invalid(value):
                    raise ValueError('invalid constant')
                try:
                    value = json.loads(text, parse_constant=invalid)
                except (ValueError, RecursionError):
                    raise ParseFailure('JSON 语法无效，请修正后重新上传') from None
                def visit(value, pointer='', depth=0):
                    if depth > 100:
                        raise ParseFailure('JSON 嵌套超过 100 层，请简化结构')
                    if isinstance(value, (dict, list)) and value:
                        add('JSON ' + (pointer or '/'), '{对象}' if isinstance(value, dict) else '[数组]')
                        for key, child in (value.items() if isinstance(value, dict) else enumerate(value)):
                            part = str(key).replace('~', '~0').replace('/', '~1')
                            visit(child, pointer + '/' + part, depth + 1)
                    else:
                        add('JSON ' + (pointer or '/'), json.dumps(value, ensure_ascii=False))
                visit(value)
            elif suffix == '.csv':
                csv.field_size_limit(MAX_TEXT)
                reader = csv.reader(io.StringIO(text, newline=''), strict=True)
                for n, row in enumerate(reader, 1):
                    add(f'记录 {n} · 至行 {reader.line_num}', '\n'.join(f'列 {i}：{value}' for i, value in enumerate(row, 1)))
            elif suffix in ('.txt', '.md'):
                # Group bounded consecutive lines, retaining physical positions.
                lines = text.splitlines()
                for start in range(0, len(lines), 30):
                    add(f'行 {start + 1}–{min(start + 30, len(lines))}', '\n'.join(lines[start:start + 30]))
            else:
                raise ParseFailure('不支持的文档格式')
    except Full:
        warnings.append('已达到 20 万字符或 2000 分段上限，剩余内容未读取')
    if not chunks:
        raise ParseFailure('没有可提取文字；扫描件或纯图片文档需转为带文字层的文件，本轮不做 OCR')
    return {'status': 'partial' if warnings else 'ready', 'chunks': chunks, 'info': {'warnings': warnings, 'characters': count, 'chunks': len(chunks), 'scope': '仅提取文字，不含图片、图表视觉内容及完整原始版式'}}


def main():
    resource_limits()
    try:
        result = extract(sys.argv[1], sys.argv[2])
    except ParseFailure as error:
        result = {'status': 'failed', 'chunks': [], 'info': {'error': str(error)}}
    except (MemoryError, RecursionError):
        result = {'status': 'failed', 'chunks': [], 'info': {'error': '文档超出解析资源限制，请拆分文件'}}
    except Exception:
        result = {'status': 'failed', 'chunks': [], 'info': {'error': '文件损坏或结构无法读取，请重新导出后上传'}}
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
